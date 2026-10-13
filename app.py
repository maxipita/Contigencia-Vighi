import io
import json
import os
import re
import secrets
import socket
import sys
import unicodedata
from datetime import datetime, timedelta
from functools import wraps
from urllib.parse import quote

from flask import (Flask, abort, flash, g, jsonify, redirect, render_template, request, send_file,
                   session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash

import d1
import db
import envios
import facturacion
import permisos
import trazabilidad
import ubicaciones as ubicaciones_csv
from estudios import NOMBRE_ETAPA, PASOS_IHQ, REGISTRO, SECTORES, TIPOS, Variante, con_decal, validar_combinacion

try:
    import informes
except ImportError as _err:
    informes = None
    print(f"AVISO: los informes en PDF están desactivados ({_err}). Instalar con: {sys.executable} -m pip install -r requirements.txt")
SEMAFOROS = {"on_time": "On time", "delayed": "Delayed", "late": "Late"}
SIN_INFORMES = "Los informes en PDF no están disponibles: falta instalar reportlab (python -m pip install -r requirements.txt) y reiniciar el sistema."

PUERTO = int(os.environ.get("CONTINGENCIA_PUERTO", "8000"))

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 40 * 1024 * 1024


def _clave_secreta():
    ruta = os.path.join(db.DATA, "secret.key")
    os.makedirs(db.DATA, exist_ok=True)
    if not os.path.exists(ruta):
        with open(ruta, "w") as f:
            f.write(secrets.token_hex(32))
    with open(ruta) as f:
        return f.read().strip()


app.secret_key = _clave_secreta()
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")
app.teardown_appcontext(db.cerrar)


@app.errorhandler(413)
def demasiado_grande(_):
    flash("Los archivos superan el máximo de 40 MB por envío. Subí menos imágenes a la vez.", "error")
    destino = request.referrer if request.referrer and request.referrer.startswith(request.host_url) else url_for("tablero")
    return redirect(destino)


@app.errorhandler(d1.ErrorConexion)
def sin_base(error):
    return ("<!doctype html><meta charset='utf-8'><title>Sin conexión</title>"
            "<body style='font-family:sans-serif;max-width:560px;margin:15vh auto;padding:0 20px'>"
            "<h1>Sin conexión con la base de datos</h1>"
            "<p>No se pudo comunicar con Cloudflare D1. Revisá la conexión a internet y volvé a intentar en unos segundos.</p>"
            f"<p style='color:#666;font-size:13px'>{str(error)[:200]}</p></body>"), 503


SQL_USUARIO = """SELECT u.*, p.nombre AS perfil, p.permisos AS perfil_permisos FROM usuarios u
                      LEFT JOIN perfiles p ON p.id=u.perfil_id WHERE u.id=? AND u.activo=1"""
SQL_BORRADORES = "SELECT COUNT(*) AS n FROM protocolos WHERE borrador=1"


@app.before_request
def adelantar_lecturas():
    if session.get("uid") and request.endpoint != "static":
        db.precargar([(SQL_USUARIO, (session["uid"],)), (SQL_BORRADORES, ())])


def usuario_actual():
    uid = session.get("uid")
    if not uid:
        return None
    if g.get("yo_id") != uid:
        f = db.uno(SQL_USUARIO, (uid,))
        g.yo = {**dict(f), "permisos": permisos.efectivos(f, f["perfil_permisos"])} if f else None
        g.yo_id = uid
    return g.yo


def puede(permiso):
    u = usuario_actual()
    return bool(u and permiso in u["permisos"])


def requiere_login(f):
    @wraps(f)
    def envuelta(*a, **k):
        u = usuario_actual()
        if not u:
            return redirect(url_for("login", next=request.path))
        if u["debe_cambiar_clave"] and not session.get("real_uid") and request.endpoint not in ("cambiar_clave", "logout"):
            return redirect(url_for("cambiar_clave"))
        return f(*a, **k)
    return envuelta


def requiere_permiso(*claves):
    def decorador(f):
        @wraps(f)
        @requiere_login
        def envuelta(*a, **k):
            if not any(puede(c) for c in claves):
                flash("No tenés permiso para esa acción. Pedíselo a un administrador.", "error")
                destino = request.referrer if request.referrer and request.referrer.startswith(request.host_url) else None
                return redirect(destino or url_for("tablero"))
            return f(*a, **k)
        return envuelta
    return decorador


def requiere_admin(f):
    @wraps(f)
    @requiere_login
    def envuelta(*a, **k):
        if not usuario_actual()["admin"]:
            abort(403)
        return f(*a, **k)
    return envuelta


@app.before_request
def verificar_csrf():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    if request.method == "POST":
        token = request.form.get("_csrf") or request.headers.get("X-CSRF")
        if token != session["csrf"]:
            abort(400, "Formulario vencido: recargá la página e intentá de nuevo.")


@app.before_request
def modo_vista():
    if session.get("real_uid") and request.method == "POST" and request.endpoint not in ("ver_como", "dejar_de_ver_como", "logout"):
        mensaje = "Estás viendo el sistema como otro usuario: es solo para mirar, no se registran cambios. Volvé a tu usuario para trabajar."
        if request.path.startswith("/api/"):
            return jsonify({"error": mensaje}), 403
        flash(mensaje, "error")
        destino = request.referrer if request.referrer and request.referrer.startswith(request.host_url) else None
        return redirect(destino or url_for("tablero"))


@app.context_processor
def globales():
    u = usuario_actual()
    pendientes = db.uno(SQL_BORRADORES)["n"] if u else 0
    vista = None
    if u and session.get("real_uid"):
        real = db.uno("SELECT id, iniciales, nombre FROM usuarios WHERE id=? AND activo=1 AND admin=1", (session["real_uid"],))
        if real:
            vista = {"admin": real, "otros": db.q("SELECT id, iniciales, nombre FROM usuarios WHERE activo=1 AND id<>? ORDER BY iniciales", (real["id"],))}
    return {"vista_como": vista, "yo": u, "csrf": session.get("csrf", ""), "TIPOS": TIPOS, "SECTORES": SECTORES, "REGISTRO": REGISTRO,
            "n_borradores": pendientes, "informes_ok": informes is not None, "SEMAFOROS": SEMAFOROS, "nombre_base": db.nombre_base(),
            "NOMBRE_ETAPA": NOMBRE_ETAPA, "puede": puede, "mis_sectores": set((u["sectores"] or "").split(",")) if u else set()}


@app.template_filter("fh")
def fecha_hora(v):
    if not v:
        return ""
    try:
        return datetime.strptime(v, "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return v


MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]


@app.template_filter("larga")
def fecha_larga(d):
    return f"{d.day} {MESES[d.month - 1]} {d.year}" if d else ""


@app.template_filter("edad")
def edad(v):
    try:
        n = datetime.strptime(v, "%Y-%m-%d")
    except (TypeError, ValueError):
        return ""
    hoy = datetime.now()
    return f"{hoy.year - n.year - ((hoy.month, hoy.day) < (n.month, n.day))} ({n.year})"


@app.template_filter("dh")
def fecha_hora_dt(d):
    return d.strftime("%d/%m/%Y %H:%M hs") if d else ""


@app.template_filter("f")
def fecha(v):
    if not v:
        return ""
    try:
        return datetime.strptime(v, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return v


def es_local():
    return request.remote_addr in ("127.0.0.1", "::1")


@app.route("/configurar", methods=["GET", "POST"])
def configurar():
    if db.uno("SELECT 1 FROM usuarios WHERE admin=1 AND clave_hash IS NOT NULL"):
        return redirect(url_for("login"))
    if not es_local():
        abort(403, "La configuración inicial solo se hace desde la PC servidor.")
    if request.method == "POST":
        ini = request.form["iniciales"].strip().upper()
        clave = request.form["clave"]
        if len(clave) < 8 or clave != request.form["clave2"]:
            flash("La clave debe tener al menos 8 caracteres y coincidir en los dos campos.", "error")
        elif not ini:
            flash("Faltan las iniciales.", "error")
        else:
            existente = db.uno("SELECT id FROM usuarios WHERE iniciales=?", (ini,))
            h = generate_password_hash(clave)
            if existente:
                db.ex("UPDATE usuarios SET admin=1, activo=1, clave_hash=?, debe_cambiar_clave=0, nombre=? WHERE id=?",
                      (h, request.form["nombre"].strip() or ini, existente["id"]))
            else:
                db.ex("INSERT INTO usuarios (iniciales, nombre, sectores, admin, clave_hash, debe_cambiar_clave) "
                      "VALUES (?,?,?,1,?,0)", (ini, request.form["nombre"].strip() or ini, ",".join(SECTORES), h))
            flash("Administrador creado. Ya podés ingresar.", "ok")
            return redirect(url_for("login"))
    return render_template("configurar.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if not db.uno("SELECT 1 FROM usuarios WHERE admin=1 AND clave_hash IS NOT NULL"):
        return redirect(url_for("configurar"))
    if request.method == "POST":
        u = db.uno("SELECT * FROM usuarios WHERE iniciales=? AND activo=1", (request.form["iniciales"].strip().upper(),))
        if u and u["clave_hash"] and check_password_hash(u["clave_hash"], request.form["clave"]):
            session.clear()
            session["uid"] = u["id"]
            db.auditar(u["id"], None, "login")
            destino = request.args.get("next", "")
            return redirect(destino if destino.startswith("/") and not destino.startswith("//") else url_for("tablero"))
        flash("Iniciales o clave incorrectas (o el usuario todavía no tiene clave: pedísela al administrador).", "error")
    return render_template("login.html")


@app.route("/usuarios/<int:uid>/ver-como", methods=["POST"])
@requiere_login
def ver_como(uid):
    real_id = session.get("real_uid") or session.get("uid")
    real = db.uno("SELECT id, iniciales FROM usuarios WHERE id=? AND activo=1 AND admin=1", (real_id,))
    if not real:
        abort(403)
    objetivo = db.uno("SELECT id, iniciales, nombre FROM usuarios WHERE id=? AND activo=1", (uid,))
    if not objetivo:
        flash("El usuario no existe o está inactivo.", "error")
        return redirect(url_for("usuarios"))
    if objetivo["id"] == real["id"]:
        session.pop("real_uid", None)
        session["uid"] = real["id"]
        return redirect(url_for("usuarios"))
    session["real_uid"] = real["id"]
    session["uid"] = objetivo["id"]
    db.auditar(real["id"], None, "ver_como", f"{objetivo['iniciales']}")
    return redirect(url_for("tablero"))


@app.route("/dejar-de-ver-como", methods=["POST"])
@requiere_login
def dejar_de_ver_como():
    real_id = session.pop("real_uid", None)
    if real_id:
        session["uid"] = real_id
        flash("Volviste a tu usuario.", "ok")
        return redirect(url_for("usuarios"))
    return redirect(url_for("tablero"))


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/clave", methods=["GET", "POST"])
@requiere_login
def cambiar_clave():
    yo = usuario_actual()
    if request.method == "POST":
        nueva = request.form["nueva"]
        if not check_password_hash(yo["clave_hash"], request.form["actual"]):
            flash("La clave actual no es correcta.", "error")
        elif len(nueva) < 8 or nueva != request.form["nueva2"]:
            flash("La nueva clave debe tener al menos 8 caracteres y coincidir en los dos campos.", "error")
        else:
            db.ex("UPDATE usuarios SET clave_hash=?, debe_cambiar_clave=0 WHERE id=?", (generate_password_hash(nueva), yo["id"]))
            db.auditar(yo["id"], None, "cambio_clave")
            flash("Clave actualizada.", "ok")
            return redirect(url_for("tablero"))
    return render_template("clave.html")


def listas():
    out = {}
    for r in db.q(SQL_LISTAS):
        out.setdefault(r["nombre"], []).append(r["valor"])
    return out


def responsables(tipo):
    sector = REGISTRO[tipo].sector_responsable
    return [u for u in db.q("SELECT id, iniciales, nombre, sectores FROM usuarios WHERE activo=1 ORDER BY iniciales")
            if sector in (u["sectores"] or "").split(",")]


SQL_ETAPAS = "SELECT estudio_id, etapa, fecha_hora FROM etapas"
SQL_CON_IHQ = "SELECT e.protocolo_id, e.tipo FROM micro m JOIN estudios e ON e.id=m.estudio_id WHERE m.solicita_ihq=1 AND e.anulado=0"
SQL_CON_DECAL = "SELECT e.protocolo_id, e.tipo FROM macro m JOIN estudios e ON e.id=m.estudio_id WHERE m.decal=1 AND e.anulado=0"
SQL_TIPOS = "SELECT e.protocolo_id, e.tipo FROM estudios e WHERE e.anulado=0"


def consultas_estudios(where="1=1", params=(), borradores=False):
    principal = f"""SELECT e.*, p.numero, p.apellido, p.nombre, p.dni, p.fecha_recoleccion, p.cargado_sistema, p.borrador,
                            p.lugar_recoleccion, p.lugar_entrega, u.iniciales AS responsable, l.codigo AS lote, t.tacos AS tacos_muestra,
                            (SELECT COUNT(*) FROM estudios x WHERE x.protocolo_id=e.protocolo_id AND x.anulado=0) AS hermanos,
                            (SELECT COUNT(*) FROM comentarios c WHERE c.protocolo_id=e.protocolo_id) AS n_comentarios
                     FROM estudios e JOIN protocolos p ON p.id=e.protocolo_id
                     LEFT JOIN usuarios u ON u.id=e.responsable_id LEFT JOIN lotes l ON l.id=e.lote_id
                     LEFT JOIN tacos_organo t ON t.organo=e.sitio AND t.tipo_lote=l.tipo_lote
                     WHERE {where}{'' if borradores else ' AND p.borrador=0'} ORDER BY p.id DESC, e.id"""
    return [(principal, tuple(params)), (SQL_ETAPAS, ()), (SQL_CON_IHQ, ()), (SQL_CON_DECAL, ()), (SQL_TIPOS, ())]


def tacos_del_estudio(e):
    if e["tipo"] != "BP" or e["tacos_muestra"] is None:
        return None
    frascos = int(e["cantidad"]) if (e["cantidad"] or "").isdigit() else 1
    return frascos * e["tacos_muestra"]


def sql_tacos(campo):
    return f"""SELECT e.{campo} AS clave,
                      SUM((CASE WHEN e.cantidad GLOB '[0-9]*' THEN CAST(e.cantidad AS INTEGER) ELSE 1 END) * COALESCE(t.tacos, 0)) AS tacos,
                      SUM(CASE WHEN t.tacos IS NULL THEN 1 ELSE 0 END) AS sin_dato
               FROM estudios e LEFT JOIN lotes l ON l.id=e.lote_id LEFT JOIN tacos_organo t ON t.organo=e.sitio AND t.tipo_lote=l.tipo_lote
               WHERE e.tipo='BP' AND e.anulado=0 AND e.{campo} IS NOT NULL{{filtro}} GROUP BY e.{campo}"""


def tacos_por(campo, valor=None):
    filas = db.q(sql_tacos(campo).format(filtro="" if valor is None else f" AND e.{campo}=?"), () if valor is None else (valor,))
    return {f["clave"]: (f["tacos"] or 0, f["sin_dato"] or 0) for f in filas}


def tacos_de_lote(lote_id):
    return tacos_por("lote_id", lote_id).get(lote_id, (0, 0))


def tacos_de_protocolo(protocolo_id):
    return tacos_por("protocolo_id", protocolo_id).get(protocolo_id, (0, 0))


def ubicar_en_lote(protocolo_id, usuario_id):
    capacidad = REGISTRO["BP"].tacos_por_lote
    with db.bloqueo("lotes"):
        actual = lote_del_protocolo(protocolo_id)
        if not actual or actual["cerrado"] or actual["tipo_lote"] not in REGISTRO["BP"].lotes:
            return None
        propios = tacos_de_protocolo(protocolo_id)[0]
        ocupados = tacos_de_lote(actual["id"])[0] - propios
        if not propios or not ocupados or ocupados + propios <= capacidad:
            return None
        destino = next((l for l in db.q("SELECT * FROM lotes WHERE tipo_lote=? AND fecha=? AND cerrado=0 AND id<>? ORDER BY numero",
                                        (actual["tipo_lote"], hoy(), actual["id"])) if tacos_de_lote(l["id"])[0] + propios <= capacidad), None)
        destino = destino or crear_lote(actual["tipo_lote"], usuario_id)
        asignar_lote(protocolo_id, destino, usuario_id)
        db.auditar(usuario_id, protocolo_id, "lote_lleno", f"{actual['codigo']} llegó a {ocupados}/{capacidad} tacos: pasó a {destino['codigo']}")
        return {"desde": actual["codigo"], "hacia": destino["codigo"], "ocupados": ocupados, "capacidad": capacidad}


def avisar_lote_lleno(resultado):
    if resultado:
        flash(f"El lote {resultado['desde']} se llenó ({resultado['ocupados']}/{resultado['capacidad']} tacos): el protocolo pasó al lote {resultado['hacia']}.", "ok")


def cargar_estudios(where="1=1", params=(), borradores=False):
    consultas = consultas_estudios(where, params, borradores)
    db.precargar(consultas)
    filas = db.q(*consultas[0])
    hechas, fechas = {}, {}
    for r in db.q(SQL_ETAPAS):
        hechas.setdefault(r["estudio_id"], set()).add(r["etapa"])
        fechas.setdefault(r["estudio_id"], {})[r["etapa"]] = r["fecha_hora"]
    con_ihq = {(r["protocolo_id"], r["tipo"]) for r in db.q(SQL_CON_IHQ)}
    con_decal_ = {(r["protocolo_id"], r["tipo"]) for r in db.q(SQL_CON_DECAL)}
    out, fer, ahora_ = [], db.feriados(), datetime.now()
    flujos = flujos_de() if filas else {}
    for f in filas:
        d = dict(f)
        d["flujos"] = flujos.get(f["protocolo_id"], "")
        tipo = REGISTRO[f["tipo"]]
        d["solicita_ihq"] = (f["protocolo_id"], f["tipo"]) in con_ihq
        decal = (f["protocolo_id"], f["tipo"]) in con_decal_
        d["decal"] = decal
        d["estado"], prox = tipo.estado(hechas.get(f["id"], set()), Variante(d["solicita_ihq"], decal), bool(f["anulado"]))
        if f["borrador"]:
            d["estado"] = "A completar"
        d["proxima"] = prox
        d["vence"] = trazabilidad.limite(f["tipo"], prox[0], f["fecha_recoleccion"] or f["creado_en"], fer, decal) if prox else None
        d["semaforo"] = (trazabilidad.semaforo_acumulado(f["tipo"], f["fecha_recoleccion"] or f["creado_en"], {"ingreso": f["creado_en"], **fechas.get(f["id"], {})},
                                                         d["vence"], ahora_, fer, decal) if d["vence"] and not f["borrador"] else "")
        d["etiqueta"] = etiqueta(f)
        d["tacos"] = tacos_del_estudio(f)
        out.append(d)
    return out


def etiqueta(e):
    if e["tipo"] == "PAP" or not e["sitio"]:
        return TIPOS[e["tipo"]]
    return f"{TIPOS[e['tipo']]} · {e['sitio']}"


def etiqueta_flujos(tipos, con_ihq):
    t = [k for k in ("BP", "CT", "PAP") if k in tipos]
    if not t:
        return ""
    return t[0] + (" c/ " + " y ".join(t[1:]) if len(t) > 1 else "") + (" + IHQ" if con_ihq else "")


def flujos_de(protocolo_ids=None):
    tipos = {}
    for r in db.q(SQL_TIPOS):
        tipos.setdefault(r["protocolo_id"], set()).add(r["tipo"])
    ihq = {r["protocolo_id"] for r in db.q(SQL_CON_IHQ)}
    todos = {pid: etiqueta_flujos(ts, pid in ihq) for pid, ts in tipos.items()}
    return todos if not protocolo_ids else {pid: todos[pid] for pid in protocolo_ids if pid in todos}


def miembros_flujo(e):
    return db.q("SELECT * FROM estudios WHERE protocolo_id=? AND tipo=? AND anulado=0 ORDER BY id", (e["protocolo_id"], e["tipo"]))


def variante_de(protocolo_id, tipo):
    r = db.uno(SQL_VARIANTE, (protocolo_id, tipo, protocolo_id, tipo))
    return Variante(r["ihq"], r["decal"])


def flujo_avanzado(protocolo_id, tipo):
    r = db.uno("""SELECT x.etapa FROM etapas x JOIN estudios e ON e.id=x.estudio_id
                  WHERE e.protocolo_id=? AND e.tipo=? AND e.anulado=0 ORDER BY x.fecha_hora DESC, x.id DESC LIMIT 1""", (protocolo_id, tipo))
    return NOMBRE_ETAPA.get(r["etapa"], r["etapa"]) if r else None


def protocolo_cerrado(pid):
    activos = [x for x in cargar_estudios("e.protocolo_id=?", (pid,), borradores=True) if not x["anulado"]]
    return bool(activos) and not any(x["borrador"] or x["proxima"] for x in activos)


def tipos_lote_fijos():
    fijos = listas().get("tipo_lote", [])
    return fijos + [t for e in REGISTRO.values() for t in e.lotes if t not in fijos]


@app.route("/")
@requiere_login
def tablero():
    yo = usuario_actual()
    f = {k: request.args.get(k, "") for k in ("q", "tipo", "ver", "sector", "etapa", "semaforo", "ubicacion")}
    f["ver"] = f["ver"] or "pendientes"
    grupos = {}
    for e in cargar_estudios():
        grupos.setdefault(e["protocolo_id"], []).append(e)
    mis = set((yo["sectores"] or "").split(","))
    conteo = {"pendientes": 0, "informados": 0, "sin_cargar": 0}
    for g in grupos.values():
        activos = [e for e in g if not e["anulado"]]
        if any(e["proxima"] for e in activos):
            conteo["pendientes"] += 1
        elif activos:
            conteo["informados"] += 1
            if not activos[0]["cargado_sistema"]:
                conteo["sin_cargar"] += 1

    def del_ver(g):
        activos = [e for e in g if not e["anulado"]]
        pendientes = [e for e in activos if e["proxima"]]
        v = f["ver"]
        if v == "pendientes":
            return bool(pendientes)
        if v == "mios":
            return any(e["proxima"][2] in mis for e in pendientes)
        if v == "informados":
            return bool(activos) and not pendientes
        if v == "sin_cargar":
            return bool(activos) and not pendientes and not activos[0]["cargado_sistema"]
        if v == "anulados":
            return any(e["anulado"] for e in g)
        return True

    def estudio_ok(e):
        if f["tipo"] and e["tipo"] != f["tipo"]:
            return False
        if f["sector"] and not (e["proxima"] and e["proxima"][2] == f["sector"]):
            return False
        if f["etapa"] and not (e["proxima"] and e["proxima"][0] == f["etapa"]):
            return False
        return True

    def con_texto(g):
        q = f["q"].lower()
        return any(q in " ".join(str(e[k] or "") for k in ("numero", "apellido", "nombre", "dni", "lote", "sitio")).lower() for e in g)

    filas = []
    for g in grupos.values():
        if not del_ver(g) or (f["q"] and not con_texto(g)):
            continue
        if f["ubicacion"] and f["ubicacion"] not in (g[0]["lugar_recoleccion"], g[0]["lugar_entrega"]):
            continue
        elegidos = [e for e in g if estudio_ok(e)]
        if f["ver"] == "anulados":
            vista = [e for e in elegidos if e["anulado"]]
        elif f["ver"] in ("pendientes", "mios"):
            vista = [e for e in elegidos if not e["anulado"] and e["proxima"] and (f["ver"] != "mios" or e["proxima"][2] in mis)]
        else:
            vista = [e for e in elegidos if not e["anulado"]] or elegidos
        if not vista:
            continue
        primero = g[0]
        estados = list(dict.fromkeys((e["estado"], SECTORES[e["proxima"][2]] if e["proxima"] else "", bool(e["anulado"]), not e["proxima"]) for e in vista))
        vencen = [e["vence"] for e in vista if e["vence"]]
        filas.append({
            "protocolo_id": primero["protocolo_id"], "numero": primero["numero"], "flujos": primero["flujos"],
            "n_comentarios": primero.get("n_comentarios"), "estudios": sum(1 for e in g if not e["anulado"]),
            "apellido": primero["apellido"], "nombre": primero["nombre"], "dni": primero["dni"],
            "muestras": list(dict.fromkeys((e["subcategoria"] or "", e["tipo_muestra"] or "") for e in vista)),
            "lote_id": primero["lote_id"], "lote": primero["lote"],
            "responsables": list(dict.fromkeys(e["responsable"] for e in vista if e["responsable"])),
            "estados": estados,
            "semaforo": max((e["semaforo"] for e in vista), key=lambda s: trazabilidad.GRAVEDAD[s or ""]),
            "vence": min(vencen) if vencen else None,
            "creado_en": min(e["creado_en"] for e in g)})
    por_semaforo = {k: sum(1 for r in filas if r["semaforo"] == k) for k in SEMAFOROS}
    if f["semaforo"]:
        filas = [r for r in filas if r["semaforo"] == f["semaforo"]]
    lugares = sorted({l for g in grupos.values() for l in (g[0]["lugar_recoleccion"], g[0]["lugar_entrega"]) if l} | ({f["ubicacion"]} if f["ubicacion"] else set()), key=db.clave_nombre)
    return render_template("tablero.html", filas=filas, f=f, conteo=conteo, por_semaforo=por_semaforo, lugares=lugares,
                           etapas=[(k, n) for k, n in NOMBRE_ETAPA.items() if k != "ingreso"])


def tipos_lote(tipo=None):
    fijos = tipos_lote_fijos()
    cito = [u["iniciales"] for u in responsables("PAP")]
    if tipo:
        return REGISTRO[tipo].tipos_lote(fijos, cito)
    return fijos + [i for i in cito if i not in fijos]


def hoy():
    return datetime.now().strftime("%Y-%m-%d")


def crear_lote(tipo_lote, usuario_id):
    f = hoy()
    for _ in range(5):
        n = (db.uno("SELECT MAX(numero) AS m FROM lotes WHERE tipo_lote=? AND fecha=?", (tipo_lote, f))["m"] or 0) + 1
        codigo = f"{tipo_lote}-{f[5:7]}{f[8:10]}.{n}"
        try:
            lid = db.ex("INSERT INTO lotes (codigo, tipo_lote, fecha, numero, creado_por, creado_en) VALUES (?,?,?,?,?,?)",
                        (codigo, tipo_lote, f, n, usuario_id, db.ahora()))
            break
        except db.IntegrityError:
            continue
    else:
        raise RuntimeError("No se pudo crear el lote, probá de nuevo.")
    db.auditar(usuario_id, None, "nuevo_lote", codigo)
    return db.uno("SELECT * FROM lotes WHERE id=?", (lid,))


def validar_lote(seleccion, tipo, actual_id=None):
    if not seleccion:
        return None, None
    if seleccion.startswith("nuevo:"):
        if seleccion[6:] not in tipos_lote(tipo):
            return None, f"Tipo de lote inválido para {TIPOS[tipo]}."
        return seleccion, None
    lote = db.uno("SELECT * FROM lotes WHERE id=?", (int(seleccion),)) if seleccion.isdigit() else None
    if not lote:
        return None, "El lote elegido no existe."
    if lote["id"] != actual_id:
        if lote["cerrado"]:
            return None, f"El lote {lote['codigo']} está cerrado."
        if lote["tipo_lote"] not in tipos_lote(tipo):
            return None, f"El lote {lote['codigo']} no corresponde a {TIPOS[tipo]}."
    return lote, None


def concretar_lote(valor, usuario_id):
    if isinstance(valor, str):
        return crear_lote(valor[6:], usuario_id)
    return valor


def lotes_para_formulario(tipo, actual_id=None):
    validos = tipos_lote(tipo)
    lotes = [l for l in db.q("SELECT * FROM lotes WHERE fecha=? AND cerrado=0 ORDER BY tipo_lote, numero", (hoy(),))
             if l["tipo_lote"] in validos]
    if actual_id and actual_id not in [l["id"] for l in lotes]:
        actual = db.uno("SELECT * FROM lotes WHERE id=?", (actual_id,))
        if actual:
            lotes.insert(0, actual)
    return lotes


def primer_estudio(protocolo_id):
    return db.uno("SELECT * FROM estudios WHERE protocolo_id=? AND anulado=0 ORDER BY id LIMIT 1", (protocolo_id,))


def lote_del_protocolo(protocolo_id):
    primero = primer_estudio(protocolo_id)
    return db.uno("SELECT * FROM lotes WHERE id=?", (primero["lote_id"],)) if primero and primero["lote_id"] else None


def asignar_lote(protocolo_id, lote, usuario_id):
    anterior = lote_del_protocolo(protocolo_id)
    db.ex("UPDATE estudios SET lote_id=? WHERE protocolo_id=? AND anulado=0", (lote["id"] if lote else None, protocolo_id))
    antes, despues = (anterior["codigo"] if anterior else None), (lote["codigo"] if lote else None)
    if antes != despues:
        db.auditar(usuario_id, protocolo_id, "lote", f"{antes or '—'} → {despues or '—'}")


CAMPOS_PROTOCOLO = ["fecha_recoleccion", "dni", "cobertura", "n_afiliado", "nombre", "apellido", "sexo", "exento",
                    "fecha_nacimiento", "email", "telefono", "medico", "lugar_recoleccion", "lugar_entrega", "observaciones"]
CAMPOS_ESTUDIO = ["categoria", "subcategoria", "sitio", "tipo_muestra", "cantidad", "citologia_hormonal", "observaciones",
                  "responsable_id", "tecnicas_especiales"]


CAMPOS_PACIENTE = ["apellido", "nombre", "fecha_nacimiento", "sexo", "cobertura", "n_afiliado", "exento", "email", "telefono"]


def sincronizar_paciente(pid, dni, campos=None):
    if not dni:
        return
    p = protocolo_o_404(pid)
    campos = [c for c in CAMPOS_PACIENTE if p[c]] if campos is None else [c for c in campos if c in CAMPOS_PACIENTE]
    if not campos:
        return
    db.ex(f"UPDATE protocolos SET {', '.join(c + '=?' for c in campos)} WHERE dni=? AND id<>?", [p[c] for c in campos] + [dni, pid])


def leer_protocolo():
    d = {k: request.form.get(k, "").strip() for k in CAMPOS_PROTOCOLO}
    d["numero"] = request.form.get("numero", "").strip()
    return d


def limpiar_valor(texto):
    return " ".join((texto or "").replace(",", " ").replace(";", " ").split())[:80]


def sin_repetidos(valores):
    vistos = {}
    for v in valores:
        if v:
            vistos.setdefault(v.lower(), v)
    return list(vistos.values())


def recordar_tecnicas(e):
    pedidas = [x for x in (e.get("tecnicas_especiales") or "").split("; ") if x]
    if not pedidas:
        return
    actuales = listas().get("tecnicas_especiales", [])
    conocidas = {x.lower() for x in actuales}
    nuevas = list(dict.fromkeys(x for x in pedidas if x.lower() not in conocidas))
    if nuevas:
        db.lote_ex([("INSERT INTO listas (nombre, orden, valor) VALUES (?,?,?)", ("tecnicas_especiales", len(actuales) + i, v)) for i, v in enumerate(nuevas)])


def leer_estudio(prefijo, tipo):
    t = REGISTRO[tipo]
    campo = lambda k: request.form.get(prefijo + k, "").strip()
    return {"tipo": tipo, "categoria": t.categoria,
            "subcategoria": t.subcategoria_fija or campo("subcategoria"),
            "sitio": t.sitio_fijo or campo("sitio"),
            "tipo_muestra": campo("tipo_muestra"),
            "cantidad": campo("cantidad") if campo("cantidad") in t.cantidades else None,
            "citologia_hormonal": campo("citologia_hormonal") if t.citologia_hormonal else None,
            "observaciones": campo("observaciones"),
            "responsable_id": int(campo("responsable_id")) if campo("responsable_id").isdigit() else None,
            "tecnicas_especiales": ("; ".join(sin_repetidos(limpiar_valor(x) for x in request.form.getlist(prefijo + "tecnicas_especiales"))) or None)
            if t.pide_tecnicas else None,
            "lote_sel": campo("lote_id"),
            "id": int(campo("id")) if campo("id").isdigit() else None}


def estudios_del_formulario():
    indices = sorted({int(k.split("-")[1]) for k in request.form
                      if k.startswith("e-") and k.endswith("-tipo") and k.split("-")[1].isdigit()})
    return [leer_estudio(f"e-{i}-", request.form[f"e-{i}-tipo"]) for i in indices
            if request.form.get(f"e-{i}-tipo") in REGISTRO]


def validar_estudios(lista, existentes=(), actual_id=None, con_lote=True):
    error = validar_combinacion(list(existentes) + lista)
    if error:
        return None, error
    if not (con_lote and lista):
        return None, None
    return validar_lote(lista[0]["lote_sel"], lista[0]["tipo"], actual_id)


def crear_estudio(protocolo_id, e, lote, usuario_id):
    eid = db.ex(f"INSERT INTO estudios (protocolo_id, tipo, creado_por, creado_en, {', '.join(CAMPOS_ESTUDIO)}) "
                f"VALUES (?,?,?,?,{','.join('?' * len(CAMPOS_ESTUDIO))})",
                [protocolo_id, e["tipo"], usuario_id, db.ahora()] + [e[k] for k in CAMPOS_ESTUDIO])
    recordar_tecnicas(e)
    db.auditar(usuario_id, protocolo_id, "nuevo_estudio", etiqueta(e), eid)
    asignar_lote(protocolo_id, lote, usuario_id)
    return eid


def contexto_formulario():
    return {"listas": listas(),
            "resp": {t: responsables(t) for t in REGISTRO},
            "lotes_abiertos": {t: lotes_para_formulario(t) for t in REGISTRO},
            "tipos_lote_de": {t: tipos_lote(t) for t in REGISTRO}}


SQL_PROTOCOLO = "SELECT p.*, u.iniciales AS creador FROM protocolos p LEFT JOIN usuarios u ON u.id=p.creado_por WHERE p.id=?"
SQL_HISTORIAL_PROTOCOLO = """SELECT a.*, u.iniciales FROM auditoria a LEFT JOIN usuarios u ON u.id=a.usuario_id
                        WHERE protocolo_id=? ORDER BY a.id DESC"""


def protocolo_o_404(pid):
    p = db.uno(SQL_PROTOCOLO, (pid,))
    if not p:
        abort(404)
    return p


@app.route("/protocolo/nuevo", methods=["GET", "POST"])
@requiere_permiso("protocolo_crear")
def protocolo_nuevo():
    yo = usuario_actual()
    p, lista = {}, []
    if request.method == "POST":
        p, lista = leer_protocolo(), estudios_del_formulario()
        lote_valor, error = validar_estudios(lista)
        if not p["numero"]:
            error = "Falta el N° de protocolo."
        elif (dup := db.uno("SELECT id, borrador FROM protocolos WHERE numero=?", (p["numero"],))) and dup["borrador"]:
            flash(f"El protocolo {p['numero']} ya fue generado desde Recepción y falta completarlo: cargá los datos acá.", "ok")
            return redirect(url_for("protocolo_editar", pid=dup["id"]))
        elif dup:
            error = f"El protocolo {p['numero']} ya existe: para sumarle un estudio usá «+ Agregar estudio» en su ficha."
        elif not lista:
            error = "Agregá al menos un estudio (PAP, Biopsia o Citología)."
        error = error or error_ubicaciones(p)
        if error:
            flash(error, "error")
        else:
            cols = ["numero"] + CAMPOS_PROTOCOLO
            pid = db.ex(f"INSERT INTO protocolos (creado_por, creado_en, {', '.join(cols)}) VALUES (?,?,{','.join('?' * len(cols))})",
                        [yo["id"], db.ahora()] + [p[c] for c in cols])
            db.auditar(yo["id"], pid, "ingreso", p["numero"])
            sincronizar_paciente(pid, p["dni"])
            lote = concretar_lote(lote_valor, yo["id"])
            for e in lista:
                crear_estudio(pid, e, lote, yo["id"])
            avisar_lote_lleno(ubicar_en_lote(pid, yo["id"]))
            flash(f"Protocolo {p['numero']} ingresado con {len(lista)} estudio(s).", "ok")
            if not numero_emitido(p["numero"]):
                flash(f"Atención: {p['numero']} no figura en ningún lote de etiquetas confirmado (menú Etiquetas).", "error")
            return redirect(url_for("protocolo", pid=pid))
    elif request.args.get("tipo") in REGISTRO:
        lista = [{"tipo": request.args["tipo"]}]
    ultimos = [r["numero"] for r in db.q("SELECT numero FROM protocolos ORDER BY id DESC LIMIT 5")]
    return render_template("protocolo_form.html", p=p, estudios=lista, nuevo=True, ultimos=ultimos, **contexto_formulario())


def completo(p, estudios):
    return bool(p["apellido"] and p["nombre"] and estudios and all(e["cantidad"] for e in estudios if not e["anulado"]))


@app.route("/protocolo/<int:pid>/editar", methods=["GET", "POST"])
@requiere_permiso("protocolo_crear", "protocolo_editar")
def protocolo_editar(pid):
    yo = usuario_actual()
    p = protocolo_o_404(pid)
    borrador = bool(p["borrador"])
    if not borrador and not puede("protocolo_editar"):
        flash("No tenés permiso para editar protocolos. Pedíselo a un administrador.", "error")
        return redirect(url_for("protocolo", pid=pid))
    actuales = {e["id"]: e for e in db.q("SELECT * FROM estudios WHERE protocolo_id=? ORDER BY id", (pid,))}
    lista = [{**dict(e), "lote_sel": str(e["lote_id"] or "")} for e in actuales.values() if not e["anulado"]] if borrador else []
    if request.method == "POST":
        d = leer_protocolo()
        error = None
        if borrador:
            lista = [e for e in estudios_del_formulario() if e["id"] is None or e["id"] in actuales]
            error = validar_combinacion(lista) if lista else "El protocolo tiene que tener al menos un estudio."
            lote_valor = None
            if not error:
                actual = lote_del_protocolo(pid)
                lote_valor, error = validar_lote(lista[0]["lote_sel"], lista[0]["tipo"], actual["id"] if actual else None)
        if not d["numero"]:
            error = "Falta el N° de protocolo."
        elif db.uno("SELECT 1 FROM protocolos WHERE numero=? AND id<>?", (d["numero"], pid)):
            error = f"Ya existe otro protocolo {d['numero']}."
        error = error or error_ubicaciones(d, p)
        if error:
            flash(error, "error")
        else:
            cols = ["numero"] + CAMPOS_PROTOCOLO
            db.ex(f"UPDATE protocolos SET {', '.join(c + '=?' for c in cols)} WHERE id=?", [d[c] for c in cols] + [pid])
            db.auditar(yo["id"], pid, "editar_protocolo")
            if not p["dni"] or p["dni"] == d["dni"]:
                sincronizar_paciente(pid, d["dni"], [c for c in CAMPOS_PACIENTE if d[c] != (p[c] or "")] if p["dni"] else None)
            if not borrador:
                flash("Datos del protocolo actualizados.", "ok")
                return redirect(url_for("protocolo", pid=pid))
            lote = concretar_lote(lote_valor, yo["id"])
            for e in lista:
                if e["id"]:
                    db.ex(f"UPDATE estudios SET {', '.join(k + '=?' for k in CAMPOS_ESTUDIO)} WHERE id=?",
                          [e[k] for k in CAMPOS_ESTUDIO] + [e["id"]])
                    recordar_tecnicas(e)
                else:
                    crear_estudio(pid, e, lote, yo["id"])
            asignar_lote(pid, lote, yo["id"])
            quitados = [k for k, e in actuales.items() if k not in {x["id"] for x in lista} and not e["anulado"]]
            for k in quitados:
                db.ex("UPDATE estudios SET anulado=1, motivo_anulacion=? WHERE id=?", ("Quitado al completar el protocolo", k))
                db.auditar(yo["id"], pid, "anular", "Quitado al completar el protocolo", k)
            if completo(d, [dict(x) for x in db.q("SELECT cantidad, anulado FROM estudios WHERE protocolo_id=?", (pid,))]):
                db.ex("UPDATE protocolos SET borrador=0 WHERE id=?", (pid,))
                db.auditar(yo["id"], pid, "completar", d["numero"])
                avisar_lote_lleno(ubicar_en_lote(pid, yo["id"]))
                flash(f"Protocolo {d['numero']} completo.", "ok")
                return redirect(url_for("protocolo", pid=pid))
            flash("Datos guardados. Para completar el protocolo faltan apellido, nombre y la cantidad de cada estudio.", "ok")
            return redirect(url_for("protocolo_editar", pid=pid))
        p = {**dict(p), **d}
    ctx = contexto_formulario()
    for e in actuales.values():
        if e["lote_id"] and all(l["id"] != e["lote_id"] for l in ctx["lotes_abiertos"][e["tipo"]]):
            ctx["lotes_abiertos"][e["tipo"]].insert(0, db.uno("SELECT * FROM lotes WHERE id=?", (e["lote_id"],)))
    return render_template("protocolo_form.html", p=p, estudios=lista, nuevo=False, completar=borrador, ultimos=[], **ctx)


CAMPOS_RAPIDOS = ["apellido", "nombre", "fecha_nacimiento", "cobertura", "medico", "lugar_recoleccion", "lugar_entrega", "fecha_recoleccion"]


def protocolos_del_paciente(base):
    if base["dni"]:
        donde, params = "p.dni=?", (base["dni"],)
    else:
        donde = "LOWER(p.apellido)=LOWER(?) AND LOWER(COALESCE(p.nombre,''))=LOWER(?) AND COALESCE(p.fecha_nacimiento,'')=?"
        params = (base["apellido"] or "", base["nombre"] or "", base["fecha_nacimiento"] or "")
    protocolos = db.q(f"SELECT p.*, u.iniciales AS creador FROM protocolos p LEFT JOIN usuarios u ON u.id=p.creado_por WHERE {donde} ORDER BY p.id DESC", params)
    return protocolos if any(x["id"] == base["id"] for x in protocolos) else [base]


@app.route("/paciente/<int:pid>/datos", methods=["POST"])
@requiere_permiso("protocolo_editar")
def paciente_datos(pid):
    base = protocolo_o_404(pid)
    d = {k: request.form.get(k, "").strip() for k in ["dni"] + CAMPOS_PACIENTE}
    if not (d["apellido"] or d["nombre"]):
        flash("Escribí al menos el apellido o el nombre del paciente.", "error")
        return redirect(url_for("paciente", pid=pid))
    protocolos = protocolos_del_paciente(base)
    ajenos = [x for x in db.q("SELECT id FROM protocolos WHERE dni=?", (d["dni"],)) if x["id"] not in {y["id"] for y in protocolos}] if d["dni"] else []
    if ajenos:
        flash("Ya hay otro paciente con ese DNI. Si es la misma persona, corregí el DNI en el protocolo que tiene el error.", "error")
        return redirect(url_for("paciente", pid=pid))
    cambios = [k for k in d if d[k] != (base[k] or "")]
    if cambios:
        cols = list(d)
        yo = usuario_actual()
        db.lote_ex([(f"UPDATE protocolos SET {', '.join(c + '=?' for c in cols)} WHERE id=?", [d[c] for c in cols] + [x["id"]]) for x in protocolos]
                   + [("INSERT INTO auditoria (fecha_hora, usuario_id, protocolo_id, accion, detalle) VALUES (?,?,?,?,?)",
                       (db.ahora(), yo["id"], x["id"], "editar_paciente", ", ".join(cambios))) for x in protocolos])
        flash(f"Datos del paciente actualizados en {len(protocolos)} protocolo(s).", "ok")
    return redirect(url_for("paciente", pid=pid))


@app.route("/paciente/<int:pid>")
@requiere_login
def paciente(pid):
    base = protocolo_o_404(pid)
    protocolos = protocolos_del_paciente(base)
    ids = [x["id"] for x in protocolos]
    marcas = ",".join("?" * len(ids))
    estudios = cargar_estudios(f"e.protocolo_id IN ({marcas})", ids, borradores=True)
    por_protocolo = {}
    for e in estudios:
        por_protocolo.setdefault(e["protocolo_id"], []).append(e)
    emitidos = {r["protocolo_id"]: r for r in db.q(
        f"SELECT id, protocolo_id, estudios FROM informes_emitidos WHERE protocolo_id IN ({marcas})", ids)}
    informes = {}
    for x in protocolos:
        activos = [e for e in por_protocolo.get(x["id"], []) if not e["anulado"]]
        r = emitidos.get(x["id"])
        cerrado = bool(activos) and not any(e["borrador"] or e["proxima"] for e in activos)
        if r and cerrado and (not r["estudios"] or set(r["estudios"].split(",")) == {str(e["id"]) for e in activos}):
            informes[x["id"]] = r["id"]
    return render_template("paciente.html", p=base, listas=listas(), protocolos=protocolos, estudios=por_protocolo, informes=informes,
                           flujos=flujos_de(ids), origen=pid)


@app.route("/protocolo/<int:pid>/datos", methods=["POST"])
@requiere_permiso("protocolo_editar")
def protocolo_datos(pid):
    p = protocolo_o_404(pid)
    if p["borrador"]:
        return redirect(url_for("protocolo_editar", pid=pid))
    d = {k: request.form[k].strip() for k in CAMPOS_RAPIDOS if k in request.form}
    cambios = [k for k in d if d[k] != (p[k] or "")]
    error = error_ubicaciones(d, p)
    if error:
        flash(error, "error")
        return redirect(url_for("protocolo", pid=pid))
    if cambios:
        db.ex(f"UPDATE protocolos SET {', '.join(k + '=?' for k in cambios)} WHERE id=?", [d[k] for k in cambios] + [pid])
        db.auditar(usuario_actual()["id"], pid, "editar_protocolo", ", ".join(cambios))
        sincronizar_paciente(pid, p["dni"], cambios)
        flash("Datos del protocolo actualizados.", "ok")
    return redirect(url_for("protocolo", pid=pid))


SQL_COMENTARIOS = """SELECT c.*, u.iniciales, u.nombre AS autor FROM comentarios c LEFT JOIN usuarios u ON u.id=c.usuario_id
                     WHERE c.protocolo_id=? ORDER BY c.id DESC"""


SQL_IMAGENES = """SELECT i.id, i.protocolo_id, i.nombre, i.descripcion, i.creado_en, u.iniciales FROM imagenes i
                  LEFT JOIN usuarios u ON u.id=i.usuario_id WHERE i.protocolo_id=? ORDER BY i.id"""
MAX_IMAGENES = 30
PERMISOS_IMAGENES = ("protocolo_editar", "macro", "micro")


def procesar_imagen(contenido):
    from PIL import Image, ImageOps
    if len(contenido) > 15 * 1024 * 1024:
        raise ValueError("pesa más de 15 MB")
    try:
        img = Image.open(io.BytesIO(contenido))
        if img.format not in ("JPEG", "PNG", "WEBP"):
            raise ValueError
        img = ImageOps.exif_transpose(img)
        if img.mode in ("RGBA", "LA", "P"):
            img = img.convert("RGBA")
            fondo = Image.new("RGB", img.size, (255, 255, 255))
            fondo.paste(img, mask=img.split()[-1])
            img = fondo
        else:
            img = img.convert("RGB")
    except Exception:
        raise ValueError("no es una imagen JPG, PNG o WEBP válida")

    def jpeg(lado, calidad):
        chica = img.copy()
        chica.thumbnail((lado, lado))
        salida = io.BytesIO()
        chica.save(salida, "JPEG", quality=calidad, optimize=True)
        return salida.getvalue()

    for lado, calidad in ((2000, 85), (1800, 80), (1600, 72), (1400, 65), (1200, 55), (1000, 45), (800, 40)):
        grande = jpeg(lado, calidad)
        if len(grande) <= 1_200_000:
            break
    return grande, jpeg(360, 78)


def imagenes_de(pid):
    return db.q(SQL_IMAGENES, (pid,))


def volver_a_imagenes(pid):
    destino = request.referrer if request.referrer and request.referrer.startswith(request.host_url) else url_for("protocolo", pid=pid)
    return redirect(destino.split("#")[0] + "#imagenes")


@app.route("/protocolo/<int:pid>/imagenes", methods=["POST"])
@requiere_permiso(*PERMISOS_IMAGENES)
def imagenes_subir(pid):
    protocolo_o_404(pid)
    archivos = [f for f in request.files.getlist("imagenes") if f and f.filename]
    if not archivos:
        flash("Elegí al menos una imagen.", "error")
        return volver_a_imagenes(pid)
    cupo = MAX_IMAGENES - db.uno("SELECT COUNT(*) AS n FROM imagenes WHERE protocolo_id=?", (pid,))["n"]
    descripcion = request.form.get("descripcion", "").strip()[:200]
    yo = usuario_actual()
    subidas, fallas = 0, []
    for f in archivos:
        if subidas >= cupo:
            fallas.append(f"{f.filename}: el protocolo ya tiene el máximo de {MAX_IMAGENES} imágenes")
            continue
        try:
            grande, mini = procesar_imagen(f.read())
        except ValueError as error:
            fallas.append(f"{f.filename}: {error}")
            continue
        except ImportError:
            flash("Falta instalar Pillow para subir imágenes (python -m pip install -r requirements.txt) y reiniciar el sistema.", "error")
            return volver_a_imagenes(pid)
        db.ex("INSERT INTO imagenes (protocolo_id, usuario_id, creado_en, nombre, descripcion, imagen, miniatura) VALUES (?,?,?,?,?,?,?)",
              (pid, yo["id"], db.ahora(), (f.filename or "")[:120], descripcion, grande, mini))
        subidas += 1
    if subidas:
        db.auditar(yo["id"], pid, "imagenes", f"{subidas} imagen(es)")
        flash(f"{subidas} imagen(es) subida(s).", "ok")
    for falla in fallas:
        flash(falla, "error")
    return volver_a_imagenes(pid)


def enviar_imagen(iid, columna):
    r = db.uno(f"SELECT {columna} AS datos FROM imagenes WHERE id=?", (iid,)) or abort(404)
    respuesta = send_file(io.BytesIO(bytes(r["datos"])), mimetype="image/jpeg", conditional=False)
    respuesta.headers["Cache-Control"] = "private, max-age=86400"
    return respuesta


@app.route("/imagen/<int:iid>")
@requiere_login
def imagen_ver(iid):
    return enviar_imagen(iid, "imagen")


@app.route("/imagen/<int:iid>/mini")
@requiere_login
def imagen_mini(iid):
    return enviar_imagen(iid, "miniatura")


@app.route("/imagen/<int:iid>/borrar", methods=["POST"])
@requiere_permiso(*PERMISOS_IMAGENES)
def imagen_borrar(iid):
    r = db.uno("SELECT id, protocolo_id, nombre FROM imagenes WHERE id=?", (iid,)) or abort(404)
    db.ex("DELETE FROM imagenes WHERE id=?", (iid,))
    db.auditar(usuario_actual()["id"], r["protocolo_id"], "imagen_borrada", r["nombre"] or "")
    flash("Imagen borrada.", "ok")
    return volver_a_imagenes(r["protocolo_id"])


def comentarios_de(pid):
    return db.q(SQL_COMENTARIOS, (pid,))


def volver_a_comentarios(pid):
    destino = request.referrer if request.referrer and request.referrer.startswith(request.host_url) else url_for("protocolo", pid=pid)
    return redirect(destino.split("#")[0] + "#comentarios")


def malignidad_por_flujo(estudios):
    salida = []
    for tipo in dict.fromkeys(x["tipo"] for x in estudios if not x["anulado"]):
        activos = [x for x in estudios if x["tipo"] == tipo and not x["anulado"]]
        valores = {x["malignidad"] for x in activos}
        salida.append({"tipo": tipo, "nombre": TIPOS[tipo], "informado": not any(x["borrador"] or x["proxima"] for x in activos),
                       "valor": valores.pop() if len(valores) == 1 else None})
    return salida


@app.route("/buscar")
@requiere_login
def buscar_protocolo():
    texto = request.args.get("q", "").strip()
    if not texto:
        return redirect(url_for("tablero"))
    m = re.fullmatch(r"[cC]?\s*0*(\d{1,6})", texto)
    numero = numero_protocolo(int(m.group(1))) if m and int(m.group(1)) > 0 else texto
    p = db.uno("SELECT id FROM protocolos WHERE numero=?", (numero,)) or db.uno("SELECT id FROM protocolos WHERE LOWER(numero)=?", (texto.lower(),))
    if p:
        return redirect(url_for("protocolo", pid=p["id"]))
    if m:
        flash(f"No existe el protocolo {numero}.", "error")
        return redirect(url_for("tablero"))
    return redirect(url_for("tablero", ver="todos", q=texto))


@app.route("/protocolo/<int:pid>")
@requiere_login
def protocolo(pid):
    db.precargar([(SQL_PROTOCOLO, (pid,)), (SQL_HISTORIAL_PROTOCOLO, (pid,)), (SQL_INFORME_EMITIDO, (pid,)),
                   (SQL_COMENTARIOS, (pid,)), (SQL_IMAGENES, (pid,))] + consultas_estudios("e.protocolo_id=?", (pid,), True))
    p = protocolo_o_404(pid)
    historial = db.q(SQL_HISTORIAL_PROTOCOLO, (pid,))
    estudios = cargar_estudios("e.protocolo_id=?", (pid,), borradores=True)
    pendientes = []
    for x in estudios:
        if not x["anulado"] and x["proxima"] and (x["tipo"], x["proxima"][1]) not in [(q[0], q[1]) for q in pendientes]:
            pendientes.append((x["tipo"], x["proxima"][1], TIPOS[x["tipo"]]))
    lugares = list(dict.fromkeys(l for l in (p["lugar_recoleccion"], p["lugar_entrega"]) if l))
    ubicaciones = {u["abreviatura"]: dict(u) for u in db.q(f"SELECT id, abreviatura, nombre, contactos FROM ubicaciones WHERE abreviatura IN ({','.join('?' * len(lugares))})", tuple(lugares))} if lugares else {}
    destinos = destinos_por_ubicacion({u["id"] for u in ubicaciones.values()}) if ubicaciones else {}
    for u in ubicaciones.values():
        u["emails"] = ", ".join(emails_de_informes(destinos.get(u["id"], [])))
    return render_template("protocolo.html", p=p, estudios=estudios, historial=historial, listas=listas(), comentarios=comentarios_de(pid), imagenes=imagenes_de(pid), ubicaciones=ubicaciones,
                           flujos_malignidad=malignidad_por_flujo(estudios), malignidades=envios.MALIGNIDADES,
                           emitido=informe_emitido(pid),
                           flujos=flujos_de([pid]).get(pid, ""),
                           informe_listo=bool(estudios) and not p["borrador"] and not pendientes and any(not x["anulado"] for x in estudios),
                           pendientes=pendientes, informe_vigente=informe_vigente(pid, informe_emitido(pid)))


@app.route("/protocolo/<int:pid>/malignidad", methods=["POST"])
@requiere_permiso("micro", "informe")
def malignidad_asignar(pid):
    yo = usuario_actual()
    protocolo_o_404(pid)
    tipo = request.form.get("tipo", "")
    valor = request.form.get("valor", "")
    volver = redirect(url_for("protocolo", pid=pid) + "#malignidad")
    if tipo not in TIPOS or (valor and valor not in envios.NOMBRE_MALIGNIDAD):
        flash("Elegí una malignidad válida.", "error")
        return volver
    activos = [x for x in cargar_estudios("e.protocolo_id=?", (pid,), borradores=True) if x["tipo"] == tipo and not x["anulado"]]
    if not activos or any(x["borrador"] or x["proxima"] for x in activos):
        flash("La malignidad se asigna cuando los estudios están informados.", "error")
        return volver
    db.ex("UPDATE estudios SET malignidad=? WHERE protocolo_id=? AND tipo=? AND anulado=0", (valor or None, pid, tipo))
    nombre = envios.NOMBRE_MALIGNIDAD.get(valor, "sin malignidad")
    db.auditar(yo["id"], pid, "malignidad", f"{TIPOS[tipo]}: {nombre}")
    flash(f"Malignidad de {TIPOS[tipo].lower()}: {nombre.lower()}.", "ok")
    return volver


@app.route("/protocolo/<int:pid>/comentario", methods=["POST"])
@requiere_login
def comentario_nuevo(pid):
    protocolo_o_404(pid)
    texto = request.form.get("texto", "").strip()
    if not texto:
        flash("Escribí el comentario.", "error")
    elif len(texto) > 2000:
        flash("El comentario es demasiado largo (máximo 2000 caracteres).", "error")
    else:
        db.ex("INSERT INTO comentarios (protocolo_id, usuario_id, fecha_hora, texto) VALUES (?,?,?,?)",
              (pid, usuario_actual()["id"], db.ahora(), texto))
        db.auditar(usuario_actual()["id"], pid, "comentario", texto[:60] + ("…" if len(texto) > 60 else ""))
        flash("Comentario agregado.", "ok")
    return volver_a_comentarios(pid)


@app.route("/comentario/<int:cid>/borrar", methods=["POST"])
@requiere_login
def comentario_borrar(cid):
    yo = usuario_actual()
    c = db.uno("SELECT * FROM comentarios WHERE id=?", (cid,)) or abort(404)
    if c["usuario_id"] != yo["id"] and not yo["admin"]:
        flash("Solo podés borrar tus propios comentarios.", "error")
    else:
        db.ex("DELETE FROM comentarios WHERE id=?", (cid,))
        db.auditar(yo["id"], c["protocolo_id"], "comentario_borrado", c["texto"][:60] + ("…" if len(c["texto"]) > 60 else ""))
        flash("Comentario borrado.", "ok")
    return volver_a_comentarios(c["protocolo_id"])


@app.route("/protocolo/<int:pid>/estudio/nuevo", methods=["GET", "POST"])
@requiere_permiso("protocolo_editar")
def estudio_nuevo(pid):
    yo = usuario_actual()
    p = protocolo_o_404(pid)
    existentes = [dict(e) for e in db.q("SELECT tipo, subcategoria, anulado FROM estudios WHERE protocolo_id=?", (pid,))]
    lista = [{"tipo": request.args["tipo"]}] if request.args.get("tipo") in REGISTRO else []
    if request.method == "POST":
        lista = estudios_del_formulario()
        lote_del_prot = lote_del_protocolo(pid)
        lote_valor, error = validar_estudios(lista, existentes, con_lote=not primer_estudio(pid))
        if not lista:
            error = "Elegí qué estudio agregar."
        for e in lista:
            if not error and (avance := flujo_avanzado(pid, e["tipo"])):
                error = (f"Ya hay {REGISTRO[e['tipo']].nombre.lower()} de este protocolo en proceso (llegó a {avance}) y comparten el mismo flujo: "
                         "para sumar otra hay que deshacer las etapas antes.")
        if error:
            flash(error, "error")
        else:
            lote = lote_del_prot if primer_estudio(pid) else concretar_lote(lote_valor, yo["id"])
            for e in lista:
                crear_estudio(pid, e, lote, yo["id"])
            avisar_lote_lleno(ubicar_en_lote(pid, yo["id"]))
            flash(f"Se agregó {', '.join(etiqueta(e) for e in lista)} al protocolo {p['numero']}.", "ok")
            return redirect(url_for("protocolo", pid=pid))
    lote_fijo = lote_del_protocolo(pid)
    return render_template("estudio_form.html", p=p, estudios=lista, editando=None, lote_editable=not primer_estudio(pid),
                           lote_fijo=lote_fijo["codigo"] if lote_fijo else None, **contexto_formulario())


@app.route("/estudio/<int:eid>/editar", methods=["GET", "POST"])
@requiere_permiso("protocolo_editar")
def estudio_editar(eid):
    yo = usuario_actual()
    actual = db.uno("SELECT * FROM estudios WHERE id=?", (eid,)) or abort(404)
    p = protocolo_o_404(actual["protocolo_id"])
    e = dict(actual)
    primero = primer_estudio(p["id"])
    es_primero = bool(primero and primero["id"] == eid)
    if request.method == "POST":
        e = {**leer_estudio("e-0-", actual["tipo"]), "id": eid}
        otros = [dict(x) for x in db.q("SELECT tipo, subcategoria, anulado FROM estudios WHERE protocolo_id=? AND id<>?",
                                       (p["id"], eid))]
        lote_valor, error = validar_estudios([{**e, "anulado": actual["anulado"]}], otros, actual["lote_id"], con_lote=es_primero)
        if error:
            flash(error, "error")
        else:
            db.ex(f"UPDATE estudios SET {', '.join(k + '=?' for k in CAMPOS_ESTUDIO)} WHERE id=?",
                  [e[k] for k in CAMPOS_ESTUDIO] + [eid])
            recordar_tecnicas(e)
            db.auditar(yo["id"], p["id"], "editar_estudio", etiqueta(e), eid)
            if es_primero:
                asignar_lote(p["id"], concretar_lote(lote_valor, yo["id"]), yo["id"])
            avisar_lote_lleno(ubicar_en_lote(p["id"], yo["id"]))
            flash("Datos del estudio actualizados.", "ok")
            return redirect(url_for("estudio", eid=eid))
    else:
        e["lote_sel"] = str(actual["lote_id"] or "")
    ctx = contexto_formulario()
    ctx["lotes_abiertos"][actual["tipo"]] = lotes_para_formulario(actual["tipo"], actual["lote_id"])
    lote_fijo = lote_del_protocolo(p["id"])
    return render_template("estudio_form.html", p=p, estudios=[e], editando=eid, lote_editable=es_primero,
                           lote_fijo=lote_fijo["codigo"] if lote_fijo else None, **ctx)


DUENO = ("(SELECT CASE WHEN d.anulado=1 THEN d.id ELSE (SELECT MIN(x.id) FROM estudios x WHERE x.protocolo_id=d.protocolo_id "
         "AND x.tipo=d.tipo AND x.anulado=0) END FROM estudios d WHERE d.id=?)")
SQL_FICHA = """SELECT e.*, p.numero, p.borrador, p.fecha_recoleccion, p.apellido, p.nombre, p.dni, p.medico, p.cobertura,
                         u.iniciales AS responsable, u2.iniciales AS creador, l.codigo AS lote, """ + DUENO.replace("?", "e.id") + """ AS dueno
                  FROM estudios e JOIN protocolos p ON p.id=e.protocolo_id
                  LEFT JOIN usuarios u ON u.id=e.responsable_id LEFT JOIN usuarios u2 ON u2.id=e.creado_por
                  LEFT JOIN lotes l ON l.id=e.lote_id WHERE e.id=?"""
SQL_ETAPAS_FICHA = """SELECT x.*, u.iniciales, u.nombre FROM etapas x
                                            JOIN usuarios u ON u.id=x.usuario_id WHERE estudio_id=?"""
SQL_ULTIMA_ETAPA = "SELECT * FROM etapas WHERE estudio_id=? ORDER BY fecha_hora DESC, id DESC LIMIT 1"
SQL_HISTORIAL_ESTUDIO = """SELECT a.*, u.iniciales FROM auditoria a LEFT JOIN usuarios u ON u.id=a.usuario_id
                        WHERE a.estudio_id IN (SELECT x.id FROM estudios x JOIN estudios d ON d.protocolo_id=x.protocolo_id AND d.tipo=x.tipo
                                               WHERE d.id=? AND (d.anulado=1 AND x.id=d.id OR d.anulado=0 AND x.anulado=0))
                        ORDER BY a.id DESC"""
SQL_MACRO = "SELECT * FROM macro WHERE estudio_id=" + DUENO
SQL_MICRO = "SELECT * FROM micro WHERE estudio_id=" + DUENO
SQL_IHQ = "SELECT * FROM ihq WHERE estudio_id=" + DUENO
SQL_TEMPLATES = "SELECT id, titulo, tipo_biopsia AS zona, organo, tipo_muestra AS tipo FROM templates WHERE clase=? ORDER BY titulo"
SQL_LISTAS = "SELECT nombre, valor FROM listas ORDER BY nombre, orden"
SQL_VARIANTE = ("SELECT (SELECT 1 FROM micro m JOIN estudios e ON e.id=m.estudio_id "
                "WHERE e.protocolo_id=? AND e.tipo=? AND e.anulado=0 AND m.solicita_ihq=1) AS ihq, "
                "(SELECT 1 FROM macro m JOIN estudios e ON e.id=m.estudio_id "
                "WHERE e.protocolo_id=? AND e.tipo=? AND e.anulado=0 AND m.decal=1) AS decal")


def consultas_ficha(eid):
    return [(SQL_FICHA, (eid,)), (SQL_MACRO, (eid,)), (SQL_MICRO, (eid,)), (SQL_IHQ, (eid,)), (SQL_ETAPAS_FICHA, (eid,))]


def ficha(eid):
    db.precargar(consultas_ficha(eid))
    e = db.uno(SQL_FICHA, (eid,))
    if not e:
        abort(404)
    macro = db.uno(SQL_MACRO, (eid,))
    micro = db.uno(SQL_MICRO, (eid,))
    ihq = db.uno(SQL_IHQ, (eid,))
    etapas = {x["etapa"]: x for x in db.q(SQL_ETAPAS_FICHA, (eid,))}
    solicita = (Variante(micro and micro["solicita_ihq"], macro and macro["decal"]) if e["anulado"]
                else variante_de(e["protocolo_id"], e["tipo"]))
    return e, macro, micro, ihq, etapas, solicita


def con_permiso(permiso, estado):
    return estado if puede(permiso) or not estado[0] else (False, "No tenés permiso para cargar esta sección.")


def habilitada(e, etapas, solicita, etapa):
    if e["anulado"]:
        return False, "El estudio está anulado."
    _, prox = REGISTRO[e["tipo"]].estado(set(etapas), solicita)
    if prox and prox[0] == etapa:
        return True, ""
    if etapa in etapas:
        x = etapas[etapa]
        return False, (f"{NOMBRE_ETAPA[etapa]} ya fue completada por {x['iniciales']} el {fecha_hora(x['fecha_hora'])}. "
                       "Para corregirla hay que deshacer las etapas hasta esa.")
    return False, (f"Se habilita cuando el estudio llegue a {NOMBRE_ETAPA[etapa]}"
                   + (f" (ahora está en {prox[1]})." if prox else "."))


def macro_editable(e, etapas, solicita):
    en_macro = habilitada(e, etapas, solicita, "macroscopia")
    if en_macro[0]:
        return con_permiso("macro", en_macro)
    en_micro = habilitada(e, etapas, solicita, "microscopia")
    if en_micro[0] and "macro" in REGISTRO[e["tipo"]].secciones:
        return (True, "") if puede("macro") or puede("micro") else (False, "No tenés permiso para cargar esta sección.")
    return en_macro


SQL_TECNICAS = """SELECT t.id, t.estudio_id, t.numero, t.fragmentos, t.inclusion, t.tacos, t.cassettes, t.comentarios, t.creado_en, u.iniciales
                  FROM tecnicas t JOIN estudios e ON e.id=t.estudio_id LEFT JOIN usuarios u ON u.id=t.usuario_id
                  WHERE e.protocolo_id=? AND e.tipo=? AND e.anulado=0 ORDER BY t.id"""
INCLUSIONES = ("Reserva", "Inclusión total")


def estudios_sin_tecnicas(miembros):
    if not miembros:
        return []
    con = {r["estudio_id"] for r in db.q(f"SELECT DISTINCT estudio_id FROM tecnicas WHERE estudio_id IN ({','.join('?' * len(miembros))})", [m["id"] for m in miembros])}
    return [m for m in miembros if m["id"] not in con]


def sumar_cassettes(e):
    total = db.uno("""SELECT SUM(t.cassettes) AS n FROM tecnicas t JOIN estudios x ON x.id=t.estudio_id
                      WHERE x.protocolo_id=? AND x.tipo=? AND x.anulado=0""", (e["protocolo_id"], e["tipo"]))["n"]
    guardar_seccion("macro", e["dueno"], {"cassettes": total})


@app.route("/estudio/<int:eid>")
@requiere_login
def estudio(eid):
    db.precargar(consultas_ficha(eid) + [(SQL_ULTIMA_ETAPA, (eid,)), (SQL_HISTORIAL_ESTUDIO, (eid,)), (SQL_TEMPLATES, ("macro",)),
                                         (SQL_TEMPLATES, ("micro",)), (SQL_LISTAS, ()), ("SELECT fecha FROM feriados", ())])
    primero = db.uno(SQL_FICHA, (eid,))
    if primero:
        db.precargar([(SQL_VARIANTE, (primero["protocolo_id"], primero["tipo"], primero["protocolo_id"], primero["tipo"])), (SQL_COMENTARIOS, (primero["protocolo_id"],)), (SQL_IMAGENES, (primero["protocolo_id"],)), (SQL_TECNICAS, (primero["protocolo_id"], primero["tipo"]))] + consultas_estudios("e.protocolo_id=?", (primero["protocolo_id"],), True))
    e, macro, micro, ihq, etapas, solicita = ficha(eid)
    tipo = REGISTRO[e["tipo"]]
    est, prox = tipo.estado(set(etapas), solicita, bool(e["anulado"]))
    ultima = db.uno(SQL_ULTIMA_ETAPA, (eid,))
    historial = db.q(SQL_HISTORIAL_ESTUDIO, (eid,))
    tpl = lambda clase: db.q(SQL_TEMPLATES, (clase,))
    pasos = tipo.flujo(solicita)
    decal = con_decal(solicita)
    traza_enc, traza = trazabilidad.calcular(e, pasos, etapas, db.feriados(), decal=decal)
    hermanos = cargar_estudios("e.protocolo_id=?", (e["protocolo_id"],), borradores=True)
    activos = [h for h in hermanos if not h["anulado"]]
    miembros = [h for h in hermanos if h["id"] == e["id"]] if e["anulado"] else [h for h in activos if h["tipo"] == e["tipo"]]
    pestanas, vistos = [], set()
    for h in hermanos:
        if h["anulado"]:
            pestanas.append({"id": h["id"], "texto": f"{h['etiqueta']} (anulado)", "activa": e["anulado"] and h["id"] == e["id"], "anulado": True})
        elif h["tipo"] not in vistos:
            vistos.add(h["tipo"])
            n = sum(1 for x in activos if x["tipo"] == h["tipo"])
            pestanas.append({"id": h["id"], "texto": h["etiqueta"] if n == 1 else f"{REGISTRO[h['tipo']].nombre} ({n} estudios)",
                             "activa": not e["anulado"] and h["tipo"] == e["tipo"], "anulado": False})
    limite_etapa = trazabilidad.limite(e["tipo"], prox[0], e["fecha_recoleccion"] or e["creado_en"], db.feriados(), decal) if prox and not e["anulado"] and not e["borrador"] else None
    semaforo = (trazabilidad.semaforo_acumulado(e["tipo"], e["fecha_recoleccion"] or e["creado_en"],
                                                {"ingreso": e["creado_en"], **{k: v["fecha_hora"] for k, v in etapas.items()}}, limite_etapa, datetime.now(), db.feriados(), decal)
                if limite_etapa else "")
    cierra_protocolo = bool(prox and not e["anulado"] and prox[0] == tipo.ultima_etapa(solicita)
                            and all(not h["proxima"] and not h["borrador"] for h in activos if h["tipo"] != e["tipo"]))
    tecnicas = db.q(SQL_TECNICAS, (e["protocolo_id"], e["tipo"])) if "macro" in tipo.secciones and not e["anulado"] else []
    return render_template("estudio.html", tecnicas=tecnicas, sin_tecnicas={m["id"] for m in estudios_sin_tecnicas(miembros)}, e=e, miembros=miembros, pestanas=pestanas, comentarios=comentarios_de(e["protocolo_id"]), imagenes=imagenes_de(e["protocolo_id"]), cierra_protocolo=cierra_protocolo, semaforo=semaforo,
                           flujos=flujos_de([e["protocolo_id"]]).get(e["protocolo_id"], ""), tipo=tipo, macro=macro, micro=micro, ihq=ihq, etapas=etapas, estado=est,
                           prox=prox, ultima=ultima, historial=historial, traza=traza, traza_enc=traza_enc,
                           hermanos=hermanos, etiqueta=etiqueta(e) if len(miembros) < 2 else tipo.nombre,
                           tpl_macro=tpl("macro"), tpl_micro=tpl("micro"), bethesda=listas().get("bethesda", []),
                           tecnicas_pedidas=sin_repetidos(x.strip() for m in miembros for x in (m["tecnicas_especiales"] or "").split(";")),
                           lista_ihq=listas().get("ihq_marcadores", []),
                           filtro_templates={"zona": e["subcategoria"] or "", "organo": e["sitio"] or "", "tipo": e["tipo_muestra"] or ""},
                           ihq_hecha=any(k in etapas for k in PASOS_IHQ),
                           edita_macro=macro_editable(e, etapas, solicita),
                           decal_cambiable=macro_editable(e, etapas, solicita)[0] and "procesamiento" not in etapas and "decal" not in etapas,
                           edita_micro=con_permiso("micro", habilitada(e, etapas, solicita, "microscopia")))


def guardar_seccion(tabla, eid, valores):
    if db.uno(f"SELECT 1 FROM {tabla} WHERE estudio_id=?", (eid,)):
        db.ex(f"UPDATE {tabla} SET {', '.join(k + '=?' for k in valores)} WHERE estudio_id=?", list(valores.values()) + [eid])
    else:
        db.ex(f"INSERT INTO {tabla} (estudio_id, {', '.join(valores)}) VALUES (?,{','.join('?' * len(valores))})",
              [eid] + list(valores.values()))


@app.route("/estudio/<int:eid>/macro", methods=["POST"])
@requiere_permiso("macro", "micro")
def guardar_macro(eid):
    e, macro, _, _, etapas, solicita = ficha(eid)
    ok, motivo = macro_editable(e, etapas, solicita)
    if "macro" not in REGISTRO[e["tipo"]].secciones or not ok:
        flash(f"No se guardó: {motivo or 'este estudio no tiene macroscopía.'}", "error")
        return redirect(url_for("estudio", eid=eid) + "#macro")
    valores = {"template": request.form.get("template", "").strip(),
               "descripcion": request.form.get("descripcion", "").strip()}
    detalle = ""
    if request.form.get("decal_presente") == "1" and "procesamiento" not in etapas and "decal" not in etapas:
        decal = 1 if request.form.get("decal") == "1" else 0
        valores["decal"] = decal
        if decal != (1 if macro and macro["decal"] else 0):
            detalle = "agrega Macro > Decal" if decal else "quita Macro > Decal"
    guardar_seccion("macro", e["dueno"], valores)
    db.auditar(usuario_actual()["id"], e["protocolo_id"], "guardar_macro", detalle, eid)
    flash("Macroscopía guardada.", "ok")
    if request.form.get("y_listo"):
        return marcar_listo(eid, "macroscopia")
    return redirect(url_for("estudio", eid=eid) + "#macro")


def entero_positivo(valor, maximo=999):
    valor = (valor or "").strip()
    return int(valor) if valor.isdigit() and 1 <= int(valor) <= maximo else None


@app.route("/estudio/<int:eid>/tecnicas", methods=["POST"])
@requiere_permiso("macro", "micro")
def tecnica_nueva(eid):
    e, _, _, _, etapas, solicita = ficha(eid)
    ok, motivo = macro_editable(e, etapas, solicita)
    destino = url_for("estudio", eid=eid) + "#tecnicas"
    if "macro" not in REGISTRO[e["tipo"]].secciones or not ok:
        flash(f"No se guardó: {motivo or 'este estudio no tiene macroscopía.'}", "error")
        return redirect(destino)
    miembro = db.uno("SELECT id FROM estudios WHERE id=? AND protocolo_id=? AND tipo=? AND anulado=0",
                     (request.form.get("estudio_id", "") if request.form.get("estudio_id", "").isdigit() else 0, e["protocolo_id"], e["tipo"]))
    inclusion = request.form.get("inclusion", "")
    fragmentos, cassettes = entero_positivo(request.form.get("fragmentos")), entero_positivo(request.form.get("cassettes"))
    numero = request.form.get("numero", "").strip()[:10]
    if not miembro:
        flash("Elegí el estudio de la técnica.", "error")
    elif inclusion not in INCLUSIONES:
        flash("Elegí el tipo de inclusión.", "error")
    elif not numero or not fragmentos or not cassettes:
        flash("Completá el número, los fragmentos y los cassettes de la técnica.", "error")
    else:
        tacos = cargar_estudios("e.id=?", (miembro["id"],))[0]["tacos"]
        db.ex("""INSERT INTO tecnicas (estudio_id, numero, fragmentos, inclusion, tacos, cassettes, comentarios, usuario_id, creado_en)
                 VALUES (?,?,?,?,?,?,?,?,?)""", (miembro["id"], numero, fragmentos, inclusion, tacos, cassettes,
                                                  request.form.get("comentarios", "").strip()[:300], usuario_actual()["id"], db.ahora()))
        sumar_cassettes(e)
        db.auditar(usuario_actual()["id"], e["protocolo_id"], "tecnica", f"{numero}: {inclusion}, {fragmentos} fragm., {cassettes} cassette(s)", miembro["id"])
        flash("Técnica agregada.", "ok")
    return redirect(destino)


@app.route("/tecnica/<int:tid>/borrar", methods=["POST"])
@requiere_permiso("macro", "micro")
def tecnica_borrar(tid):
    t = db.uno("SELECT id, estudio_id, numero FROM tecnicas WHERE id=?", (tid,)) or abort(404)
    e, _, _, _, etapas, solicita = ficha(t["estudio_id"])
    ok, motivo = macro_editable(e, etapas, solicita)
    destino = url_for("estudio", eid=t["estudio_id"]) + "#tecnicas"
    if not ok:
        flash(f"No se borró: {motivo}", "error")
        return redirect(destino)
    db.ex("DELETE FROM tecnicas WHERE id=?", (tid,))
    sumar_cassettes(e)
    db.auditar(usuario_actual()["id"], e["protocolo_id"], "tecnica_borrada", t["numero"] or "", t["estudio_id"])
    flash("Técnica borrada.", "ok")
    return redirect(destino)


@app.route("/estudio/<int:eid>/micro", methods=["POST"])
@requiere_permiso("micro")
def guardar_micro(eid):
    e, _, micro, _, etapas, solicita = ficha(eid)
    ok, motivo = habilitada(e, etapas, solicita, "microscopia")
    if not ok:
        flash(f"No se guardó: {motivo}", "error")
        return redirect(url_for("estudio", eid=eid) + "#micro")
    con_ihq = "ihq" in REGISTRO[e["tipo"]].secciones
    pide_ihq = 1 if con_ihq and request.form.get("solicita_ihq") == "1" else 0
    guardar_seccion("micro", e["dueno"], {"template": request.form.get("template", "").strip(),
                                   "descripcion": request.form.get("descripcion", "").strip(),
                                   "conclusion": request.form.get("conclusion", "").strip(),
                                   "bethesda": request.form.get("bethesda", "").strip(),
                                   "solicita_ihq": pide_ihq})
    db.auditar(usuario_actual()["id"], e["protocolo_id"], "guardar_micro", "solicita IHQ" if pide_ihq else "", eid)
    flash("Microscopía guardada.", "ok")
    if request.form.get("y_listo"):
        return marcar_listo(eid, "microscopia")
    return redirect(url_for("estudio", eid=eid) + "#micro")


@app.route("/estudio/<int:eid>/ihq", methods=["POST"])
@requiere_permiso("micro")
def guardar_ihq(eid):
    e = ficha(eid)[0]
    guardar_seccion("ihq", e["dueno"], {"marcadores": request.form.get("marcadores", "").strip(),
                                 "resultado": request.form.get("resultado", "").strip()})
    db.auditar(usuario_actual()["id"], e["protocolo_id"], "guardar_ihq", "", eid)
    flash("IHQ guardada.", "ok")
    if request.form.get("y_listo"):
        return marcar_listo(eid, "interp_ihq")
    return redirect(url_for("estudio", eid=eid) + "#ihq")


def motivo_no_listo(yo, tipo, etapa, prox, solicita, macro, micro, ihq, sin_tecnicas):
    falta = None if etapa in PASOS_IHQ and not solicita else REGISTRO[tipo].requisito(etapa, macro, micro, ihq)
    if not falta and sin_tecnicas:
        falta = "Faltan las técnicas de: " + ", ".join(etiqueta(m) for m in sin_tecnicas) + "."
    if not prox or prox[0] != etapa:
        return "Esa etapa no es la próxima pendiente de este estudio."
    if not yo["admin"] and prox[2] not in (yo["sectores"] or "").split(","):
        return f"La etapa {prox[1]} la marca el sector {SECTORES[prox[2]]}."
    return falta


@app.route("/estudio/<int:eid>/listo/<etapa>", methods=["POST"])
@requiere_permiso("etapas")
def marcar_listo(eid, etapa):
    yo = usuario_actual()
    e, macro, micro, ihq, etapas, solicita = ficha(eid)
    if e["borrador"]:
        flash("Primero completá los datos del paciente de este protocolo.", "error")
        return redirect(url_for("protocolo_editar", pid=e["protocolo_id"]))
    tipo = REGISTRO[e["tipo"]]
    _, prox = tipo.estado(set(etapas), solicita, bool(e["anulado"]))
    miembros = miembros_flujo(e)
    sin = estudios_sin_tecnicas(miembros) if etapa == "macroscopia" and "macro" in tipo.secciones else []
    motivo = motivo_no_listo(yo, e["tipo"], etapa, prox, solicita, macro, micro, ihq, sin)
    if motivo:
        flash(motivo, "error")
    else:
        ahora_ = db.ahora()
        for m in miembros:
            if not db.uno("SELECT 1 FROM etapas WHERE estudio_id=? AND etapa=?", (m["id"], etapa)):
                db.ex("INSERT INTO etapas (estudio_id, etapa, usuario_id, fecha_hora) VALUES (?,?,?,?)", (m["id"], etapa, yo["id"], ahora_))
        db.auditar(yo["id"], e["protocolo_id"], "listo", NOMBRE_ETAPA[etapa] + (f" (en {len(miembros)} estudios {tipo.nombre.lower()})" if len(miembros) > 1 else ""), eid)
        flash(f"✔ {NOMBRE_ETAPA[etapa]} registrada por {yo['iniciales']}"
              + (f" en los {len(miembros)} estudios de {tipo.nombre.lower()} del protocolo." if len(miembros) > 1 else "."), "ok")
        if request.form.get("firmar_yo") == "1" and protocolo_cerrado(e["protocolo_id"]):
            db.ex("UPDATE protocolos SET firmante_id=? WHERE id=?", (yo["id"], e["protocolo_id"]))
            db.auditar(yo["id"], e["protocolo_id"], "firmante", f"{yo['iniciales']} firmará el informe", eid)
            flash("Vas a figurar como firmante del informe.", "ok")
    return redirect(request.form.get("volver") or url_for("estudio", eid=eid))


SQL_INFORME_EMITIDO = """SELECT i.id, i.version, i.creado_en, i.comentario, i.estudios, u.iniciales, u.nombre AS generado_por, f.nombre AS firmante
                     FROM informes_emitidos i LEFT JOIN usuarios u ON u.id=i.usuario_id LEFT JOIN usuarios f ON f.id=i.firmante_id
                     WHERE i.protocolo_id=?"""


def informe_emitido(pid):
    return db.uno(SQL_INFORME_EMITIDO, (pid,))


def informe_vigente(pid, emitido):
    if not emitido or not protocolo_cerrado(pid):
        return False
    if not emitido["estudios"]:
        return True
    activos = {str(x["id"]) for x in cargar_estudios("e.protocolo_id=?", (pid,), borradores=True) if not x["anulado"]}
    return set(emitido["estudios"].split(",")) == activos


@app.route("/estudio/<int:eid>/informe")
@requiere_permiso("informe")
def informe_de_estudio(eid):
    return redirect(url_for("informe", pid=ficha(eid)[0]["protocolo_id"]))


@app.route("/protocolo/<int:pid>/informe")
@requiere_permiso("informe")
def informe(pid):
    p = protocolo_o_404(pid)
    if informes is None:
        flash(SIN_INFORMES, "error")
        return redirect(url_for("protocolo", pid=pid))
    try:
        d = informes.armar(pid)
    except informes.InformeError as err:
        flash(str(err), "error")
        return redirect(url_for("protocolo", pid=pid))
    emitido = informe_emitido(pid)
    return render_template("informe.html", d=d, p=p, emitido=emitido, informe_vigente=informe_vigente(pid, emitido))


@app.route("/protocolo/<int:pid>/informe.pdf", methods=["POST"])
@requiere_permiso("informe")
def informe_pdf(pid):
    protocolo_o_404(pid)
    if informes is None:
        flash(SIN_INFORMES, "error")
        return redirect(url_for("protocolo", pid=pid))
    yo = usuario_actual()
    try:
        d = informes.armar(pid, request.form.get("comentario", ""))
    except informes.InformeError as err:
        flash(str(err), "error")
        return redirect(url_for("protocolo", pid=pid))
    pdf = informes.generar_pdf(d)
    previo = db.uno("SELECT i.id, i.version, i.creado_en, u.iniciales FROM informes_emitidos i LEFT JOIN usuarios u ON u.id=i.usuario_id WHERE i.protocolo_id=?", (pid,))
    campos = (yo["id"], d["medico"]["id"], db.ahora(), d["comentario"], informes.nombre_archivo(d), pdf, ",".join(str(i) for s in d["estudios"] for i in s["eids"]))
    try:
        if previo:
            db.ex("UPDATE informes_emitidos SET usuario_id=?, firmante_id=?, creado_en=?, comentario=?, archivo=?, pdf=?, estudios=?, version=version+1 WHERE id=?",
                  campos + (previo["id"],))
            iid, version = previo["id"], previo["version"] + 1
        else:
            iid, version = db.ex("INSERT INTO informes_emitidos (usuario_id, firmante_id, creado_en, comentario, archivo, pdf, estudios, protocolo_id, version) "
                                 "VALUES (?,?,?,?,?,?,?,?,1)", campos + (pid,)), 1
    except db.IntegrityError:
        flash("Alguien más generó el informe al mismo tiempo, revisalo y generalo de nuevo si hace falta.", "error")
        return redirect(url_for("informe", pid=pid))
    reemplaza = f" · reemplaza el de {previo['iniciales'] or '—'} del {fecha_hora(previo['creado_en'])}" if previo else ""
    db.auditar(yo["id"], pid, "informe_pdf",
               f"{d['titulo'].title()} generado por {yo['iniciales']} · firma de {d['medico']['nombre']}" + (f" (emisión {version}{reemplaza})" if previo else ""))
    return redirect(url_for("informe_ver", iid=iid))


@app.route("/informe/<int:iid>.pdf")
@requiere_permiso("informe")
def informe_ver(iid):
    r = db.uno("SELECT protocolo_id, estudios, archivo, pdf FROM informes_emitidos WHERE id=?", (iid,))
    if not r:
        abort(404)
    if not informe_vigente(r["protocolo_id"], r):
        flash("Ese informe ya no se puede entregar: el protocolo tiene estudios pendientes o cambió desde que se generó. Hay que generarlo de nuevo.", "error")
        return redirect(url_for("protocolo", pid=r["protocolo_id"]))
    return send_file(io.BytesIO(bytes(r["pdf"])), mimetype="application/pdf", download_name=r["archivo"], max_age=0)


SQL_ENVIOS = """SELECT v.id, v.canal, v.rol, v.destino, v.estado, v.detalle, v.version, v.creado_en, u.iniciales
                FROM envios v LEFT JOIN usuarios u ON u.id=v.usuario_id WHERE v.protocolo_id=? ORDER BY v.id DESC"""
ROLES_ENVIO = {"paciente": "Paciente", "ubicacion": "Ubicación", "otro": "Otro"}


def textos_envio(p):
    centro = informes.CENTRO if informes else {}
    contacto = " o por ".join(x for x in (f"teléfono al {centro['telefono']}" if centro.get("telefono") else "",
                                           f"e-mail a {centro['email']}" if centro.get("email") else "") if x)
    firma = (centro.get("nombre", "CAP Vighi") + (" · " + centro["subtitulo"].title() if centro.get("subtitulo") else "")).strip()
    cuerpo = (f"Buenos días,\n\nAdjuntamos el informe de anatomía patológica correspondiente al protocolo {p['numero']}."
              + (f"\n\nAnte cualquier consulta pueden comunicarse con el laboratorio por {contacto}." if contacto else "")
              + f"\n\nSaludos cordiales,\n{firma}")
    return {"asunto": f"Informe de anatomía patológica · Protocolo {p['numero']}", "cuerpo": cuerpo,
            "whatsapp": f"Hola, te enviamos el informe de anatomía patológica del protocolo {p['numero']}. Ante cualquier consulta, escribinos por acá. {centro.get('nombre', 'CAP Vighi')}"}


def estudios_para_envio(pid):
    return [(r["tipo"], r["malignidad"]) for r in db.q("SELECT tipo, malignidad FROM estudios WHERE protocolo_id=? AND anulado=0", (pid,))]


def destinos_por_ubicacion(ids=None):
    filas = db.q("SELECT id, ubicacion_id, canal, destino, informes, alcance FROM ubicacion_destinos ORDER BY canal, destino")
    salida = {}
    for f in filas:
        if ids is None or f["ubicacion_id"] in ids:
            salida.setdefault(f["ubicacion_id"], []).append(f)
    return salida


def emails_de_informes(destinos):
    return [d["destino"] for d in destinos if d["canal"] == "email" and d["informes"]]


def lugares_de_envio(p):
    claves = list(dict.fromkeys(l for l in (p["lugar_entrega"], p["lugar_recoleccion"]) if l))
    if not claves:
        return []
    filas = {u["abreviatura"]: u for u in db.q(f"SELECT id, abreviatura, nombre, contactos FROM ubicaciones WHERE abreviatura IN ({','.join('?' * len(claves))})", tuple(claves))}
    destinos = destinos_por_ubicacion({u["id"] for u in filas.values()})
    estudios = estudios_para_envio(p["id"])
    salida = []
    for clave in claves:
        u = filas.get(clave)
        propios = destinos.get(u["id"], []) if u else []
        salida.append({"abreviatura": clave, "nombre": (u["nombre"] if u else "") or "", "contactos": (u["contactos"] if u else "") or "", "id": u["id"] if u else None,
                       "emails": [dict(d, estado=envios.evaluar(d["alcance"], estudios), texto=envios.alcance_texto(d["alcance"])) for d in propios if d["canal"] == "email" and d["informes"]],
                       "telefonos": [dict(d, estado=envios.evaluar(d["alcance"], estudios), texto=envios.alcance_texto(d["alcance"]), visible=envios.whatsapp_visible(d["destino"]))
                                     for d in propios if d["canal"] == "whatsapp" and d["informes"]],
                       "ocultos": sum(1 for d in propios if not d["informes"]), "es_entrega": clave == p["lugar_entrega"]})
    return salida


def informe_para_enviar(pid):
    emitido = informe_emitido(pid)
    return emitido, bool(emitido) and informe_vigente(pid, emitido)


@app.route("/protocolo/<int:pid>/enviar")
@requiere_permiso("informe")
def envio(pid):
    db.precargar([(SQL_PROTOCOLO, (pid,)), (SQL_INFORME_EMITIDO, (pid,)), (SQL_ENVIOS, (pid,))])
    p = protocolo_o_404(pid)
    if p["borrador"]:
        flash("El protocolo todavía está por completar.", "error")
        return redirect(url_for("protocolo", pid=pid))
    emitido, vigente = informe_para_enviar(pid)
    historial = db.q(SQL_ENVIOS, (pid,))
    mandados = {}
    for v in historial:
        if v["estado"] == "enviado" and v["canal"] == "email" and emitido and v["version"] == emitido["version"]:
            mandados.setdefault(v["destino"], v["creado_en"])
    return render_template("envio.html", flujos_malignidad=malignidad_por_flujo(cargar_estudios("e.protocolo_id=?", (pid,), borradores=True)), malignidades=envios.NOMBRE_MALIGNIDAD,
                           p=p, emitido=emitido, vigente=vigente, historial=historial, mandados=mandados, lugares=lugares_de_envio(p),
                           textos=textos_envio(p), email_ok=envios.configurado(), faltante=envios.faltante(), roles=ROLES_ENVIO,
                           telefono_paciente=envios.whatsapp_visible(envios.whatsapp_numero(p["telefono"])) or (p["telefono"] or ""))


@app.route("/protocolo/<int:pid>/enviar", methods=["POST"])
@requiere_permiso("informe")
def envio_enviar(pid):
    yo = usuario_actual()
    p = protocolo_o_404(pid)
    volver = redirect(url_for("envio", pid=pid))
    emitido, vigente = informe_para_enviar(pid)
    if not vigente:
        flash("Primero hay que generar el informe: el guardado no existe o cambió el protocolo.", "error")
        return volver
    if not envios.configurado():
        flash("El envío por e-mail no está configurado en este servidor.", "error")
        return volver
    pedidos = []
    if request.form.get("usar_paciente") == "1":
        correo = request.form.get("paciente_email", "").strip().lower()
        if not envios.email_valido(correo):
            flash("El e-mail del paciente no es válido.", "error")
            return volver
        pedidos.append(("paciente", correo))
    for correo in request.form.getlist("ubic_email"):
        if not envios.email_valido(correo):
            flash(f"El e-mail {correo[:60]} no es válido.", "error")
            return volver
        pedidos.append(("ubicacion", correo.strip().lower()))
    otros = request.form.get("otro_email", "").strip()
    if otros:
        sueltos = [x for x in re.split(r"[\s,;]+", otros) if x]
        malos = [x for x in sueltos if not envios.email_valido(x)]
        if malos:
            flash(f"El e-mail {malos[0][:60]} no es válido.", "error")
            return volver
        pedidos += [("otro", x.lower()) for x in sueltos]
    vistos = {}
    for rol, correo in pedidos:
        vistos.setdefault(correo, rol)
    if not vistos:
        flash("Elegí a quién enviar el informe.", "error")
        return volver
    if len(vistos) > envios.MAX_DESTINOS:
        flash(f"Se pueden enviar hasta {envios.MAX_DESTINOS} e-mails por vez.", "error")
        return volver
    asunto = " ".join(request.form.get("asunto", "").split())[:150] or textos_envio(p)["asunto"]
    cuerpo = request.form.get("cuerpo", "").replace("\r\n", "\n").strip()[:3000] or textos_envio(p)["cuerpo"]
    guardar_correo = (request.form.get("guardar_paciente") == "1" and request.form.get("usar_paciente") == "1" and puede("protocolo_editar")
                      and request.form.get("paciente_email", "").strip().lower() != (p["email"] or "").lower())
    pdf = db.uno("SELECT pdf, archivo FROM informes_emitidos WHERE id=?", (emitido["id"],))
    resultados = envios.enviar_emails(list(vistos), asunto, cuerpo, bytes(pdf["pdf"]), pdf["archivo"] or "informe.pdf")
    ahora_ = db.ahora()
    sentencias = [("INSERT INTO envios (protocolo_id, informe_id, version, canal, rol, destino, estado, detalle, usuario_id, creado_en) VALUES (?,?,?,?,?,?,?,?,?,?)",
                   (pid, emitido["id"], emitido["version"], "email", vistos[destino], destino, "enviado" if ok else "error", detalle or None, yo["id"], ahora_))
                  for destino, ok, detalle in resultados]
    bien = [r[0] for r in resultados if r[1]]
    mal = [r for r in resultados if not r[1]]
    sentencias.append(("INSERT INTO auditoria (fecha_hora, usuario_id, protocolo_id, estudio_id, accion, detalle) VALUES (?,?,?,?,?,?)",
                       (ahora_, yo["id"], pid, None, "informe_envio",
                        f"Informe enviado por e-mail a {', '.join(bien)}" + (f" · falló: {', '.join(r[0] for r in mal)}" if mal else "") if bien
                        else f"No se pudo enviar el informe por e-mail a {', '.join(r[0] for r in mal)}")))
    db.lote_ex(sentencias)
    if guardar_correo:
        nuevo = request.form.get("paciente_email", "").strip().lower()
        db.ex("UPDATE protocolos SET email=? WHERE id=?", (nuevo, pid))
        sincronizar_paciente(pid, p["dni"], ["email"])
        db.auditar(yo["id"], pid, "editar_paciente", f"E-mail del paciente: {nuevo}")
    if bien:
        flash(f"✔ Informe enviado a {len(bien)} destinatario{'s' if len(bien) != 1 else ''}: {', '.join(bien)}.", "ok")
    for destino, _, detalle in mal:
        flash(f"No se pudo enviar a {destino}: {detalle}", "error")
    return volver


@app.route("/protocolo/<int:pid>/enviar/whatsapp", methods=["POST"])
@requiere_permiso("informe")
def envio_whatsapp(pid):
    yo = usuario_actual()
    protocolo_o_404(pid)
    emitido, vigente = informe_para_enviar(pid)
    if not vigente:
        return jsonify(error="Primero hay que generar el informe: el guardado no existe o cambió el protocolo."), 400
    numero = envios.whatsapp_numero(request.form.get("numero", ""))
    if not numero:
        return jsonify(error="El número no es válido. Escribilo con el código de área, por ejemplo 11 5015 1099."), 400
    rol = request.form.get("rol") if request.form.get("rol") in ROLES_ENVIO else "otro"
    texto = " ".join(request.form.get("texto", "").split())[:700]
    db.lote_ex([("INSERT INTO envios (protocolo_id, informe_id, version, canal, rol, destino, estado, detalle, usuario_id, creado_en) VALUES (?,?,?,?,?,?,?,?,?,?)",
                 (pid, emitido["id"], emitido["version"], "whatsapp", rol, "+" + numero, "abierto", "Se abrió WhatsApp: el PDF lo adjunta quien envía", yo["id"], db.ahora())),
                ("INSERT INTO auditoria (fecha_hora, usuario_id, protocolo_id, estudio_id, accion, detalle) VALUES (?,?,?,?,?,?)",
                 (db.ahora(), yo["id"], pid, None, "informe_envio", f"Se abrió WhatsApp para enviar el informe a +{numero}"))])
    return jsonify(enlace=f"https://wa.me/{numero}" + (f"?text={quote(texto)}" if texto else ""))


@app.route("/usuarios/<int:uid>/firma")
@requiere_admin
def usuario_firma(uid):
    f = db.uno("SELECT imagen FROM firmas WHERE usuario_id=?", (uid,))
    if not f:
        abort(404)
    return send_file(io.BytesIO(bytes(f["imagen"])), mimetype="image/png", max_age=0)


@app.route("/estudio/<int:eid>/deshacer", methods=["POST"])
@requiere_permiso("etapas")
def deshacer(eid):
    yo = usuario_actual()
    e = ficha(eid)[0]
    ultima = db.uno("SELECT * FROM etapas WHERE estudio_id=? ORDER BY fecha_hora DESC, id DESC LIMIT 1", (eid,))
    if not ultima:
        flash("No hay etapas para deshacer.", "error")
    elif ultima["usuario_id"] != yo["id"] and not puede("etapas_deshacer"):
        flash("Solo quien la registró o alguien con permiso puede deshacer la etapa.", "error")
    else:
        miembros = miembros_flujo(e)
        for m in miembros or [e]:
            db.ex("DELETE FROM etapas WHERE estudio_id=? AND etapa=?", (m["id"], ultima["etapa"]))
        db.ex("UPDATE protocolos SET firmante_id=NULL WHERE id=?", (e["protocolo_id"],))
        db.auditar(yo["id"], e["protocolo_id"], "deshacer", NOMBRE_ETAPA.get(ultima["etapa"], ultima["etapa"]), eid)
        flash(f"Se deshizo {NOMBRE_ETAPA.get(ultima['etapa'])}.", "ok")
    return redirect(url_for("estudio", eid=eid))


@app.route("/estudio/<int:eid>/anular", methods=["POST"])
@requiere_permiso("estudio_anular")
def anular(eid):
    e = ficha(eid)[0]
    motivo = request.form.get("motivo", "").strip()
    nuevo = 0 if e["anulado"] else 1
    otros = [dict(x) for x in db.q("SELECT tipo, subcategoria, anulado FROM estudios WHERE protocolo_id=? AND id<>?",
                                   (e["protocolo_id"], eid))]
    if nuevo and not motivo:
        flash("Indicá el motivo de la anulación.", "error")
    elif not nuevo and (error := validar_combinacion(otros + [{"tipo": e["tipo"], "subcategoria": e["subcategoria"]}])):
        flash(f"No se puede reactivar: {error}", "error")
    elif not nuevo and (avance := flujo_avanzado(e["protocolo_id"], e["tipo"])):
        flash(f"No se puede reactivar: hay {REGISTRO[e['tipo']].nombre.lower()} de este protocolo en proceso (llegó a {avance}) y comparten el mismo "
              "flujo. Primero hay que deshacer las etapas.", "error")
    else:
        db.ex("UPDATE estudios SET anulado=?, motivo_anulacion=? WHERE id=?", (nuevo, motivo if nuevo else None, eid))
        db.consolidar_diagnosticos(db.con(), e["protocolo_id"], e["tipo"], extra=[eid] if nuevo else [])
        db.olvidar_cache()
        db.auditar(usuario_actual()["id"], e["protocolo_id"], "anular" if nuevo else "reactivar", motivo, eid)
        flash("Estudio anulado." if nuevo else "Estudio reactivado.", "ok")
        if not nuevo:
            avisar_lote_lleno(ubicar_en_lote(e["protocolo_id"], usuario_actual()["id"]))
    return redirect(url_for("estudio", eid=eid))


@app.route("/protocolo/<int:pid>/sistema", methods=["POST"])
@requiere_permiso("sistema")
def cargado_sistema(pid):
    protocolo_o_404(pid)
    cargado = 1 if request.form.get("cargado_sistema") == "1" else 0
    db.ex("UPDATE protocolos SET protocolo_sistema=?, cargado_sistema=? WHERE id=?",
          (request.form.get("protocolo_sistema", "").strip(), cargado, pid))
    db.auditar(usuario_actual()["id"], pid, "sistema", "cargado" if cargado else "no cargado")
    flash("Datos de carga en el sistema actualizados.", "ok")
    return redirect(url_for("protocolo", pid=pid))


@app.route("/lotes")
@requiere_login
def lotes():
    fecha_sel = request.args.get("fecha") or hoy()
    filas = db.q("""SELECT l.*, u.iniciales AS creador, COUNT(e.id) AS cantidad
                    FROM lotes l LEFT JOIN usuarios u ON u.id=l.creado_por
                    LEFT JOIN estudios e ON e.lote_id=l.id AND e.anulado=0
                    WHERE l.fecha=? GROUP BY l.id ORDER BY l.cerrado, l.tipo_lote, l.numero""", (fecha_sel,))
    dias = db.q("SELECT fecha, COUNT(*) AS n FROM lotes GROUP BY fecha ORDER BY fecha DESC LIMIT 15")
    return render_template("lotes.html", lotes=filas, fecha_sel=fecha_sel, es_hoy=fecha_sel == hoy(), dias=dias,
                           tacos=tacos_por("lote_id"), lotes_bp=REGISTRO["BP"].lotes, capacidad=REGISTRO["BP"].tacos_por_lote,
                           grupos_lote=[("Biopsias", tipos_lote("BP")), ("Citologías", tipos_lote("CT")),
                                        ("PAP · citotécnico", tipos_lote("PAP"))])


@app.route("/lotes/nuevo", methods=["POST"])
@requiere_permiso("lotes_armar")
def nuevo_lote():
    tipo_lote = request.form.get("tipo_lote", "")
    if tipo_lote not in tipos_lote():
        flash("Elegí el tipo de lote.", "error")
        return redirect(url_for("lotes"))
    lote = crear_lote(tipo_lote, usuario_actual()["id"])
    flash(f"Lote {lote['codigo']} creado.", "ok")
    return redirect(url_for("lote", lote_id=lote["id"]))


def lote_o_404(lote_id):
    l = db.uno("""SELECT l.*, u.iniciales AS creador, u2.iniciales AS cerrador FROM lotes l
                  LEFT JOIN usuarios u ON u.id=l.creado_por LEFT JOIN usuarios u2 ON u2.id=l.cerrado_por
                  WHERE l.id=?""", (lote_id,))
    if not l:
        abort(404)
    return l


def orden_numero(numero):
    m = re.fullmatch(r"(\D*)(\d+)(.*)", numero or "")
    return (m.group(1), int(m.group(2)), m.group(3)) if m else (numero or "", 0, "")


def plan_avance(lote_id):
    return plan_avance_de(cargar_estudios("e.lote_id=?", (lote_id,), borradores=True))


def plan_avance_de(estudios):
    activos = [x for x in estudios if not x["anulado"]]
    if not activos:
        return None, "El lote no tiene estudios para avanzar."
    borradores = sorted({x["numero"] for x in activos if x["borrador"]}, key=orden_numero)
    if borradores:
        return None, "Hay protocolos a completar en el lote (" + ", ".join(borradores[:8]) + (", …" if len(borradores) > 8 else "") + "): completalos antes de avanzar."
    flujos = {}
    for x in activos:
        flujos.setdefault((x["protocolo_id"], x["tipo"]), []).append(x)
    por_etapa = {}
    for f in flujos.values():
        por_etapa.setdefault(f[0]["proxima"][0] if f[0]["proxima"] else None, []).append(f)
    if len(por_etapa) > 1:
        partes = [f"{NOMBRE_ETAPA[k] if k else 'Informado'}: {len(v)}" for k, v in sorted(por_etapa.items(), key=lambda kv: -len(kv[1]))]
        return None, "Las muestras del lote no están todas en la misma etapa (" + " · ".join(partes) + ")."
    etapa = next(iter(por_etapa))
    if etapa is None:
        return None, "Todos los estudios del lote ya están informados."
    return {"etapa": etapa, "flujos": list(flujos.values()), "sector": por_etapa[etapa][0][0]["proxima"][2]}, None


@app.route("/lotes/<int:lote_id>/avanzar", methods=["POST"])
@requiere_permiso("etapas")
def lote_avanzar(lote_id):
    yo = usuario_actual()
    l = lote_o_404(lote_id)
    destino = url_for("lote", lote_id=lote_id)
    plan, error = plan_avance(lote_id)
    if not error and request.form.get("etapa") != plan["etapa"]:
        error = "El lote cambió de etapa mientras tenías la pantalla abierta: revisá y volvé a intentar."
    if not error and not yo["admin"] and plan["sector"] not in (yo["sectores"] or "").split(","):
        error = f"La etapa {NOMBRE_ETAPA[plan['etapa']]} la marca el sector {SECTORES[plan['sector']]}."
    if error:
        flash(error, "error")
        return redirect(destino)
    etapa = plan["etapa"]
    db.precargar([(f"SELECT m.* FROM {t} m JOIN estudios e ON e.id=m.estudio_id WHERE e.lote_id=?", (lote_id,)) for t in ("macro", "micro", "ihq")]
                 + [("SELECT DISTINCT t.estudio_id FROM tecnicas t JOIN estudios e ON e.id=t.estudio_id WHERE e.lote_id=?", (lote_id,))])
    por_tabla = {t: {r["estudio_id"]: r for r in db.q(f"SELECT m.* FROM {t} m JOIN estudios e ON e.id=m.estudio_id WHERE e.lote_id=?", (lote_id,))}
                 for t in ("macro", "micro", "ihq")}
    con_tecnicas = {r["estudio_id"] for r in db.q("SELECT DISTINCT t.estudio_id FROM tecnicas t JOIN estudios e ON e.id=t.estudio_id WHERE e.lote_id=?", (lote_id,))}
    problemas = []
    for flujo in plan["flujos"]:
        cabecera = min(m["id"] for m in flujo)
        primero = flujo[0]
        sin = [m for m in flujo if m["id"] not in con_tecnicas] if etapa == "macroscopia" and "macro" in REGISTRO[primero["tipo"]].secciones else []
        motivo = motivo_no_listo(yo, primero["tipo"], etapa, primero["proxima"], Variante(primero["solicita_ihq"], primero["decal"]),
                                 por_tabla["macro"].get(cabecera), por_tabla["micro"].get(cabecera), por_tabla["ihq"].get(cabecera), sin)
        if motivo:
            problemas.append(f"{primero['numero']}: {motivo}")
    if problemas:
        flash(f"No se avanzó el lote ({NOMBRE_ETAPA[etapa]}): " + " · ".join(problemas[:6]) + (f" · y {len(problemas) - 6} más" if len(problemas) > 6 else ""), "error")
        return redirect(destino)
    ahora_ = db.ahora()
    estudios = [m for f in plan["flujos"] for m in f]
    sentencias = []
    for i in range(0, len(estudios), 25):
        grupo = estudios[i:i + 25]
        sentencias.append(("INSERT INTO etapas (estudio_id, etapa, usuario_id, fecha_hora) VALUES " + ",".join(["(?,?,?,?)"] * len(grupo)),
                           tuple(v for m in grupo for v in (m["id"], etapa, yo["id"], ahora_))))
    for i in range(0, len(plan["flujos"]), 16):
        grupo = plan["flujos"][i:i + 16]
        sentencias.append(("INSERT INTO auditoria (fecha_hora, usuario_id, protocolo_id, estudio_id, accion, detalle) VALUES " + ",".join(["(?,?,?,?,?,?)"] * len(grupo)),
                           tuple(v for f in grupo for v in (ahora_, yo["id"], f[0]["protocolo_id"], f[0]["id"], "listo",
                                                            f"{NOMBRE_ETAPA[etapa]} (lote {l['codigo']})"))))
    try:
        db.lote_ex(sentencias)
    except db.IntegrityError:
        flash("No se avanzó el lote: alguna etapa se registró en el medio. Recargá la pantalla y volvé a intentar.", "error")
        return redirect(destino)
    flash(f"✔ {NOMBRE_ETAPA[etapa]} registrada por {yo['iniciales']} en {len(plan['flujos'])} protocolo(s) / {len(estudios)} estudio(s) del lote {l['codigo']}.", "ok")
    return redirect(destino)


@app.route("/lotes/<int:lote_id>")
@requiere_login
def lote(lote_id):
    l = lote_o_404(lote_id)
    del_lote = cargar_estudios("e.lote_id=?", (lote_id,), borradores=True)
    del_lote.sort(key=lambda x: (orden_numero(x["numero"]), x["id"]))
    primeros = {r["id"] for r in db.q("SELECT MIN(id) AS id FROM estudios WHERE anulado=0 GROUP BY protocolo_id")}
    sin_lote = [x for x in cargar_estudios("e.lote_id IS NULL AND e.anulado=0")
                if x["id"] in primeros and l["tipo_lote"] in tipos_lote(x["tipo"])][:300]
    en_bp = l["tipo_lote"] in REGISTRO["BP"].lotes
    tacos_total, sin_dato = tacos_de_lote(lote_id) if en_bp else (0, 0)
    avance, sin_avance = plan_avance_de(del_lote)
    return render_template("lote.html", avance=avance, sin_avance=sin_avance, l=l, estudios=del_lote, sin_lote=sin_lote, capacidad=REGISTRO["BP"].tacos_por_lote if en_bp else None,
                           tacos_total=tacos_total, sin_dato=sin_dato,
                           etiquetas_lab=(l["tipo_lote"] in tipos_lote("BP") or l["tipo_lote"] in tipos_lote("PAP"))
                           and any(not x["anulado"] and not x["borrador"] for x in del_lote))


@app.route("/lotes/<int:lote_id>/agregar", methods=["POST"])
@requiere_permiso("lotes_armar")
def lote_agregar(lote_id):
    yo = usuario_actual()
    l = lote_o_404(lote_id)
    elegido = request.form.get("estudio_id", "")
    numero = request.form.get("protocolo", "").strip()
    if elegido.isdigit():
        p = db.uno("SELECT p.* FROM protocolos p JOIN estudios e ON e.protocolo_id=p.id WHERE e.id=?", (int(elegido),))
    else:
        p = db.uno("SELECT * FROM protocolos WHERE numero=?", (numero,)) if numero else None
    primero = primer_estudio(p["id"]) if p else None
    if l["cerrado"]:
        flash("El lote está cerrado.", "error")
    elif not elegido and not numero:
        flash("Elegí un protocolo de la lista o escribí su N°.", "error")
    elif not p:
        flash(f"No existe el protocolo {numero}. Primero hay que ingresarlo.", "error")
    elif not primero:
        flash(f"El protocolo {p['numero']} no tiene estudios activos.", "error")
    elif l["tipo_lote"] not in tipos_lote(primero["tipo"]):
        flash(f"El protocolo {p['numero']} empieza con {etiqueta(primero)} y el lote lo marca su primer flujo: no va en un lote {l['tipo_lote']}.", "error")
    elif primero["lote_id"] == lote_id:
        flash(f"{p['numero']} ya está en este lote.", "error")
    else:
        previo = lote_del_protocolo(p["id"])
        asignar_lote(p["id"], l, yo["id"])
        flash(f"{p['numero']} agregado al lote" + (f" (estaba en {previo['codigo']})." if previo else "."), "ok")
        capacidad = REGISTRO["BP"].tacos_por_lote
        total = tacos_de_lote(lote_id)[0] if l["tipo_lote"] in REGISTRO["BP"].lotes else 0
        if total > capacidad:
            flash(f"Atención: el lote {l['codigo']} queda en {total}/{capacidad} tacos.", "error")
    return redirect(url_for("lote", lote_id=lote_id))


@app.route("/lotes/<int:lote_id>/quitar/<int:eid>", methods=["POST"])
@requiere_permiso("lotes_armar")
def lote_quitar(lote_id, eid):
    l = lote_o_404(lote_id)
    e = db.uno("SELECT protocolo_id FROM estudios WHERE id=? AND lote_id=?", (eid, lote_id))
    if l["cerrado"]:
        flash("El lote está cerrado.", "error")
    elif e:
        asignar_lote(e["protocolo_id"], None, usuario_actual()["id"])
        flash("Protocolo quitado del lote.", "ok")
    return redirect(url_for("lote", lote_id=lote_id))


@app.route("/lotes/<int:lote_id>/cerrar", methods=["POST"])
@requiere_permiso("lotes_cerrar")
def lote_cerrar(lote_id):
    yo = usuario_actual()
    l = lote_o_404(lote_id)
    if l["cerrado"]:
        db.ex("UPDATE lotes SET cerrado=0, cerrado_por=NULL, cerrado_en=NULL WHERE id=?", (lote_id,))
        db.auditar(yo["id"], None, "reabrir_lote", l["codigo"])
        flash(f"Lote {l['codigo']} reabierto.", "ok")
    else:
        db.ex("UPDATE lotes SET cerrado=1, cerrado_por=?, cerrado_en=? WHERE id=?", (yo["id"], db.ahora(), lote_id))
        db.auditar(yo["id"], None, "cerrar_lote", l["codigo"])
        flash(f"Lote {l['codigo']} cerrado.", "ok")
    return redirect(url_for("lote", lote_id=lote_id))


@app.route("/lotes/<int:lote_id>/observaciones", methods=["POST"])
@requiere_permiso("lotes_armar")
def lote_observaciones(lote_id):
    lote_o_404(lote_id)
    db.ex("UPDATE lotes SET observaciones=? WHERE id=?", (request.form.get("observaciones", "").strip(), lote_id))
    flash("Observaciones guardadas.", "ok")
    return redirect(url_for("lote", lote_id=lote_id))


@app.route("/lotes/<int:lote_id>/eliminar", methods=["POST"])
@requiere_permiso("lotes_cerrar")
def lote_eliminar(lote_id):
    l = lote_o_404(lote_id)
    if db.uno("SELECT 1 FROM estudios WHERE lote_id=?", (lote_id,)):
        flash("Solo se puede eliminar un lote vacío.", "error")
        return redirect(url_for("lote", lote_id=lote_id))
    db.ex("DELETE FROM lotes WHERE id=?", (lote_id,))
    db.auditar(usuario_actual()["id"], None, "eliminar_lote", l["codigo"])
    flash(f"Lote {l['codigo']} eliminado.", "ok")
    return redirect(url_for("lotes", fecha=l["fecha"]))


@app.route("/api/catalogo/<tipo>")
@requiere_login
def api_catalogo(tipo):
    arbol = {}
    for r in db.q("SELECT subcategoria, sitio, tipo_muestra FROM catalogo WHERE tipo=?", (tipo,)):
        arbol.setdefault(r["subcategoria"], {}).setdefault(r["sitio"], []).append(r["tipo_muestra"])
    return jsonify(arbol)


def sin_acentos(texto):
    t = unicodedata.normalize("NFD", texto or "")
    return " ".join("".join(c for c in t if unicodedata.category(c) != "Mn").lower().split())


def tipos_de_muestra_bp():
    cuentas = {}
    for r in db.q("SELECT tipo_muestra, COUNT(*) AS n FROM catalogo WHERE tipo='BP' AND tipo_muestra IS NOT NULL AND tipo_muestra<>'' GROUP BY tipo_muestra"):
        cuentas.setdefault(sin_acentos(r["tipo_muestra"]), []).append((r["n"], r["tipo_muestra"]))
    return sorted(max(v)[1] for v in cuentas.values())


@app.route("/templates", methods=["GET", "POST"])
@requiere_admin
def templates_admin():
    yo = usuario_actual()
    catalogo = db.q("SELECT DISTINCT subcategoria, sitio FROM catalogo WHERE tipo='BP' ORDER BY subcategoria, sitio")
    zonas = sorted({r["subcategoria"] for r in catalogo})
    tipos = tipos_de_muestra_bp()
    if request.method == "POST":
        ids = [int(x) for x in request.form.getlist("ids") if x.isdigit()]
        cambios = {}
        for campo, columna in (("zona", "tipo_biopsia"), ("organo", "organo"), ("tipo", "tipo_muestra")):
            valor = request.form.get(campo, "")
            if valor == "__vaciar__":
                cambios[columna] = None
            elif valor:
                cambios[columna] = valor
        organo = cambios.get("organo")
        if organo and "tipo_biopsia" not in cambios:
            propias = {r["subcategoria"] for r in catalogo if r["sitio"] == organo}
            if len(propias) == 1:
                cambios["tipo_biopsia"] = next(iter(propias))
        if not ids:
            flash("Marcá al menos un template.", "error")
        elif not cambios:
            flash("Elegí qué querés asignar (zona, órgano o tipo).", "error")
        else:
            sets = ", ".join(f"{k}=?" for k in cambios)
            for i in range(0, len(ids), 40):
                db.lote_ex([(f"UPDATE templates SET {sets} WHERE id=?", tuple(cambios.values()) + (tid,)) for tid in ids[i:i + 40]])
            db.auditar(yo["id"], None, "templates", f"{len(ids)} templates: " + ", ".join(f"{k}={v or '(vacío)'}" for k, v in cambios.items()))
            flash(f"Se actualizaron {len(ids)} templates.", "ok")
        return redirect(url_for("templates_admin", **{k: v for k, v in request.args.items()}))
    f = {k: request.args.get(k, "") for k in ("clase", "zona", "organo", "tipo", "estado", "q")}
    where, par = ["1=1"], []
    if f["clase"] in ("macro", "micro"):
        where.append("clase=?")
        par.append(f["clase"])
    if f["zona"]:
        where.append("tipo_biopsia=?")
        par.append(f["zona"])
    if f["organo"]:
        where.append("organo=?")
        par.append(f["organo"])
    if f["tipo"]:
        where.append("tipo_muestra=?")
        par.append(f["tipo"])
    if f["estado"] == "sin_organo":
        where.append("(organo IS NULL OR organo='')")
    elif f["estado"] == "sin_tipo":
        where.append("(tipo_muestra IS NULL OR tipo_muestra='')")
    elif f["estado"] == "sin_zona":
        where.append("(tipo_biopsia IS NULL OR tipo_biopsia='')")
    if f["q"]:
        where.append("LOWER(titulo) LIKE ?")
        par.append("%" + f["q"].lower() + "%")
    total = db.uno("SELECT COUNT(*) AS n FROM templates")["n"]
    filas = db.q(f"SELECT id, clase, titulo, tipo_biopsia, organo, tipo_muestra FROM templates WHERE {' AND '.join(where)} ORDER BY clase, titulo LIMIT 400", par)
    sin_clasificar = db.uno("SELECT COUNT(*) AS n FROM templates WHERE organo IS NULL OR organo=''")["n"]
    return render_template("templates_admin.html", filas=filas, f=f, zonas=zonas, tipos=tipos, total=total, sin_clasificar=sin_clasificar,
                           organos=[(z, [r["sitio"] for r in catalogo if r["subcategoria"] == z]) for z in zonas])


@app.route("/api/template/<int:tid>")
@requiere_login
def api_template(tid):
    t = db.uno("SELECT titulo, texto, conclusion FROM templates WHERE id=?", (tid,))
    return jsonify(dict(t)) if t else (jsonify({}), 404)


@app.route("/api/medicos")
@requiere_login
def api_medicos():
    r = jsonify([m["nombre"] for m in db.q("SELECT nombre FROM medicos ORDER BY nombre")])
    r.headers["Cache-Control"] = "private, max-age=3600"
    return r


@app.route("/api/ubicaciones")
@requiere_login
def api_ubicaciones():
    r = jsonify([{"v": u["abreviatura"], "n": u["nombre"] or ""} for u in db.q("SELECT abreviatura, nombre FROM ubicaciones ORDER BY abreviatura")])
    r.headers["Cache-Control"] = "private, max-age=60"
    return r


CAMPOS_LUGAR = (("lugar_recoleccion", "Lugar recolección"), ("lugar_entrega", "Lugar entrega"))
SQL_UBICACIONES = """SELECT u.*, (SELECT COUNT(*) FROM ubicacion_medicos m WHERE m.ubicacion_id=u.id) AS n_medicos
                     FROM ubicaciones u ORDER BY u.abreviatura"""


def sql_protocolos_por_lugar(campo):
    return f"SELECT {campo} AS lugar, COUNT(*) AS n FROM protocolos WHERE borrador=0 AND {campo} IS NOT NULL GROUP BY {campo}"


def error_ubicaciones(d, actual=None):
    for campo, rotulo in CAMPOS_LUGAR:
        v = d.get(campo)
        if not v or (actual is not None and v == (actual[campo] or "")):
            continue
        if not db.uno("SELECT 1 FROM ubicaciones WHERE abreviatura=?", (v,)):
            return f"{rotulo}: elegí una ubicación de la lista. Si falta, se crea en Ubicaciones."
    return None


def ubicacion_o_404(uid):
    return db.uno("SELECT * FROM ubicaciones WHERE id=?", (uid,)) or abort(404)


@app.route("/ubicaciones")
@requiere_login
def ubicaciones_lista():
    q = request.args.get("q", "").strip()
    db.precargar([(SQL_UBICACIONES, ()), (sql_protocolos_por_lugar("lugar_recoleccion"), ()), (sql_protocolos_por_lugar("lugar_entrega"), ())])
    filas = db.q(SQL_UBICACIONES)
    if q:
        palabras = db.clave_nombre(q).split()
        por_medico = {r["ubicacion_id"] for r in db.q(
            "SELECT DISTINCT m.ubicacion_id FROM ubicacion_medicos m JOIN medicos d ON d.id=m.medico_id WHERE "
            + " AND ".join(["d.nombre LIKE ?"] * len(palabras)), tuple(f"%{p}%" for p in palabras))} if palabras else set()
        filas = [u for u in filas if u["id"] in por_medico or all(p in db.clave_nombre(f"{u['abreviatura']} {u['nombre'] or ''} {u['direccion'] or ''} {u['contactos'] or ''}")
                                                                  for p in palabras)]
    destinos = destinos_por_ubicacion()
    filas = [dict(u, emails=", ".join(emails_de_informes(destinos.get(u["id"], [])))) for u in filas]
    return render_template("ubicaciones.html", ubicaciones=filas, q=q,
                           recoleccion={r["lugar"]: r["n"] for r in db.q(sql_protocolos_por_lugar("lugar_recoleccion"))},
                           entrega={r["lugar"]: r["n"] for r in db.q(sql_protocolos_por_lugar("lugar_entrega"))})


def leer_ubicacion():
    return {k: request.form.get(k, "").strip() for k in ("abreviatura", "nombre", "direccion", "contactos")}


def error_ubicacion(d, uid=None):
    if not d["abreviatura"]:
        return "Falta la abreviatura de la ubicación."
    if db.uno("SELECT 1 FROM ubicaciones WHERE abreviatura=? AND id<>?", (d["abreviatura"], uid or 0)):
        return f"Ya existe una ubicación llamada «{d['abreviatura']}»."
    return None


@app.route("/ubicaciones/nueva", methods=["GET", "POST"])
@requiere_permiso("ubicaciones")
def ubicacion_nueva():
    d = {}
    if request.method == "POST":
        d = leer_ubicacion()
        error = error_ubicacion(d)
        if error:
            flash(error, "error")
        else:
            uid = db.ex("INSERT INTO ubicaciones (abreviatura, nombre, direccion, contactos) VALUES (?,?,?,?)",
                        (d["abreviatura"], d["nombre"], d["direccion"], d["contactos"]))
            db.auditar(usuario_actual()["id"], None, "ubicacion", f"Alta: {d['abreviatura']}")
            flash("Ubicación creada. Ahora podés cargar a quién se envían los informes y asociarle médicos.", "ok")
            return redirect(url_for("ubicacion", uid=uid))
    return render_template("ubicacion.html", u=d, nueva=True, medicos=[], destinos=[], tipos=envios.TIPOS_ESTUDIO, malignidades=envios.MALIGNIDADES, whatsapp_visible=envios.whatsapp_visible, alcance_texto=envios.alcance_texto, recoleccion=0, entrega=0)


@app.route("/ubicaciones/<int:uid>")
@requiere_login
def ubicacion(uid):
    u = ubicacion_o_404(uid)
    medicos = db.q("SELECT d.id, d.nombre FROM ubicacion_medicos m JOIN medicos d ON d.id=m.medico_id WHERE m.ubicacion_id=? ORDER BY d.nombre", (uid,))
    cuenta = lambda campo: db.uno(f"SELECT COUNT(*) AS n FROM protocolos WHERE borrador=0 AND {campo}=?", (u["abreviatura"],))["n"]
    destinos = db.q("SELECT id, canal, destino, informes, alcance FROM ubicacion_destinos WHERE ubicacion_id=? ORDER BY canal, destino", (uid,))
    return render_template("ubicacion.html", u=u, nueva=False, medicos=medicos, destinos=destinos, tipos=envios.TIPOS_ESTUDIO, malignidades=envios.MALIGNIDADES, whatsapp_visible=envios.whatsapp_visible, alcance_texto=envios.alcance_texto,
                           recoleccion=cuenta("lugar_recoleccion"), entrega=cuenta("lugar_entrega"))


@app.route("/ubicaciones/<int:uid>/editar", methods=["POST"])
@requiere_permiso("ubicaciones")
def ubicacion_editar(uid):
    u = ubicacion_o_404(uid)
    d = leer_ubicacion()
    error = error_ubicacion(d, uid)
    if error:
        flash(error, "error")
        return redirect(url_for("ubicacion", uid=uid))
    sentencias = [("UPDATE ubicaciones SET abreviatura=?, nombre=?, direccion=?, contactos=? WHERE id=?", (d["abreviatura"], d["nombre"], d["direccion"], d["contactos"], uid))]
    if d["abreviatura"] != u["abreviatura"]:
        sentencias += [(f"UPDATE protocolos SET {campo}=? WHERE {campo}=?", (d["abreviatura"], u["abreviatura"])) for campo, _ in CAMPOS_LUGAR]
    db.lote_ex(sentencias)
    detalle = f"{u['abreviatura']} → {d['abreviatura']}" if d["abreviatura"] != u["abreviatura"] else d["abreviatura"]
    db.auditar(usuario_actual()["id"], None, "ubicacion", f"Edición: {detalle}")
    flash("Ubicación actualizada." + (" Se actualizó la abreviatura en sus protocolos." if d["abreviatura"] != u["abreviatura"] else ""), "ok")
    return redirect(url_for("ubicacion", uid=uid))


@app.route("/ubicaciones/<int:uid>/eliminar", methods=["POST"])
@requiere_permiso("ubicaciones")
def ubicacion_eliminar(uid):
    u = ubicacion_o_404(uid)
    if db.uno("SELECT 1 FROM protocolos WHERE lugar_recoleccion=? OR lugar_entrega=? LIMIT 1", (u["abreviatura"], u["abreviatura"])):
        flash("No se puede eliminar: hay protocolos con esta ubicación.", "error")
        return redirect(url_for("ubicacion", uid=uid))
    db.lote_ex([("DELETE FROM ubicacion_medicos WHERE ubicacion_id=?", (uid,)), ("DELETE FROM ubicacion_destinos WHERE ubicacion_id=?", (uid,)), ("DELETE FROM ubicaciones WHERE id=?", (uid,))])
    db.auditar(usuario_actual()["id"], None, "ubicacion", f"Baja: {u['abreviatura']}")
    flash(f"Ubicación «{u['abreviatura']}» eliminada.", "ok")
    return redirect(url_for("ubicaciones_lista"))


def alcance_del_formulario():
    return envios.alcance_valido(request.form.getlist("alcance"))


@app.route("/ubicaciones/<int:uid>/destino", methods=["POST"])
@requiere_permiso("ubicaciones")
def ubicacion_destino_agregar(uid):
    u = ubicacion_o_404(uid)
    volver = redirect(url_for("ubicacion", uid=uid) + "#envio")
    canal = "whatsapp" if request.form.get("canal") == "whatsapp" else "email"
    bruto = request.form.get("destino", "").strip()
    if canal == "email":
        destino = bruto.lower()
        if not envios.email_valido(destino):
            flash("El e-mail no es válido.", "error")
            return volver
    else:
        destino = envios.whatsapp_numero(bruto)
        if not destino:
            flash("El número no es válido: escribilo con el código de área, por ejemplo 11 5015 1099.", "error")
            return volver
    if db.uno("SELECT 1 FROM ubicacion_destinos WHERE ubicacion_id=? AND canal=? AND destino=?", (uid, canal, destino)):
        flash("Ese destino ya está cargado en la ubicación.", "error")
        return volver
    informes = 1 if request.form.get("informes") == "1" else 0
    db.ex("INSERT INTO ubicacion_destinos (ubicacion_id, canal, destino, informes, alcance) VALUES (?,?,?,?,?)", (uid, canal, destino, informes, alcance_del_formulario()))
    db.auditar(usuario_actual()["id"], None, "ubicacion", f"{u['abreviatura']}: se agregó {destino} ({'WhatsApp' if canal == 'whatsapp' else 'e-mail'})")
    flash("Destino agregado.", "ok")
    return volver


@app.route("/ubicaciones/<int:uid>/destino/<int:did>", methods=["POST"])
@requiere_permiso("ubicaciones")
def ubicacion_destino_editar(uid, did):
    u = ubicacion_o_404(uid)
    d = db.uno("SELECT * FROM ubicacion_destinos WHERE id=? AND ubicacion_id=?", (did, uid)) or abort(404)
    if request.form.get("accion") == "borrar":
        db.ex("DELETE FROM ubicacion_destinos WHERE id=?", (did,))
        db.auditar(usuario_actual()["id"], None, "ubicacion", f"{u['abreviatura']}: se quitó {d['destino']}")
        flash("Destino quitado.", "ok")
    else:
        informes = 1 if request.form.get("informes") == "1" else 0
        alcance = alcance_del_formulario()
        db.ex("UPDATE ubicacion_destinos SET informes=?, alcance=? WHERE id=?", (informes, alcance, did))
        db.auditar(usuario_actual()["id"], None, "ubicacion",
                   f"{u['abreviatura']}: {d['destino']} → {'recibe informes' if informes else 'no recibe informes'} · {envios.alcance_texto(alcance)}")
        flash("Destino actualizado.", "ok")
    return redirect(url_for("ubicacion", uid=uid) + "#envio")


@app.route("/ubicaciones/<int:uid>/medico", methods=["POST"])
@requiere_permiso("ubicaciones")
def ubicacion_medico_agregar(uid):
    u = ubicacion_o_404(uid)
    nombre = request.form.get("medico", "").strip()
    clave = db.clave_nombre(nombre)
    hallados = [m for m in db.q("SELECT id, nombre FROM medicos WHERE nombre=?", (nombre,))] if nombre else []
    if not hallados and clave:
        hallados = [m for m in db.q("SELECT id, nombre FROM medicos ORDER BY id") if db.clave_nombre(m["nombre"]) == clave]
    if not hallados:
        flash("No encontré ese médico: elegilo de la lista que se despliega al escribir.", "error")
    elif db.uno("SELECT 1 FROM ubicacion_medicos WHERE ubicacion_id=? AND medico_id=?", (uid, hallados[0]["id"])):
        flash(f"{hallados[0]['nombre']} ya está asociado a esta ubicación.", "error")
    else:
        db.ex("INSERT INTO ubicacion_medicos (ubicacion_id, medico_id) VALUES (?,?)", (uid, hallados[0]["id"]))
        db.auditar(usuario_actual()["id"], None, "ubicacion", f"{u['abreviatura']}: se asoció a {hallados[0]['nombre']}")
        flash(f"{hallados[0]['nombre']} asociado a la ubicación.", "ok")
    return redirect(url_for("ubicacion", uid=uid) + "#medicos")


@app.route("/ubicaciones/<int:uid>/medicos/quitar", methods=["POST"])
@requiere_permiso("ubicaciones")
def ubicacion_medico_quitar(uid):
    u = ubicacion_o_404(uid)
    ids = [int(i) for i in request.form.getlist("medico") if i.isdigit()]
    if not ids:
        flash("Marcá los médicos que querés quitar.", "error")
    else:
        db.lote_ex([("DELETE FROM ubicacion_medicos WHERE ubicacion_id=? AND medico_id=?", (uid, i)) for i in ids])
        db.auditar(usuario_actual()["id"], None, "ubicacion", f"{u['abreviatura']}: se quitaron {len(ids)} médico(s)")
        flash(f"Se quitaron {len(ids)} médico(s) de la ubicación.", "ok")
    return redirect(url_for("ubicacion", uid=uid) + "#medicos")


@app.route("/ubicaciones/importar", methods=["POST"])
@requiere_permiso("ubicaciones")
def ubicaciones_importar():
    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        flash("Elegí el archivo CSV de ubicaciones.", "error")
        return redirect(url_for("ubicaciones_lista"))
    try:
        filas = ubicaciones_csv.leer_csv(archivo.read())
    except ubicaciones_csv.ErrorArchivo as e:
        flash(str(e), "error")
        return redirect(url_for("ubicaciones_lista"))
    if not filas:
        flash("El archivo no tiene ubicaciones.", "error")
        return redirect(url_for("ubicaciones_lista"))
    r = db.ubicaciones_importar(filas)
    db.auditar(usuario_actual()["id"], None, "ubicaciones_importar", f"{r['nuevas']} nuevas, {r['actualizadas']} actualizadas, {r['vinculos']} médicos asociados")
    flash(f"Importación lista: {r['nuevas']} ubicaciones nuevas, {r['actualizadas']} actualizadas, {r['vinculos']} médicos asociados"
          + (f", {r['medicos_nuevos']} médicos agregados a la lista." if r["medicos_nuevos"] else "."), "ok")
    return redirect(url_for("ubicaciones_lista"))


@app.route("/usuarios")
@requiere_admin
def usuarios():
    filas = db.q("""SELECT u.*, p.nombre AS perfil FROM usuarios u LEFT JOIN perfiles p ON p.id=u.perfil_id
                    ORDER BY u.activo DESC, u.iniciales""")
    return render_template("usuarios.html", usuarios=filas)


@app.route("/usuarios/nuevo", methods=["GET", "POST"])
@app.route("/usuarios/<int:uid>", methods=["GET", "POST"])
@requiere_admin
def usuario(uid=None):
    yo = usuario_actual()
    u = db.uno("SELECT * FROM usuarios WHERE id=?", (uid,)) if uid else None
    if request.method == "POST":
        ini = request.form["iniciales"].strip().upper()
        sectores = ",".join(s for s in SECTORES if request.form.get("s_" + s))
        perfil_id = int(request.form["perfil_id"]) if request.form.get("perfil_id", "").isdigit() else None
        ajuste = {k: request.form.get("p_" + k, "") for k in permisos.PERMISOS}
        titulo = request.form.get("titulo", "") if request.form.get("titulo") in ("Dr.", "Dra.") else ""
        mn, mp = request.form.get("mn", "").strip()[:20], request.form.get("mp", "").strip()[:20]
        firma_png, firma_error = None, None
        archivo = request.files.get("firma")
        if archivo and archivo.filename and informes is None:
            firma_error = SIN_INFORMES
        elif archivo and archivo.filename:
            try:
                firma_png = informes.procesar_firma(archivo.read())
            except ValueError as err:
                firma_error = str(err)
        vals = (ini, request.form.get("nombre", "").strip() or ini, sectores,
                1 if request.form.get("admin") else 0, 1 if request.form.get("activo") else 0, perfil_id,
                permisos.texto({k for k, v in ajuste.items() if v == "mas"}),
                permisos.texto({k for k, v in ajuste.items() if v == "menos"}))
        dup = db.uno("SELECT id FROM usuarios WHERE iniciales=? AND id<>?", (ini, uid or 0))
        clave = request.form.get("clave", "")
        if not ini or dup:
            flash("Iniciales vacías o ya usadas por otra persona.", "error")
        elif clave and len(clave) < 8:
            flash("La clave temporal debe tener al menos 8 caracteres.", "error")
        elif perfil_id and not db.uno("SELECT 1 FROM perfiles WHERE id=?", (perfil_id,)):
            flash("El perfil elegido no existe.", "error")
        elif uid == yo["id"] and not (vals[3] and vals[4]):
            flash("No podés quitarte el rol de administrador ni desactivarte a vos mismo.", "error")
        elif firma_error:
            flash(f"Firma: {firma_error}", "error")
        else:
            if u:
                db.ex("UPDATE usuarios SET iniciales=?, nombre=?, sectores=?, admin=?, activo=?, perfil_id=?, permisos_mas=?, "
                      "permisos_menos=? WHERE id=?", vals + (uid,))
            else:
                uid = db.ex("INSERT INTO usuarios (iniciales, nombre, sectores, admin, activo, perfil_id, permisos_mas, "
                            "permisos_menos) VALUES (?,?,?,?,?,?,?,?)", vals)
            if clave:
                db.ex("UPDATE usuarios SET clave_hash=?, debe_cambiar_clave=1 WHERE id=?", (generate_password_hash(clave), uid))
            db.ex("UPDATE usuarios SET titulo=?, mn=?, mp=? WHERE id=?", (titulo, mn, mp, uid))
            if firma_png:
                db.ex("DELETE FROM firmas WHERE usuario_id=?", (uid,))
                db.ex("INSERT INTO firmas (usuario_id, imagen, actualizada_en) VALUES (?,?,?)", (uid, firma_png, db.ahora()))
                db.auditar(yo["id"], None, "firma", f"{ini}: firma cargada")
            elif request.form.get("quitar_firma"):
                db.ex("DELETE FROM firmas WHERE usuario_id=?", (uid,))
                db.auditar(yo["id"], None, "firma", f"{ini}: firma quitada")
            detalle = (f"perfil {db.uno('SELECT nombre FROM perfiles WHERE id=?', (perfil_id,))['nombre'] if perfil_id else '—'}"
                       + (f", +{vals[6]}" if vals[6] else "") + (f", -{vals[7]}" if vals[7] else ""))
            db.auditar(yo["id"], None, "usuario", f"{ini}: {detalle}{' (clave temporal)' if clave else ''}")
            flash(f"Usuario {ini} guardado." + (" Al ingresar va a tener que cambiar la clave." if clave else ""), "ok")
            return redirect(url_for("usuarios"))
        u = {**(dict(u) if u else {}), "iniciales": ini, "nombre": vals[1], "sectores": sectores, "admin": vals[3],
             "activo": vals[4], "perfil_id": perfil_id, "permisos_mas": vals[6], "permisos_menos": vals[7],
             "clave_hash": u["clave_hash"] if u else None, "titulo": titulo, "mn": mn, "mp": mp}
    lista_perfiles = perfiles_con_permisos()
    tiene_firma = bool(uid and db.uno("SELECT 1 FROM firmas WHERE usuario_id=?", (uid,)))
    return render_template("usuario.html", u=u, perfiles=lista_perfiles, tiene_firma=tiene_firma, PERMISOS=permisos.PERMISOS,
                           GRUPOS=permisos.GRUPOS, lista_permisos=permisos.lista,
                           mapa_perfiles={p["id"]: sorted(p["set"]) for p in lista_perfiles})


def perfiles_con_permisos():
    return [{**dict(f), "set": permisos.lista(f["permisos"])} for f in db.q(
        """SELECT p.*, (SELECT COUNT(*) FROM usuarios u WHERE u.perfil_id=p.id AND u.activo=1) AS usuarios
           FROM perfiles p ORDER BY p.nombre""")]


@app.route("/perfiles")
@requiere_admin
def perfiles():
    return render_template("perfiles.html", perfiles=perfiles_con_permisos(), PERMISOS=permisos.PERMISOS,
                           GRUPOS=permisos.GRUPOS)


@app.route("/perfiles/nuevo", methods=["GET", "POST"])
@app.route("/perfiles/<int:pid>", methods=["GET", "POST"])
@requiere_admin
def perfil(pid=None):
    yo = usuario_actual()
    p = db.uno("SELECT * FROM perfiles WHERE id=?", (pid,)) if pid else None
    if pid and not p:
        abort(404)
    if request.method == "POST":
        if request.form.get("eliminar"):
            if db.uno("SELECT 1 FROM usuarios WHERE perfil_id=?", (pid,)):
                flash("No se puede eliminar: hay usuarios con este perfil. Cambiales el perfil primero.", "error")
                return redirect(url_for("perfil", pid=pid))
            db.ex("DELETE FROM perfiles WHERE id=?", (pid,))
            db.auditar(yo["id"], None, "eliminar_perfil", p["nombre"])
            flash(f"Perfil {p['nombre']} eliminado.", "ok")
            return redirect(url_for("perfiles"))
        nombre = request.form.get("nombre", "").strip()
        elegidos = permisos.texto({k for k in permisos.PERMISOS if request.form.get("p_" + k)})
        if not nombre or db.uno("SELECT 1 FROM perfiles WHERE nombre=? AND id<>?", (nombre, pid or 0)):
            flash("El nombre está vacío o ya lo usa otro perfil.", "error")
        else:
            if p:
                db.ex("UPDATE perfiles SET nombre=?, permisos=? WHERE id=?", (nombre, elegidos, pid))
            else:
                pid = db.ex("INSERT INTO perfiles (nombre, permisos) VALUES (?,?)", (nombre, elegidos))
            db.auditar(yo["id"], None, "perfil", f"{nombre}: {elegidos or 'sin permisos'}")
            flash(f"Perfil {nombre} guardado.", "ok")
            return redirect(url_for("perfiles"))
        p = {**(dict(p) if p else {}), "nombre": nombre, "permisos": elegidos}
    usuarios_del = db.q("SELECT iniciales, nombre, activo FROM usuarios WHERE perfil_id=? ORDER BY iniciales", (pid,)) if pid else []
    return render_template("perfil.html", p=p, marcados=permisos.lista(p["permisos"] if p else ""), usuarios=usuarios_del,
                           PERMISOS=permisos.PERMISOS, GRUPOS=permisos.GRUPOS)


@app.route("/feriados", methods=["GET", "POST"])
@requiere_permiso("feriados")
def feriados():
    yo = usuario_actual()
    if request.method == "POST":
        f, desc = request.form.get("fecha", ""), request.form.get("descripcion", "").strip()
        if request.form.get("borrar"):
            db.ex("DELETE FROM feriados WHERE fecha=?", (request.form["borrar"],))
            db.auditar(yo["id"], None, "feriado_borrado", request.form["borrar"])
        elif len(f) == 10:
            db.ex("DELETE FROM feriados WHERE fecha=?", (f,))
            db.ex("INSERT INTO feriados VALUES (?,?)", (f, desc))
            db.auditar(yo["id"], None, "feriado", f"{f} {desc}")
            flash("Feriado guardado.", "ok")
        return redirect(url_for("feriados"))
    return render_template("feriados.html", feriados=db.q("SELECT * FROM feriados ORDER BY fecha"))


@app.template_filter("peso")
def peso(v):
    try:
        n = float(v or 0)
    except (TypeError, ValueError):
        return ""
    texto = f"{abs(n):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if n < 0 else "") + "$ " + texto


@app.template_filter("peso0")
def peso_entero(v):
    try:
        n = round(float(v or 0))
    except (TypeError, ValueError):
        return ""
    return ("-" if n < 0 else "") + "$ " + f"{abs(n):,}".replace(",", ".")


def coberturas_lista():
    return listas().get("cobertura", [])


def factura_o_404(fid):
    return db.uno("""SELECT f.*, uc.iniciales AS creador, ue.iniciales AS emisor, ub.iniciales AS cobrador, ua.iniciales AS anulador
                     FROM facturas f LEFT JOIN usuarios uc ON uc.id=f.creado_por LEFT JOIN usuarios ue ON ue.id=f.emitida_por
                     LEFT JOIN usuarios ub ON ub.id=f.cobrada_por LEFT JOIN usuarios ua ON ua.id=f.anulada_por WHERE f.id=?""", (fid,)) or abort(404)


def items_de_factura(fid):
    return db.q("SELECT * FROM factura_items WHERE factura_id=? ORDER BY numero_protocolo, id", (fid,))


@app.route("/facturacion")
@requiere_permiso("facturacion")
def facturacion_lista():
    f = {k: request.args.get(k, "") for k in ("cobertura", "anio", "mes", "estado")}
    todas = db.q("SELECT f.*, u.iniciales AS creador FROM facturas f LEFT JOIN usuarios u ON u.id=f.creado_por ORDER BY f.id DESC")
    visibles = []
    for r in todas:
        referencia = (r["fecha_factura"] or r["creado_en"] or "")[:10]
        if f["cobertura"] and r["cobertura"] != f["cobertura"]:
            continue
        if f["anio"] and referencia[:4] != f["anio"]:
            continue
        if f["mes"] and referencia[5:7] != f["mes"].zfill(2):
            continue
        if f["estado"] == "sincobrar":
            if r["estado"] != "emitida":
                continue
        elif f["estado"]:
            if r["estado"] != f["estado"]:
                continue
        elif r["estado"] == "anulada":
            continue
        visibles.append(r)
    acumulado = {k: round(sum(r[k] or 0 for r in visibles if r["estado"] != "anulada"), 2) for k in ("exento", "gravado", "iva", "total")}
    anios = sorted({(r["fecha_factura"] or r["creado_en"] or "")[:4] for r in todas if (r["fecha_factura"] or r["creado_en"])} | {hoy()[:4]}, reverse=True)
    return render_template("facturacion.html", facturas=visibles, f=f, acumulado=acumulado, anios=anios, coberturas=coberturas_lista(),
                           estados=facturacion.ESTADOS, clases=facturacion.CLASES_ESTADO)


MESES_CORTOS = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]


@app.route("/administracion")
@requiere_permiso("facturacion")
def administracion():
    hoy_d = datetime.now().date()
    actual = (hoy_d.year, hoy_d.month)
    siguiente = (hoy_d.year + (hoy_d.month == 12), hoy_d.month % 12 + 1)
    pendientes = db.q("SELECT id, numero, cobertura, total, fecha_estimada_cobro, fec_tipo FROM facturas WHERE estado='emitida' ORDER BY fecha_estimada_cobro, id")
    columnas = {actual: [], siguiente: []}
    sin_fecha = 0
    for r in pendientes:
        if not r["fecha_estimada_cobro"]:
            sin_fecha += 1
            continue
        d = datetime.strptime(r["fecha_estimada_cobro"], "%Y-%m-%d").date()
        atraso = (hoy_d - d).days
        color = "atr5" if atraso > 5 else ("atr" if atraso > 0 else "ok")
        simbolo = "tent" if r["fec_tipo"] == "tentativa" else "conf"
        titulo = {"ok": "En fecha", "atr": "Cobro atrasado", "atr5": "Cobro atrasado más de 5 días"}[color] + " · " + \
                 ("fecha tentativa" if simbolo == "tent" else "fecha confirmada" if r["fec_tipo"] == "confirmada" else "fecha de cobro cargada")
        clave = max(actual, (d.year, d.month))
        columnas.setdefault(clave, []).append({"id": r["id"], "numero": r["numero"], "cobertura": r["cobertura"], "total": r["total"], "color": color,
                                               "simbolo": simbolo, "titulo": titulo, "fecha": f"{d.day:02d} {MESES_CORTOS[d.month - 1]}.", "orden": d})
    meses = []
    for clave in sorted(columnas):
        filas = sorted(columnas[clave], key=lambda x: (x["orden"], x["cobertura"]))
        nombre = MESES[clave[1] - 1].capitalize() + (f" {clave[0]}" if clave[0] != hoy_d.year else "")
        meses.append({"nombre": nombre, "total": round(sum(x["total"] or 0 for x in filas), 2), "filas": filas, "actual": clave == actual})
    return render_template("administracion.html", meses=meses, sin_fecha=sin_fecha, cantidad=sum(len(m["filas"]) for m in meses))


def facturacion_buscar(f):
    proto = {r["id"]: r for r in db.q("SELECT id, cobertura, n_afiliado, exento, lugar_recoleccion FROM protocolos")}
    facturados = facturacion.ids_facturados()
    ctx = facturacion.contexto()
    q = f["q"].strip().lower()
    cobertura = facturacion.sin_acentos(f["cobertura"])
    out = []
    for e in cargar_estudios():
        p = proto.get(e["protocolo_id"])
        if e["anulado"] or not p:
            continue
        if f["tipo"] and e["tipo"] != f["tipo"]:
            continue
        if f["flujo"] and not (e["solicita_ihq"] if f["flujo"] == "IHQ" else e["tipo"] == f["flujo"]):
            continue
        if f["clase"] == "biop" and e["tipo"] != "BP":
            continue
        if f["clase"] == "cito" and e["tipo"] not in ("CT", "PAP"):
            continue
        if f["clase"] == "ihq" and not e["solicita_ihq"]:
            continue
        if f["prioridad"] == "ihq" and not e["solicita_ihq"]:
            continue
        if f["prioridad"] == "te" and not (e["tecnicas_especiales"] or "").strip():
            continue
        if f["zona"] and (e["subcategoria"] or "") != f["zona"]:
            continue
        if f["procedencia"] and (p["lugar_recoleccion"] or "").strip() != f["procedencia"]:
            continue
        if cobertura and facturacion.sin_acentos(p["cobertura"]) != cobertura:
            continue
        fecha_e = facturacion.fecha_de_estudio(e)
        if f["desde"] and fecha_e < f["desde"]:
            continue
        if f["hasta"] and fecha_e > f["hasta"]:
            continue
        informado = not e["proxima"]
        if (f["informados"] == "si" and not informado) or (f["informados"] == "no" and informado):
            continue
        ya = facturados.get(e["id"])
        if (f["facturados"] == "no" and ya) or (f["facturados"] == "si" and not ya):
            continue
        if q and q not in " ".join(str(x or "") for x in (e["numero"], e["apellido"], e["nombre"], e["dni"], p["n_afiliado"])).lower():
            continue
        tarifa = facturacion.resolver(ctx, p["cobertura"], e["tipo"], e["sitio"], e["tipo_muestra"], fecha_e or hoy())
        iva = facturacion.iva_de(tarifa, p["exento"])
        if f["iva"] and iva != f["iva"]:
            continue
        cantidad = facturacion.cantidad_de(e["cantidad"])
        out.append({"id": e["id"], "protocolo_id": e["protocolo_id"], "numero": e["numero"], "fecha": fecha_e, "tipo": e["tipo"],
                    "paciente": (e["apellido"] or "") + (", " + e["nombre"] if e["nombre"] else ""), "afiliado": p["n_afiliado"] or "",
                    "cobertura": p["cobertura"] or "", "etiqueta": e["etiqueta"], "informado": informado, "factura_id": ya, "tarifa": tarifa["ok"],
                    "motivo": tarifa.get("motivo", ""), "codigo": tarifa.get("codigo", ""), "cantidad": cantidad,
                    "unitario": tarifa.get("importe"), "iva": iva,
                    "importe": facturacion.importe_linea(cantidad, tarifa["importe"]) if tarifa["ok"] else None})
    out.sort(key=lambda r: (orden_numero(r["numero"]), r["id"]))
    return out


@app.route("/facturar", methods=["GET", "POST"])
@requiere_permiso("facturacion")
def facturar():
    yo = usuario_actual()
    if request.method == "POST":
        ids = [int(x) for x in request.form.getlist("ids") if x.isdigit()]
        volver = url_for("facturar", **{k: v for k, v in request.form.items() if k in ("desde", "hasta", "cobertura", "procedencia", "iva", "flujo", "zona", "clase", "informados", "facturados", "prioridad", "q")})
        if not ids:
            flash("Marcá al menos un estudio para facturar.", "error")
            return redirect(volver)
        with db.bloqueo("facturas"):
            proto = {r["id"]: r for r in db.q("SELECT id, cobertura, n_afiliado, exento, lugar_recoleccion FROM protocolos")}
            ya = facturacion.ids_facturados()
            ctx = facturacion.contexto()
            elegidos = [e for e in cargar_estudios() if e["id"] in set(ids)]
            problemas, items, coberturas = [], [], set()
            if len(elegidos) != len(set(ids)):
                problemas.append("Algún estudio elegido ya no existe.")
            for e in elegidos:
                p = proto[e["protocolo_id"]]
                if e["anulado"]:
                    problemas.append(f"{e['numero']} · {e['etiqueta']}: el estudio está anulado.")
                    continue
                if e["id"] in ya:
                    problemas.append(f"{e['numero']} · {e['etiqueta']}: ya está en la factura N° {ya[e['id']]}.")
                    continue
                if not (p["cobertura"] or "").strip():
                    problemas.append(f"{e['numero']}: el protocolo no tiene cobertura.")
                    continue
                fecha_e = facturacion.fecha_de_estudio(e)
                tarifa = facturacion.resolver(ctx, p["cobertura"], e["tipo"], e["sitio"], e["tipo_muestra"], fecha_e or hoy())
                if not tarifa["ok"]:
                    problemas.append(f"{e['numero']} · {e['etiqueta']}: {tarifa['motivo']}.")
                    continue
                coberturas.add(facturacion.sin_acentos(p["cobertura"]))
                cantidad = facturacion.cantidad_de(e["cantidad"])
                items.append({"estudio_id": e["id"], "protocolo_id": e["protocolo_id"], "numero_protocolo": e["numero"],
                              "paciente": ((e["apellido"] or "") + (", " + e["nombre"] if e["nombre"] else ""))[:160], "afiliado": p["n_afiliado"] or "",
                              "fecha": fecha_e, "estudio": e["etiqueta"], "codigo": tarifa["codigo"], "descripcion": tarifa["descripcion"],
                              "cantidad": cantidad, "unitario": tarifa["importe"], "iva": facturacion.iva_de(tarifa, p["exento"]),
                              "importe": facturacion.importe_linea(cantidad, tarifa["importe"]), "cobertura": p["cobertura"]})
            if len(coberturas) > 1:
                problemas.append("Los estudios elegidos son de coberturas distintas: una factura es de una sola cobertura.")
            if problemas or not items:
                flash("No se generó la factura: " + " · ".join(problemas[:6] or ["no hay estudios válidos."]), "error")
                return redirect(volver)
            fid = db.ex("INSERT INTO facturas (estado, cobertura, alicuota, creado_por, creado_en) VALUES ('borrador',?,?,?,?)",
                        (items[0]["cobertura"], facturacion.ALICUOTA_IVA, yo["id"], db.ahora()))
            try:
                db.lote_ex(facturacion.sentencias_items(fid, items))
            except Exception:
                db.ex("DELETE FROM facturas WHERE id=?", (fid,))
                raise
            facturacion.recalcular(fid)
        db.auditar(yo["id"], None, "factura_nueva", f"Borrador {fid} · {items[0]['cobertura']} · {len(items)} ítems")
        flash(f"Se generó el borrador de factura con {len(items)} ítems.", "ok")
        return redirect(url_for("factura", fid=fid))
    claves = ("desde", "hasta", "cobertura", "procedencia", "iva", "flujo", "zona", "clase", "prioridad", "q")
    f = {k: request.args.get(k, "") for k in claves}
    f["tipo"] = ""
    buscado = "buscar" in request.args
    hoy_d = datetime.now()
    if not buscado:
        f["desde"] = hoy_d.replace(day=1).strftime("%Y-%m-%d")
        f["hasta"] = ((hoy_d.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)).strftime("%Y-%m-%d")
    f["informados"] = request.args.get("informados", "si")
    f["facturados"] = request.args.get("facturados", "no")
    filas = facturacion_buscar(f) if buscado else []
    procedencias = sorted({(r["lugar_recoleccion"] or "").strip() for r in db.q("SELECT DISTINCT lugar_recoleccion FROM protocolos")} - {""})
    zonas = sorted({r["subcategoria"] for r in db.q("SELECT DISTINCT subcategoria FROM estudios WHERE subcategoria IS NOT NULL AND subcategoria<>''")})
    return render_template("facturar.html", f=f, filas=filas[:500], total=len(filas), buscado=buscado, coberturas=coberturas_lista(),
                           tipos_iva=facturacion.TIPOS_IVA, procedencias=procedencias, zonas=zonas)


@app.route("/facturacion/<int:fid>")
@requiere_permiso("facturacion")
def factura(fid):
    return render_template("factura.html", f=factura_o_404(fid), items=items_de_factura(fid), estados=facturacion.ESTADOS,
                           clases=facturacion.CLASES_ESTADO, tipos_iva=facturacion.TIPOS_IVA, hoy=hoy(), tipos_fec=facturacion.TIPOS_FEC)


@app.route("/facturacion/<int:fid>/guardar", methods=["POST"])
@requiere_permiso("facturacion")
def factura_guardar(fid):
    f = factura_o_404(fid)
    if f["estado"] != "borrador":
        flash("Solo se puede editar una factura en borrador.", "error")
        return redirect(url_for("factura", fid=fid))
    errores, sentencias = [], []
    for it in items_de_factura(fid):
        cantidad = facturacion.numero(request.form.get(f"c_{it['id']}"), None)
        unitario = facturacion.numero(request.form.get(f"u_{it['id']}"), None)
        iva = request.form.get(f"i_{it['id']}", it["iva"])
        if cantidad is None or cantidad < 1 or cantidad != int(cantidad) or unitario is None or unitario < 0 or iva not in facturacion.TIPOS_IVA:
            errores.append(it["numero_protocolo"])
            continue
        sentencias.append(("UPDATE factura_items SET codigo=?, cantidad=?, unitario=?, iva=?, importe=? WHERE id=?",
                           (request.form.get(f"k_{it['id']}", it["codigo"]).strip()[:20], int(cantidad), unitario, iva,
                            facturacion.importe_linea(int(cantidad), unitario), it["id"])))
    if errores:
        flash("No se guardó: revisá cantidad e importe de " + ", ".join(dict.fromkeys(errores)) + ".", "error")
        return redirect(url_for("factura", fid=fid))
    if sentencias:
        db.lote_ex(sentencias)
    db.ex("UPDATE facturas SET observaciones=? WHERE id=?", (request.form.get("observaciones", "").strip()[:500] or None, fid))
    facturacion.recalcular(fid)
    flash("Factura guardada.", "ok")
    return redirect(url_for("factura", fid=fid))


@app.route("/facturacion/<int:fid>/quitar/<int:iid>", methods=["POST"])
@requiere_permiso("facturacion")
def factura_quitar(fid, iid):
    f = factura_o_404(fid)
    if f["estado"] != "borrador":
        flash("Solo se pueden quitar ítems de una factura en borrador.", "error")
    else:
        db.ex("DELETE FROM factura_items WHERE id=? AND factura_id=?", (iid, fid))
        facturacion.recalcular(fid)
        flash("Ítem quitado de la factura.", "ok")
    return redirect(url_for("factura", fid=fid))


@app.route("/facturacion/<int:fid>/emitir", methods=["POST"])
@requiere_permiso("facturacion")
def factura_emitir(fid):
    yo = usuario_actual()
    f = factura_o_404(fid)
    numero_factura = request.form.get("numero", "").strip()[:40]
    fecha_factura = request.form.get("fecha_factura", "") or hoy()
    fecha_estimada = request.form.get("fecha_estimada_cobro", "").strip()
    fec_tipo = request.form.get("fec_tipo", "tentativa")
    if fec_tipo not in facturacion.TIPOS_FEC:
        fec_tipo = "tentativa"
    if f["estado"] != "borrador":
        flash("La factura ya no está en borrador.", "error")
    elif not items_de_factura(fid):
        flash("La factura no tiene ítems.", "error")
    elif not numero_factura:
        flash("Indicá el número de factura para emitirla.", "error")
    elif not facturacion.fecha_valida(fecha_factura) or (fecha_estimada and not facturacion.fecha_valida(fecha_estimada)):
        flash("La fecha de factura o la fecha estimada de cobro no es válida.", "error")
    elif db.uno("SELECT 1 FROM facturas WHERE numero=? AND id<>? AND estado<>'anulada'", (numero_factura, fid)):
        flash(f"Ya existe otra factura con el número {numero_factura}.", "error")
    else:
        db.ex("UPDATE facturas SET estado='emitida', numero=?, fecha_factura=?, fecha_estimada_cobro=?, fec_tipo=?, emitida_por=?, emitida_en=? WHERE id=?",
              (numero_factura, fecha_factura, fecha_estimada or None, fec_tipo, yo["id"], db.ahora(), fid))
        db.auditar(yo["id"], None, "factura_emitida", f"{numero_factura} · {f['cobertura']} · {fecha_factura}")
        flash(f"Factura {numero_factura} emitida.", "ok")
    return redirect(url_for("factura", fid=fid))


@app.route("/facturacion/<int:fid>/fec", methods=["POST"])
@requiere_permiso("facturacion")
def factura_fec(fid):
    yo = usuario_actual()
    f = factura_o_404(fid)
    fecha_estimada = request.form.get("fecha_estimada_cobro", "").strip()
    fec_tipo = request.form.get("fec_tipo", "tentativa")
    if f["estado"] != "emitida":
        flash("La fecha estimada de cobro se cambia solo en una factura emitida.", "error")
    elif fecha_estimada and not facturacion.fecha_valida(fecha_estimada):
        flash("La fecha estimada de cobro no es válida.", "error")
    elif fec_tipo not in facturacion.TIPOS_FEC:
        flash("Elegí el tipo de fecha.", "error")
    else:
        db.ex("UPDATE facturas SET fecha_estimada_cobro=?, fec_tipo=? WHERE id=?", (fecha_estimada or None, fec_tipo, fid))
        db.auditar(yo["id"], None, "factura_fec", f"{f['numero']} · {fecha_estimada or 'sin fecha'} · {fec_tipo}")
        flash("Fecha estimada de cobro actualizada.", "ok")
    return redirect(url_for("factura", fid=fid))


@app.route("/facturacion/<int:fid>/cobrar", methods=["POST"])
@requiere_permiso("facturacion")
def factura_cobrar(fid):
    yo = usuario_actual()
    f = factura_o_404(fid)
    fecha_cobro = request.form.get("fecha_cobro", "") or hoy()
    if f["estado"] != "emitida":
        flash("Solo se puede registrar el cobro de una factura emitida.", "error")
    elif not facturacion.fecha_valida(fecha_cobro):
        flash("La fecha de cobro no es válida.", "error")
    else:
        db.ex("UPDATE facturas SET estado='cobrada', fecha_cobro=?, cobrada_por=?, cobrada_en=? WHERE id=?", (fecha_cobro, yo["id"], db.ahora(), fid))
        db.auditar(yo["id"], None, "factura_cobrada", f"{f['numero']} · {fecha_cobro}")
        flash("Cobro registrado.", "ok")
    return redirect(url_for("factura", fid=fid))


@app.route("/facturacion/<int:fid>/anular", methods=["POST"])
@requiere_permiso("facturacion")
def factura_anular(fid):
    yo = usuario_actual()
    f = factura_o_404(fid)
    motivo = request.form.get("motivo", "").strip()[:200]
    if f["estado"] not in ("borrador", "emitida"):
        flash("Solo se puede anular una factura en borrador o emitida.", "error")
    elif not motivo:
        flash("Indicá el motivo de la anulación.", "error")
    else:
        db.ex("UPDATE facturas SET estado='anulada', anulada_por=?, anulada_en=?, anulada_motivo=? WHERE id=?", (yo["id"], db.ahora(), motivo, fid))
        db.auditar(yo["id"], None, "factura_anulada", f"{f['numero'] or 'borrador ' + str(fid)} · {motivo}")
        flash("Factura anulada: sus estudios vuelven a estar disponibles para facturar.", "ok")
    return redirect(url_for("factura", fid=fid))


@app.route("/facturacion/<int:fid>.csv")
@requiere_permiso("facturacion")
def factura_csv(fid):
    f = factura_o_404(fid)
    datos = facturacion.csv_de_factura(f, items_de_factura(fid)).encode("utf-8")
    nombre = re.sub(r"[^A-Za-z0-9_-]+", "_", f["numero"] or f"borrador_{fid}")
    return send_file(io.BytesIO(datos), mimetype="text/csv", as_attachment=True, download_name=f"factura_{nombre}.csv")


def dos_decimales(v):
    return f"{v:.2f}".replace(".", ",")


@app.route("/tarifario", methods=["GET", "POST"])
@requiere_permiso("facturacion")
def tarifario():
    yo = usuario_actual()
    tarifarios = db.q("SELECT * FROM tarifarios ORDER BY nombre")
    general = next((t for t in tarifarios if t["nombre"] == facturacion.GENERAL), tarifarios[0] if tarifarios else None)
    elegido = request.values.get("tarifario", "")
    actual = next((t for t in tarifarios if str(t["id"]) == elegido), general)
    ver = request.values.get("ver", "vigentes")
    destino = url_for("tarifario", tarifario=actual["id"], ver=ver)
    hoy_ = hoy()
    if request.method == "POST":
        accion = request.form.get("accion", "")
        if accion == "crear_tarifario":
            nombre = " ".join(request.form.get("nombre", "").split())[:80].upper()
            iva = request.form.get("iva", "exento")
            if not nombre or iva not in facturacion.TIPOS_IVA:
                flash("Indicá el nombre y el IVA de la tarifa.", "error")
            elif db.uno("SELECT 1 FROM tarifarios WHERE nombre=?", (nombre,)):
                flash("Ya existe una tarifa con ese nombre.", "error")
            else:
                nuevo_id = db.ex("INSERT INTO tarifarios (nombre, iva) VALUES (?,?)", (nombre, iva))
                db.auditar(yo["id"], None, "tarifario_nuevo", nombre)
                flash(f"Tarifa {nombre} creada: cargale los importes y asignale las coberturas.", "ok")
                destino = url_for("tarifario", tarifario=nuevo_id)
        elif accion == "guardar_tarifario":
            iva = request.form.get("iva", actual["iva"])
            if iva in facturacion.TIPOS_IVA:
                db.ex("UPDATE tarifarios SET iva=? WHERE id=?", (iva, actual["id"]))
            if actual["id"] != general["id"]:
                marcadas = set(request.form.getlist("coberturas")) & set(coberturas_lista())
                sentencias = [("DELETE FROM tarifario_coberturas WHERE tarifario_id=?", (actual["id"],))]
                sentencias += [("DELETE FROM tarifario_coberturas WHERE cobertura=?", (c,)) for c in marcadas]
                sentencias += [("INSERT INTO tarifario_coberturas (cobertura, tarifario_id) VALUES (?,?)", (c, actual["id"])) for c in marcadas]
                db.lote_ex(sentencias)
            db.auditar(yo["id"], None, "tarifario", f"{actual['nombre']}: IVA {iva}")
            flash("Tarifa guardada.", "ok")
        elif accion == "valor":
            codigo = request.form.get("codigo", "").strip()[:20]
            importe = facturacion.numero(request.form.get("importe"), None)
            desde = facturacion.fecha_flexible(request.form.get("vigente_desde")) or hoy_
            descripcion = request.form.get("descripcion", "").strip()[:300]
            if not codigo or importe is None or importe < 0:
                flash("Indicá el nomenclador y un importe válido.", "error")
            else:
                sentencias = []
                if descripcion:
                    sentencias += [("DELETE FROM nomencladores WHERE codigo=?", (codigo,)),
                                   ("INSERT INTO nomencladores (codigo, descripcion) VALUES (?,?)", (codigo, descripcion))]
                elif not db.uno("SELECT 1 FROM nomencladores WHERE codigo=?", (codigo,)):
                    sentencias.append(("INSERT INTO nomencladores (codigo, descripcion) VALUES (?,'')", (codigo,)))
                sentencias += [("DELETE FROM tarifario_valores WHERE tarifario_id=? AND codigo=? AND vigente_desde=?", (actual["id"], codigo, desde)),
                               ("INSERT INTO tarifario_valores (tarifario_id, codigo, importe, vigente_desde) VALUES (?,?,?,?)", (actual["id"], codigo, importe, desde))]
                db.lote_ex(sentencias)
                db.auditar(yo["id"], None, "tarifa_valor", f"{actual['nombre']} · {codigo} · {importe} desde {desde}")
                flash(f"Importe de {codigo} guardado, vigente desde {desde[8:]}/{desde[5:7]}/{desde[:4]}.", "ok")
        elif accion == "bloque":
            porcentaje = facturacion.numero(request.form.get("porcentaje"), None)
            desde = facturacion.fecha_flexible(request.form.get("vigente_desde")) or hoy_
            if porcentaje is None or porcentaje <= -100:
                flash("Indicá un porcentaje válido (por ejemplo 5 o 5,5).", "error")
            else:
                vigentes, vistos = [], set()
                for r in db.q("SELECT codigo, importe, vigente_desde FROM tarifario_valores WHERE tarifario_id=? AND vigente_desde<=? ORDER BY codigo, vigente_desde DESC", (actual["id"], desde)):
                    if r["codigo"] not in vistos:
                        vistos.add(r["codigo"])
                        vigentes.append(r)
                sentencias = []
                for r in vigentes:
                    nuevo_importe = round(r["importe"] * (1 + porcentaje / 100), 2)
                    sentencias += [("DELETE FROM tarifario_valores WHERE tarifario_id=? AND codigo=? AND vigente_desde=?", (actual["id"], r["codigo"], desde)),
                                   ("INSERT INTO tarifario_valores (tarifario_id, codigo, importe, vigente_desde) VALUES (?,?,?,?)", (actual["id"], r["codigo"], nuevo_importe, desde))]
                for i in range(0, len(sentencias), 40):
                    db.lote_ex(sentencias[i:i + 40])
                db.auditar(yo["id"], None, "tarifa_bloque", f"{actual['nombre']}: {porcentaje:+}% desde {desde} ({len(vigentes)} importes)")
                flash(f"Se actualizaron {len(vigentes)} importes ({porcentaje:+g} %), vigentes desde el {desde[8:]}/{desde[5:7]}/{desde[:4]}.", "ok")
        elif accion == "importar":
            nuevos, errores = facturacion.leer_valores(request.form.get("texto", ""), hoy_)
            if errores:
                flash("No se importó nada: " + " · ".join(errores[:6]) + (f" · y {len(errores) - 6} más" if len(errores) > 6 else ""), "error")
            elif not nuevos:
                flash("No hay líneas para importar.", "error")
            else:
                sentencias = []
                for n in nuevos:
                    if n["descripcion"]:
                        sentencias += [("DELETE FROM nomencladores WHERE codigo=?", (n["codigo"],)),
                                       ("INSERT INTO nomencladores (codigo, descripcion) VALUES (?,?)", (n["codigo"], n["descripcion"]))]
                    sentencias += [("DELETE FROM tarifario_valores WHERE tarifario_id=? AND codigo=? AND vigente_desde=?", (actual["id"], n["codigo"], n["vigente_desde"])),
                                   ("INSERT INTO tarifario_valores (tarifario_id, codigo, importe, vigente_desde) VALUES (?,?,?,?)", (actual["id"], n["codigo"], n["importe"], n["vigente_desde"]))]
                for i in range(0, len(sentencias), 40):
                    db.lote_ex(sentencias[i:i + 40])
                db.ex("INSERT INTO nomencladores (codigo, descripcion) SELECT DISTINCT codigo, '' FROM tarifario_valores WHERE codigo NOT IN (SELECT codigo FROM nomencladores)")
                db.auditar(yo["id"], None, "tarifas_importadas", f"{actual['nombre']}: {len(nuevos)} importes")
                flash(f"Se importaron {len(nuevos)} importes.", "ok")
        elif accion == "borrar_valor" and request.form.get("id", "").isdigit():
            db.ex("DELETE FROM tarifario_valores WHERE id=? AND tarifario_id=?", (int(request.form["id"]), actual["id"]))
            flash("Importe borrado.", "ok")
        return redirect(destino)
    filas = db.q("""SELECT v.id, v.codigo, v.importe, v.vigente_desde, n.descripcion FROM tarifario_valores v
                    LEFT JOIN nomencladores n ON n.codigo=v.codigo WHERE v.tarifario_id=? ORDER BY v.codigo, v.vigente_desde DESC""", (actual["id"],))
    if ver != "todas":
        vistos, vigentes = set(), []
        for r in filas:
            if r["codigo"] in vistos or r["vigente_desde"] > hoy_:
                continue
            vistos.add(r["codigo"])
            vigentes.append(r)
        filas = vigentes
    asignadas = {r["cobertura"]: r["tarifario_id"] for r in db.q("SELECT cobertura, tarifario_id FROM tarifario_coberturas")}
    nombres = {t["id"]: t["nombre"] for t in tarifarios}
    return render_template("tarifario.html", tarifarios=tarifarios, actual=actual, general=general, ver="todas" if ver == "todas" else "vigentes", filas=filas, hoy=hoy_,
                           coberturas=coberturas_lista(), asignadas=asignadas, nombres=nombres, tipos_iva=facturacion.TIPOS_IVA,
                           nomencladores=db.q("SELECT codigo, descripcion FROM nomencladores ORDER BY codigo"), dos=dos_decimales)


@app.route("/nomencladores", methods=["GET", "POST"])
@requiere_permiso("facturacion")
def nomencladores():
    yo = usuario_actual()
    if request.method == "POST":
        accion = request.form.get("accion", "")
        if accion == "guardar_nomenclador":
            codigo = request.form.get("codigo", "").strip()[:20]
            descripcion = request.form.get("descripcion", "").strip()[:300]
            if not codigo:
                flash("Indicá el código del nomenclador.", "error")
            else:
                db.lote_ex([("DELETE FROM nomencladores WHERE codigo=?", (codigo,)), ("INSERT INTO nomencladores (codigo, descripcion) VALUES (?,?)", (codigo, descripcion))])
                db.auditar(yo["id"], None, "nomenclador", f"{codigo} · {descripcion}")
                flash("Nomenclador guardado.", "ok")
        elif accion == "borrar_nomenclador":
            codigo = request.form.get("codigo", "")
            if db.uno("SELECT 1 FROM tarifario_valores WHERE codigo=?", (codigo,)) or db.uno("SELECT 1 FROM nomenclador_muestras WHERE codigo=?", (codigo,)):
                flash("No se puede borrar: tiene importes o tipos de muestra asociados.", "error")
            else:
                db.ex("DELETE FROM nomencladores WHERE codigo=?", (codigo,))
                flash("Nomenclador borrado.", "ok")
        elif accion == "asignar":
            tipo = request.form.get("tipo", "")
            tipo_muestra = " ".join(request.form.get("tipo_muestra", "").split())[:80]
            sitio = " ".join(request.form.get("sitio", "").split())[:80]
            codigo = request.form.get("codigo", "").strip()
            if tipo not in facturacion.TIPOS_ESTUDIO or not db.uno("SELECT 1 FROM nomencladores WHERE codigo=?", (codigo,)):
                flash("Elegí el tipo de estudio y un nomenclador que exista.", "error")
            else:
                db.lote_ex([("DELETE FROM nomenclador_muestras WHERE tipo=? AND sitio=? AND tipo_muestra=?", (tipo, sitio, tipo_muestra)),
                            ("INSERT INTO nomenclador_muestras (tipo, sitio, tipo_muestra, codigo) VALUES (?,?,?,?)", (tipo, sitio, tipo_muestra, codigo))])
                db.auditar(yo["id"], None, "nomenclador_muestra", f"{tipo} · {tipo_muestra or 'cualquier muestra'} · {sitio or 'cualquier sitio'} → {codigo}")
                flash("Asignación guardada.", "ok")
        elif accion == "borrar_asignacion" and request.form.get("id", "").isdigit():
            db.ex("DELETE FROM nomenclador_muestras WHERE id=?", (int(request.form["id"]),))
            flash("Asignación borrada.", "ok")
        elif accion == "importar":
            nuevas, errores = facturacion.leer_asignaciones(request.form.get("texto", ""))
            existentes = {r["codigo"] for r in db.q("SELECT codigo FROM nomencladores")}
            errores += [f"El nomenclador {n['codigo']} no existe." for n in nuevas if n["codigo"] not in existentes]
            if errores:
                flash("No se importó nada: " + " · ".join(dict.fromkeys(errores[:6])), "error")
            elif not nuevas:
                flash("No hay líneas para importar.", "error")
            else:
                sentencias = []
                for n in nuevas:
                    sentencias += [("DELETE FROM nomenclador_muestras WHERE tipo=? AND sitio=? AND tipo_muestra=?", (n["tipo"], n["sitio"], n["tipo_muestra"])),
                                   ("INSERT INTO nomenclador_muestras (tipo, sitio, tipo_muestra, codigo) VALUES (?,?,?,?)", (n["tipo"], n["sitio"], n["tipo_muestra"], n["codigo"]))]
                for i in range(0, len(sentencias), 40):
                    db.lote_ex(sentencias[i:i + 40])
                flash(f"Se importaron {len(nuevas)} asignaciones.", "ok")
        return redirect(url_for("nomencladores"))
    lista = db.q("SELECT codigo, descripcion FROM nomencladores ORDER BY codigo")
    uso = {r["codigo"]: r["n"] for r in db.q("SELECT codigo, COUNT(*) AS n FROM nomenclador_muestras GROUP BY codigo")}
    asignaciones = db.q("""SELECT m.id, m.tipo, m.sitio, m.tipo_muestra, m.codigo, n.descripcion FROM nomenclador_muestras m
                           LEFT JOIN nomencladores n ON n.codigo=m.codigo ORDER BY m.tipo, m.tipo_muestra, m.sitio""")
    cubiertos = {(r["tipo"], facturacion.sin_acentos(r["tipo_muestra"])) for r in asignaciones if not r["sitio"]}
    sin_asignar = [t for t in tipos_de_muestra_bp() if ("BP", facturacion.sin_acentos(t)) not in cubiertos]
    sitios = sorted({r["sitio"] for r in db.q("SELECT DISTINCT sitio FROM catalogo WHERE tipo='BP'")})
    return render_template("nomencladores.html", nomencladores=lista, uso=uso, asignaciones=asignaciones, tipos=facturacion.TIPOS_ESTUDIO,
                           sin_asignar=sin_asignar, tipos_muestra=tipos_de_muestra_bp(), sitios=sitios)


LISTADOS = {"tecnicas_especiales": ("Técnicas especiales", "Se tildan al agregar un estudio (biopsias y citologías) y se pueden elegir al informar la microscopía."),
            "ihq_marcadores": ("Marcadores de IHQ", "Se eligen al cargar la IHQ de una biopsia, cuando el diagnóstico la solicita.")}


def limpiar_listado(texto):
    return sin_repetidos(limpiar_valor(linea) for linea in (texto or "").splitlines())


@app.route("/listados", methods=["GET", "POST"])
@requiere_admin
def listados():
    yo = usuario_actual()
    if request.method == "POST":
        clave = request.form.get("lista", "")
        if clave in LISTADOS:
            items = limpiar_listado(request.form.get("valores", ""))
            db.lote_ex([("DELETE FROM listas WHERE nombre=?", (clave,))]
                       + [("INSERT INTO listas (nombre, orden, valor) VALUES (?,?,?)", (clave, i, v)) for i, v in enumerate(items)])
            db.auditar(yo["id"], None, "listado", f"{LISTADOS[clave][0]}: {len(items)} valores")
            flash(f"{LISTADOS[clave][0]}: {len(items)} valores guardados.", "ok")
        return redirect(url_for("listados") + f"#{clave}")
    actuales = listas()
    return render_template("listados.html", listados=[(k, t, d, actuales.get(k, [])) for k, (t, d) in LISTADOS.items()])


TIPOS_ETIQUETA = {"bp": "Recepción", "pap": "PAP-Laboratorio", "lab": "BP-Laboratorio"}
GRUPO_NUMERACION = ("bp", "pap")
ESTUDIO_LAB = {"lab": "BP", "pap": "PAP"}
LOTES_ETIQUETA_EXTRA = ["PAPS"]
MUESTRAS_PAP = ["EXO", "ENDO", "ENDO/EXO", "PAPURG", "CUPULA", "DERRAME"]
CODIGO_MUESTRA_PAP = {"citologia endocervical": "ENDO", "citologia exocervical": "EXO", "cupula vaginal": "CUPULA", "liquido": "DERRAME"}
MAX_PROTOCOLO = 999999
DIAS_ETIQUETADOS = 3


def codigo_muestra_pap(tipo_muestra, cantidad=""):
    if (cantidad or "").strip() == "1/2":
        return "ENDO/EXO"
    normal = unicodedata.normalize("NFD", tipo_muestra or "").encode("ascii", "ignore").decode().strip().lower()
    return CODIGO_MUESTRA_PAP.get(normal, (tipo_muestra or "").strip().upper())


def numero_protocolo(n):
    return f"C{n:06d}"


def estudio_de_lote(tipo_lote):
    return {"PAPS": "PAP", "CT": "CT"}.get(tipo_lote, "BP")


def numero_emitido(protocolo):
    m = re.fullmatch(r"C(\d{6})", protocolo or "")
    if not m:
        return True
    n = int(m.group(1))
    return db.uno("SELECT 1 FROM etiquetas_lotes WHERE tipo IN ('bp','pap') AND desde <= ? AND hasta >= ?", (n, n)) is not None


def personal(sector):
    filas = db.q("SELECT iniciales, nombre, sectores FROM usuarios WHERE activo=1 ORDER BY iniciales")
    return [{"iniciales": u["iniciales"], "nombre": u["nombre"] if (u["nombre"] or u["iniciales"]) != u["iniciales"] else ""}
            for u in filas if sector in (u["sectores"] or "").split(",")]


def tipos_lote_etiqueta():
    return sorted(set(tipos_lote_fijos()) | set(LOTES_ETIQUETA_EXTRA), key=str.casefold)


def listas_etiquetas():
    return {"tiposLote": tipos_lote_etiqueta(), "muestrasPap": MUESTRAS_PAP,
            "citotecnicos": personal("citotecnico"), "patologos": personal("firmante")}


def etiquetas_historial():
    progreso = {r["etiqueta_lote_id"]: r["ok"] for r in db.q(
        "SELECT etiqueta_lote_id, SUM(CASE WHEN borrador=0 THEN 1 ELSE 0 END) AS ok FROM protocolos "
        "WHERE etiqueta_lote_id IS NOT NULL GROUP BY etiqueta_lote_id")}
    out = []
    for f in db.q("""SELECT e.*, u.iniciales AS creador, l.codigo AS lote_codigo FROM etiquetas_lotes e
                     LEFT JOIN usuarios u ON u.id=e.creado_por LEFT JOIN lotes l ON l.id=e.lote_id ORDER BY e.id"""):
        out.append({"id": f["id"], "tipo": f["tipo"], "desde": f["desde"], "hasta": f["hasta"], "loteId": f["lote_id"],
                    "loteCodigo": f["lote_codigo"], "n": f["n"], "porProto": f["por_proto"], "etiquetas": f["etiquetas"],
                    "creadoEn": f["creado_en"], "creador": f["creador"] or "", "params": json.loads(f["params"]),
                    "completados": progreso.get(f["id"]) if f["tipo"] == "bp" else None})
    return out


def etiquetas_cantidad(cantidad):
    return int(cantidad) if (cantidad or "").isdigit() else (1 if cantidad else 0)


def tecnicas_por_estudio():
    out = {}
    for r in db.q("SELECT estudio_id, numero, cassettes FROM tecnicas ORDER BY id"):
        out.setdefault(r["estudio_id"], []).append(r)
    return out


def rotulos_cassettes(tecnicas):
    out = []
    for t in tecnicas:
        n = t["cassettes"] or 1
        numero = (t["numero"] or "").strip()
        out += [numero] if n == 1 else [f"{numero} {k}/{n}".strip() for k in range(1, n + 1)]
    return out


def etiquetas_casos():
    limite = (datetime.now() - timedelta(days=DIAS_ETIQUETADOS)).strftime("%Y-%m-%d")
    filas = db.q("""SELECT e.id, e.tipo, p.numero, p.apellido, p.nombre, e.cantidad, e.lote_id, l.codigo AS lote,
                           p.fecha_recoleccion, e.lab_etiquetado_en, e.tipo_muestra, u.iniciales AS resp
                    FROM estudios e JOIN protocolos p ON p.id=e.protocolo_id
                    LEFT JOIN usuarios u ON u.id=e.responsable_id LEFT JOIN lotes l ON l.id=e.lote_id
                    WHERE p.borrador=0 AND e.anulado=0 AND e.tipo IN ('BP','PAP') AND e.cantidad IS NOT NULL AND e.cantidad<>''
                      AND (e.lab_etiquetado_en IS NULL OR e.lab_etiquetado_en >= ?) ORDER BY e.lote_id, e.id""", (limite,))
    tecnicas = tecnicas_por_estudio()
    rotulos = {f["id"]: rotulos_cassettes(tecnicas.get(f["id"], [])) if f["tipo"] == "BP" else [] for f in filas}
    return [{"id": f["id"], "estudio": f["tipo"], "protocolo": f["numero"], "apellido": f["apellido"] or "",
             "nombre": f["nombre"] or "", "cantidad": len(rotulos[f["id"]]) or etiquetas_cantidad(f["cantidad"]),
             "nros": rotulos[f["id"]], "frascos": etiquetas_cantidad(f["cantidad"]), "loteId": f["lote_id"],
             "loteCodigo": f["lote"] or "", "fechaRec": f["fecha_recoleccion"] or "", "resp": f["resp"] or "",
             "muestra": codigo_muestra_pap(f["tipo_muestra"], f["cantidad"]) if f["tipo"] == "PAP" else "",
             "etiquetado": bool(f["lab_etiquetado_en"])}
            for f in sorted(filas, key=lambda f: (f["lote_id"] or 0, orden_numero(f["numero"]), f["id"]))]


def siguiente_libre():
    marcas = ",".join("?" * len(GRUPO_NUMERACION))
    fila = db.uno(f"SELECT MAX(hasta) AS m FROM etiquetas_lotes WHERE tipo IN ({marcas})", GRUPO_NUMERACION)
    return (fila["m"] or 0) + 1


def _entero(valor, minimo, maximo):
    try:
        n = int(valor)
    except (TypeError, ValueError):
        return None
    return n if minimo <= n <= maximo else None


def _lista_valida(valor, total, permitidos):
    return isinstance(valor, list) and len(valor) == total and all(v in permitidos for v in valor)


def datos_etiquetas():
    return {"listas": listas_etiquetas(), "historial": etiquetas_historial(), "casos": etiquetas_casos()}


@app.route("/etiquetas")
@requiere_permiso("etiquetas_recepcion", "etiquetas_lab")
def etiquetas():
    preset = None
    if request.args.get("lote"):
        lote = db.uno("SELECT * FROM lotes WHERE id=?", (int(request.args["lote"]),)) if request.args["lote"].isdigit() else None
        if not lote:
            flash("El lote no existe.", "error")
            return redirect(url_for("lotes"))
        if not puede("etiquetas_lab"):
            flash("No tenés permiso para las etiquetas de laboratorio.", "error")
            return redirect(url_for("lote", lote_id=lote["id"]))
        if lote["tipo_lote"] in tipos_lote("BP"):
            preset = {"tab": "lab", "loteId": lote["id"]}
        elif lote["tipo_lote"] in tipos_lote("PAP"):
            preset = {"tab": "pap", "loteId": lote["id"]}
        else:
            flash("Las etiquetas de laboratorio son para lotes de biopsias y de PAP.", "error")
            return redirect(url_for("lote", lote_id=lote["id"]))
    return render_template("etiquetas.html", datos={**datos_etiquetas(), "nombres": TIPOS_ETIQUETA, "preset": preset})


@app.route("/api/etiquetas/estado")
@requiere_permiso("etiquetas_recepcion", "etiquetas_lab")
def api_etiquetas_estado():
    return jsonify({"historial": etiquetas_historial(), "casos": etiquetas_casos()})


@app.route("/api/etiquetas/confirmar", methods=["POST"])
@requiere_permiso("etiquetas_recepcion", "etiquetas_lab")
def api_etiquetas_confirmar():
    yo = usuario_actual()
    d = request.get_json(silent=True) or {}
    p = d.get("params") if isinstance(d.get("params"), dict) else {}
    tipo = d.get("tipo")

    def error(texto, estado=400, **extra):
        return jsonify({"ok": False, "error": texto, **datos_etiquetas(), **extra}), estado

    if tipo not in TIPOS_ETIQUETA:
        return error("Tipo de etiqueta inválido.")
    if not puede("etiquetas_recepcion" if tipo == "bp" else "etiquetas_lab"):
        return error(f"No tenés permiso para etiquetas de {TIPOS_ETIQUETA[tipo]}.", 403)
    citos = {x["iniciales"] for x in personal("citotecnico")}
    desde = hasta = None
    borradores = etiquetados = None

    if tipo == "bp":
        n = _entero(d.get("n"), 1, 1000)
        desde = _entero(d.get("inicio"), 1, MAX_PROTOCOLO)
        if not n or not desde:
            return error("Cargá la cantidad de protocolos (1 a 1000) y el primer número.")
        hasta = desde + n - 1
        if hasta > MAX_PROTOCOLO:
            return error(f"El rango se pasa de {numero_protocolo(MAX_PROTOCOLO)}. Bajá la cantidad o el primer número.")
        try:
            datetime.strptime(str(p.get("fecha", "")), "%Y-%m-%d")
        except ValueError:
            return error("Elegí la fecha de recolección antes de confirmar.")
        if p.get("lote") not in tipos_lote_etiqueta():
            return error("Tipo de lote inválido.")
        paps = p["lote"] == "PAPS"
        if paps and p.get("cito") not in citos:
            return error("Elegí el citotécnico: con lote PAPS va en la segunda línea de cada etiqueta.")
        cito = p["cito"] if paps else ""
        params = {"fecha": p["fecha"], "lote": p["lote"], "cito": cito}
        est = REGISTRO[estudio_de_lote(p["lote"])]
        borradores = {"estudio": est.clave, "categoria": est.categoria, "subcategoria": est.subcategoria_fija,
                      "sitio": est.sitio_fijo, "tipo_lote": cito if paps else p["lote"], "fecha_lote": hoy(),
                      "fecha_rec": p["fecha"], "protocolos": [numero_protocolo(k) for k in range(desde, hasta + 1)],
                      "responsable_id": (db.uno("SELECT id FROM usuarios WHERE iniciales=?", (cito,))["id"] if paps else None)}
        por_proto, total = 2, n * 2
    else:
        ids = d.get("items")
        if not isinstance(ids, list) or not ids or len(ids) > 500 or len(set(ids)) != len(ids) or \
                not all(isinstance(i, int) for i in ids):
            return error("Marcá al menos un protocolo para etiquetar.")
        filas = {f["id"]: f for f in db.q(
            f"""SELECT e.*, p.numero, p.borrador, p.fecha_recoleccion, u.iniciales AS resp, l.codigo AS lote_codigo
                FROM estudios e JOIN protocolos p ON p.id=e.protocolo_id
                LEFT JOIN usuarios u ON u.id=e.responsable_id LEFT JOIN lotes l ON l.id=e.lote_id
                WHERE e.id IN ({','.join('?' * len(ids))})""", ids)}
        items = []
        tecnicas = tecnicas_por_estudio() if tipo == "lab" else {}
        for i in ids:
            f = filas.get(i)
            if not f or f["anulado"] or f["borrador"] or f["tipo"] != ESTUDIO_LAB[tipo] or not etiquetas_cantidad(f["cantidad"]):
                return error(f"El protocolo {f['numero'] if f else i} no está listo para {TIPOS_ETIQUETA[tipo]}: "
                             "tiene que estar completo (paciente y cantidad) y ser del estudio que corresponde.")
            nros = rotulos_cassettes(tecnicas.get(f["id"], [])) if tipo == "lab" else []
            items.append({"casoId": f["id"], "protocolo": f["numero"], "cantidad": len(nros) or etiquetas_cantidad(f["cantidad"]),
                          "loteCodigo": f["lote_codigo"] or "", "resp": f["resp"] or "", "fechaRec": f["fecha_recoleccion"] or "",
                          "muestra": codigo_muestra_pap(f["tipo_muestra"], f["cantidad"]) if tipo == "pap" else "", "nros": nros})
        n, total, por_proto = len(items), sum(i["cantidad"] for i in items), 1
        params = {"items": items}
        por_etiqueta = lambda campo: [i[campo] for i in items for _ in range(i["cantidad"])]
        if tipo == "lab":
            params["patos"] = por_etiqueta("resp")
            params["nros"] = [x for i in items for x in (i["nros"] + [""] * (i["cantidad"] - len(i["nros"])))]
        else:
            params["muestras"] = por_etiqueta("muestra")
        etiquetados = ids

    r = db.etiquetas_confirmar(tipo, GRUPO_NUMERACION, desde, hasta, n, por_proto, total,
                               json.dumps(params, ensure_ascii=False), yo["id"], borradores, etiquetados)
    if r.get("error") == "choque":
        f = r["fila"]
        return error(f"Ese rango se pisa con el lote de etiquetas N° {f['id']} ({TIPOS_ETIQUETA[f['tipo']]}), ya confirmado: "
                     f"{numero_protocolo(f['desde'])} a {numero_protocolo(f['hasta'])}. Cambiá el primer número: el siguiente "
                     f"libre es {numero_protocolo(siguiente_libre())}.", 409, conflicto=True)
    if r.get("error") == "existentes":
        ej = ", ".join(r["protocolos"][:3]) + ("…" if len(r["protocolos"]) > 3 else "")
        return error(f"Ya existen protocolos con esos números ({ej}), cargados a mano. "
                     "Elegí otro primer número.", 409)
    if r["lote"]:
        if r["lote"]["nuevo"]:
            db.auditar(yo["id"], None, "nuevo_lote", r["lote"]["codigo"])
        detalle = f"{numero_protocolo(desde)} a {numero_protocolo(hasta)} → lote {r['lote']['codigo']}"
    else:
        detalle = f"{n} protocolos"
    db.auditar(yo["id"], None, "etiquetas", f"{TIPOS_ETIQUETA[tipo]} {detalle} ({total} etiquetas)")
    return jsonify({"ok": True, "id": r["id"], **datos_etiquetas()})


@app.route("/completar")
@requiere_permiso("protocolo_crear")
def completar():
    lote_sel = request.args.get("lote", "")
    filtro = ("e.lote_id=?", (int(lote_sel),)) if lote_sel.isdigit() else ("1=1", ())
    protocolos = {}
    for e in cargar_estudios("p.borrador=1 AND " + filtro[0], filtro[1], borradores=True):
        protocolos.setdefault(e["protocolo_id"], {"id": e["protocolo_id"], "numero": e["numero"],
                                                  "fecha_recoleccion": e["fecha_recoleccion"], "estudios": []})["estudios"].append(e)
    lotes_con = db.q("""SELECT l.id, l.codigo, COUNT(DISTINCT p.id) AS n FROM estudios e JOIN protocolos p ON p.id=e.protocolo_id
                        JOIN lotes l ON l.id=e.lote_id WHERE p.borrador=1 GROUP BY l.id, l.codigo, l.fecha ORDER BY l.fecha DESC, l.codigo""")
    return render_template("completar.html", protocolos=sorted(protocolos.values(), key=lambda x: x["numero"]),
                           lotes_con=lotes_con, lote_sel=lote_sel)


@app.route("/exportar")
@requiere_permiso("exportar")
def exportar():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    etapas = {}
    for x in db.q("SELECT x.estudio_id, x.etapa, x.fecha_hora, u.iniciales FROM etapas x JOIN usuarios u ON u.id=x.usuario_id"):
        etapas.setdefault(x["estudio_id"], {})[x["etapa"]] = x
    macro = {r["estudio_id"]: r for r in db.q("SELECT * FROM macro")}
    micro = {r["estudio_id"]: r for r in db.q("SELECT * FROM micro")}
    ihq = {r["estudio_id"]: r for r in db.q("SELECT * FROM ihq")}
    protocolos = {r["id"]: r for r in db.q("SELECT * FROM protocolos")}
    orden_etapas = [k for k in NOMBRE_ETAPA if k != "ingreso"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estudios"
    cab = (["N° Protocolo", "Estudio", "Estado", "Lote", "Anulado", "Protocolo sistema", "Cargado en sistema", "Responsable",
            "Ingresado", "Fecha recolección", "DNI", "Apellido", "Nombre", "Sexo", "Fecha nacimiento", "Cobertura",
            "N° afiliado", "Exento", "E-mail", "Teléfono", "Médico", "Lugar recolección", "Lugar entrega",
            "Observaciones protocolo", "Categoría", "Subcategoría", "Sitio", "Tipo de muestra", "Cantidad",
            "Citología hormonal", "Observaciones estudio",
            "Template macro", "Descripción macro", "Cassettes", "Template micro", "Descripción micro", "Conclusión",
            "Bethesda", "Técnicas especiales", "Solicita IHQ", "Marcadores IHQ", "Resultado IHQ"]
           + [f"{NOMBRE_ETAPA[k]} - {x}" for k in orden_etapas for x in ("usuario", "fecha/hora")])
    ws.append(cab)
    cabecera = {}
    for r in db.q("SELECT id, protocolo_id, tipo FROM estudios WHERE anulado=0 ORDER BY id"):
        cabecera.setdefault((r["protocolo_id"], r["tipo"]), r["id"])
    for e in cargar_estudios():
        p = protocolos[e["protocolo_id"]]
        dueno = e["id"] if e["anulado"] else cabecera.get((e["protocolo_id"], e["tipo"]), e["id"])
        ma, mi, ih, et = macro.get(dueno), micro.get(dueno), ihq.get(dueno), etapas.get(e["id"], {})
        g = lambda r, k: (r[k] if r else "") or ""
        fila = [p["numero"], e["etiqueta"], e["estado"], e["lote"], "Si" if e["anulado"] else "", p["protocolo_sistema"],
                "Si" if p["cargado_sistema"] else "No", e["responsable"], fecha_hora(e["creado_en"]),
                fecha(p["fecha_recoleccion"]), p["dni"], p["apellido"], p["nombre"], p["sexo"], fecha(p["fecha_nacimiento"]),
                p["cobertura"], p["n_afiliado"], p["exento"], p["email"], p["telefono"], p["medico"], p["lugar_recoleccion"],
                p["lugar_entrega"], p["observaciones"], e["categoria"], e["subcategoria"], e["sitio"], e["tipo_muestra"],
                e["cantidad"], e["citologia_hormonal"], e["observaciones"],
                g(ma, "template"), g(ma, "descripcion"), g(ma, "cassettes"),
                g(mi, "template"), g(mi, "descripcion"), g(mi, "conclusion"), g(mi, "bethesda"),
                e["tecnicas_especiales"] or g(mi, "tecnicas_especiales"), "Si" if mi and mi["solicita_ihq"] else "", g(ih, "marcadores"), g(ih, "resultado")]
        for k in orden_etapas:
            x = et.get(k)
            fila += [x["iniciales"] if x else "", fecha_hora(x["fecha_hora"]) if x else ""]
        ws.append(fila)
    for i, cell in enumerate(ws[1], start=1):
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="5B2C83")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    db.auditar(usuario_actual()["id"], None, "exportar")
    return send_file(buf, as_attachment=True, download_name=f"contingencia_{datetime.now():%Y%m%d_%H%M}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def ips_locales():
    ips = set()
    try:
        for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET):
            ips.add(info[4][0])
    except OSError:
        pass
    return sorted(ip for ip in ips if not ip.startswith("127."))


if __name__ == "__main__":
    try:
        db.inicializar()
    except (RuntimeError, d1.ErrorD1) as error:
        print(f"No se puede iniciar el sistema: {error}")
        raise SystemExit(1)
    if os.environ.get("CONTINGENCIA_DESARROLLO") == "1":
        app.jinja_env.auto_reload = True
        recarga = os.environ.get("CONTINGENCIA_SIN_RECARGA") != "1"
        if os.environ.get("WERKZEUG_RUN_MAIN") == "true" or not recarga:
            print("=" * 64)
            print(" MODO DESARROLLO — " + ("recarga automática al guardar" if recarga else "con depurador, SIN recarga (reiniciar tras cambiar un .py)"))
            print(f" Abrir:             http://localhost:{PUERTO}")
            print(f" Base de datos:     {db.nombre_base()}")
            print("=" * 64)
        app.run(host="127.0.0.1", port=PUERTO, debug=True, use_reloader=recarga)
        raise SystemExit
    from waitress import serve
    print("=" * 64)
    print(" Sistema de contingencia CAP Vighi")
    print(f" En esta PC:        http://localhost:{PUERTO}")
    for ip in ips_locales():
        print(f" Desde otras PCs:   http://{ip}:{PUERTO}")
    print(f" Base de datos:     {db.nombre_base()}")
    print(" Para detenerlo: cerrar esta ventana.")
    print("=" * 64)
    serve(app, host="0.0.0.0", port=PUERTO, threads=12)
