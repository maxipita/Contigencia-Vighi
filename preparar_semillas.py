"""Genera los datos iniciales (seed/*.json) a partir del Excel de contingencia y los CSV de templates.

Uso:  py preparar_semillas.py
Se corre solo cuando cambian los templates, el catálogo o los usuarios del Excel.
"""
import csv
import json
import os
import warnings
from collections import defaultdict

import openpyxl

warnings.filterwarnings("ignore")
BASE = os.path.dirname(os.path.abspath(__file__))
ONEDRIVE = os.path.expanduser(r"~\OneDrive - Centro de diagnóstico Susana Vighi SRL")
EXCEL_LISTAS = os.path.join(ONEDRIVE, "Archivos de chat de Microsoft Teams", "plan_contingencia_CAPVighi.xlsx")
EXCEL_USUARIOS = os.path.join(ONEDRIVE, "Archivos de chat de Microsoft Teams", "plan_contingencia_CAPVighi_v3.xlsx")
CSV_DIR = os.path.join(ONEDRIVE, "Escritorio", "Automatización Web", "CM Pueyrredon")
SEED = os.path.join(BASE, "seed")
PENDIENTE = "Pendiente de definir"


def guardar(nombre, datos):
    os.makedirs(SEED, exist_ok=True)
    with open(os.path.join(SEED, nombre), "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=1)
    print(f"  {nombre}: {len(datos)}")


def valores_de_nombre(wb, nombre):
    """Valores de un rango con nombre (ej. SITIO_Biopsias_Digestiva)."""
    dn = wb.defined_names.get(nombre)
    if dn is None:
        return []
    out = []
    for hoja, ref in dn.destinations:
        for fila in wb[hoja][ref.replace("$", "")]:
            for c in fila:
                if c.value not in (None, ""):
                    out.append(str(c.value).strip())
    return out


def columna(ws, col):
    return [str(c.value).strip() for c in ws[col][1:] if c.value not in (None, "")]


def catalogo():
    wb = openpyxl.load_workbook(EXCEL_LISTAS)
    ws = wb["Listas"]
    rango_sitio = dict(zip(columna(ws, "Q"), columna(ws, "R")))   # "Cat|Sub" -> nombre de rango
    rango_tm = dict(zip(columna(ws, "AU"), columna(ws, "AV")))     # "Cat|Sub|Sitio" -> nombre de rango
    filas = []
    for tipo, cat, subs in [("BP", "Biopsias", columna(ws, "F")), ("CT", "Citologías", columna(ws, "J"))]:
        for sub in subs:
            sitios = valores_de_nombre(wb, rango_sitio.get(f"{cat}|{sub}", "")) or [PENDIENTE]
            for sitio in sitios:
                tms = valores_de_nombre(wb, rango_tm.get(f"{cat}|{sub}|{sitio}", "")) or [PENDIENTE]
                for tm in tms:
                    filas.append({"tipo": tipo, "categoria": cat, "subcategoria": sub, "sitio": sitio, "tipo_muestra": tm})
    for tm in columna(ws, "I"):
        filas.append({"tipo": "PAP", "categoria": "Citologías", "subcategoria": "Ginecológica", "sitio": "Vagina", "tipo_muestra": tm})
    guardar("catalogo.json", filas)
    guardar("listas.json", {
        "cobertura": columna(ws, "B"),
        "sexo": columna(ws, "C"),
        "tipo_lote": ["NO ONCO", "ENDO", "ONCO", "PAPURG", "TACOS", "HPM", "CT"],
        "bethesda": ["NILM", "ASC-US", "ASC-H", "LSIL", "HSIL", "Carcinoma escamoso", "AGC", "AIS",
                     "Adenocarcinoma", "Insatisfactoria"],
    })


def templates(archivo, clase, col_texto, col_conclusion=None):
    with open(os.path.join(CSV_DIR, archivo), encoding="utf-8-sig") as f:
        filas = list(csv.DictReader(f))
    vistos = defaultdict(list)
    for x in filas:
        texto = (x.get(col_texto) or "").strip()
        concl = (x.get(col_conclusion) or "").strip() if col_conclusion else ""
        if not texto and not concl:
            continue                                   # sin contenido: no sirve como template
        t = x["titulo"].strip()
        if (texto, concl) in [(a["texto"], a["conclusion"]) for a in vistos[t]]:
            continue                                   # duplicado idéntico
        vistos[t].append({"texto": texto, "conclusion": concl, "tipo_biopsia": (x.get("tipo_biopsia") or "").strip()})
    out = []
    for t, versiones in vistos.items():
        for i, v in enumerate(versiones, start=1):
            out.append({"clase": clase, "titulo": f"{t} ({i})" if len(versiones) > 1 else t, **v})
    out.sort(key=lambda z: z["titulo"].casefold())
    guardar(f"templates_{clase}.json", out)
    print(f"    ({len(filas) - sum(len(v) for v in vistos.values())} filas sin texto o repetidas)")


def templates_macro_desde_excel():
    """Los textos de macro buenos quedaron en la hoja oculta Templates_Macro del Excel v3."""
    ws = openpyxl.load_workbook(EXCEL_USUARIOS)["Templates_Macro"]
    out = []
    for titulo, tipo, texto in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        if titulo and texto:
            out.append({"clase": "macro", "titulo": str(titulo).strip(), "texto": str(texto).strip(),
                        "conclusion": "", "tipo_biopsia": (tipo or "").strip()})
    guardar("templates_macro.json", out)


def usuarios():
    ws = openpyxl.load_workbook(EXCEL_USUARIOS)["Usuarios"]
    sectores = ["ingreso", "laboratorio", "macroscopia", "traslados", "citotecnico", "firmante"]
    u = defaultdict(set)
    for fila in ws.iter_rows(min_row=3, max_col=6, values_only=True):
        for sector, v in zip(sectores, fila):
            if v and str(v).strip():
                u[str(v).strip().upper()].add(sector)
    guardar("usuarios.json", [{"iniciales": k, "sectores": sorted(v)} for k, v in sorted(u.items())])


if __name__ == "__main__":
    print("Generando semillas en", SEED)
    catalogo()
    templates_macro_desde_excel()
    templates("templates_micro.csv", "micro", "micro", "observaciones")
    usuarios()
