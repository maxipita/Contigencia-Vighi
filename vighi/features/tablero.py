import io
from datetime import datetime

from flask import Blueprint, render_template, request, send_file

from vighi.core import db
from vighi.core.sesion import puede, requiere_login, requiere_permiso, usuario_actual
from vighi.dominio.estudios import NOMBRE_ETAPA
from vighi.servicios.consultas import cargar_estudios
from vighi.servicios.exportacion import CABECERA, filas_exportables
from vighi.servicios.protocolos import ediciones_en_curso
from vighi.servicios.tablero import armar, pagina_pedida
from vighi.servicios.validacion import pendientes as pendientes_de_validar

bp = Blueprint("tablero", __name__)


def dia_valido(texto):
    try:
        datetime.strptime(texto, "%Y-%m-%d")
        return texto
    except ValueError:
        return ""


@bp.route("/")
@requiere_login
def tablero():
    yo = usuario_actual()
    f = {k: request.args.get(k, "") for k in ("q", "tipo", "ver", "sector", "etapa", "semaforo", "ubicacion", "mf", "desde", "hasta")}
    f["ver"] = f["ver"] or "todos"
    f["desde"], f["hasta"] = dia_valido(f["desde"]), dia_valido(f["hasta"])
    sectores = (yo["sectores"] or "").split(",")
    propios = yo["iniciales"] if "firmante" in sectores and not yo["admin"] and not puede("tablero_todos") else None
    datos = armar(f, sectores, pagina_pedida(request.args.get("pagina")), propios=propios)
    return render_template("tablero/tablero.html", f=f, etapas=[(k, n) for k, n in NOMBRE_ETAPA.items() if k != "ingreso"], **datos)


@bp.route("/completar")
@requiere_permiso("protocolo_crear")
def completar():
    lote_sel = request.args.get("lote", "")
    filtro = ("e.lote_id=?", (int(lote_sel),)) if lote_sel.isdigit() else ("1=1", ())
    protocolos = {}
    for e in cargar_estudios("p.borrador=1 AND " + filtro[0], filtro[1], borradores=True):
        protocolos.setdefault(e["protocolo_id"], {"id": e["protocolo_id"], "numero": e["numero"],
                                                  "fecha_recoleccion": e["fecha_recoleccion"], "estudios": []})["estudios"].append(e)
    lotes_con = db.q("""SELECT l.id, l.codigo, COUNT(DISTINCT p.id) AS n FROM estudios e JOIN protocolos p ON p.id=e.protocolo_id
                        JOIN lotes l ON l.id=e.lote_id WHERE p.borrador=1 GROUP BY l.id, l.codigo, l.fecha ORDER BY l.fecha DESC, l.codigo""")
    return render_template("tablero/completar.html", protocolos=sorted(protocolos.values(), key=lambda x: x["numero"]),
                           lotes_con=lotes_con, lote_sel=lote_sel, ediciones=ediciones_en_curso(), n_validar=len(pendientes_de_validar()))


@bp.route("/exportar")
@requiere_permiso("exportar")
def exportar():
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Estudios")
    ws.freeze_panes = "C2"
    cabecera = []
    for i, titulo in enumerate(CABECERA, start=1):
        celda = WriteOnlyCell(ws, value=titulo)
        celda.font, celda.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="5B2C83")
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        cabecera.append(celda)
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.append(cabecera)
    n = 1
    for fila in filas_exportables():
        ws.append(fila)
        n += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECERA))}{n}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    db.auditar(usuario_actual()["id"], None, "exportar")
    return send_file(buf, as_attachment=True, download_name=f"contingencia_{datetime.now():%Y%m%d_%H%M}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
