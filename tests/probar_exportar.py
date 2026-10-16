import io
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from datos_informe import preparar

A, db, ids, admin = preparar()
import openpyxl
from openpyxl.utils import get_column_letter
from vighi.servicios import consultas, exportacion

fallas = []


def chequear(n, ok, det=''):
    print(('  OK   ' if ok else '  FALLA'), n, '' if ok else det)
    if not ok:
        fallas.append(n)


with A.app.app_context():
    borrador = db.ex("INSERT INTO protocolos (numero, fecha_recoleccion, apellido, nombre, borrador, creado_por, creado_en) VALUES ('C009999','2026-10-01','Borrador','Caso',1,?,?)", (admin, db.ahora()))
    fl = db.ex("INSERT INTO flujos (protocolo_id, tipo, creado_en) VALUES (?,?,?)", (borrador, 'BP', db.ahora()))
    db.ex("INSERT INTO estudios (protocolo_id, flujo_id, tipo, categoria, sitio, tipo_muestra, cantidad, creado_por, creado_en) VALUES (?,?,?,?,?,?,?,?,?)",
          (borrador, fl, 'BP', 'Biopsias', 'Mama', 'Biopsia', '1', admin, db.ahora()))
    sin_permiso = db.ex("INSERT INTO usuarios (iniciales, nombre, sectores, admin, activo, debe_cambiar_clave) VALUES ('ZZSIN','Sin permiso','ingreso',0,1,0)")
    esperados = [dict(r) for r in db.q("SELECT p.numero, e.id FROM estudios e JOIN protocolos p ON p.id=e.protocolo_id WHERE p.borrador=0 ORDER BY p.id DESC, e.id")]
    de_servicio = consultas.cargar_estudios()
    auditorias_antes = db.uno("SELECT COUNT(*) AS n FROM auditoria WHERE accion='exportar'")['n']

c = A.app.test_client()
with c.session_transaction() as s:
    s['uid'] = admin


def descargar():
    r = c.get('/exportar')
    libro = openpyxl.load_workbook(io.BytesIO(r.get_data()))
    hoja = libro['Estudios']
    return r, hoja, [[celda.value for celda in fila] for fila in hoja.iter_rows()]


print('1) contenido del Excel')
r, hoja, filas = descargar()
chequear('responde con un .xlsx', r.status_code == 200 and 'spreadsheetml' in r.headers['Content-Type'] and 'contingencia_' in r.headers['Content-Disposition'], r.status_code)
chequear('la primera fila es el encabezado de siempre', filas[0] == exportacion.CABECERA, len(filas[0]))
datos = filas[1:]
chequear('una fila por estudio de protocolos que no son borradores', len(datos) == len(esperados) and len(datos) >= 8, (len(datos), len(esperados)))
chequear('el borrador no sale', 'C009999' not in [f[0] for f in datos])
chequear('mismo orden de siempre: protocolo más nuevo primero y los estudios por número', [f[0] for f in datos] == [e['numero'] for e in esperados])
chequear('el estudio, el estado, el lote y el responsable salen de la misma cuenta que usa el resto del sistema',
         all(f[1] == e['etiqueta'] and f[2] == e['estado'] and (f[3] or None) == e['lote'] and (f[7] or None) == e['responsable'] and (f[4] or '') == ('Si' if e['anulado'] else '')
             for f, e in zip(datos, de_servicio)))
chequear('hay estudios informados y pendientes', {'INFORMADO'} <= {f[2] for f in datos} and any(f[2].startswith('Pend.') for f in datos), sorted({f[2] for f in datos}))
with A.app.app_context():
    bp = db.uno("SELECT e.id, e.flujo_id, p.numero, p.apellido FROM estudios e JOIN protocolos p ON p.id=e.protocolo_id WHERE p.numero='C000101'")
    macro = db.uno("SELECT descripcion FROM macro WHERE flujo_id=?", (bp['flujo_id'],))
    micro = db.uno("SELECT descripcion, conclusion FROM micro WHERE flujo_id=?", (bp['flujo_id'],))
    etapas = {x['etapa']: x for x in db.q("SELECT x.etapa, x.fecha_hora, u.iniciales FROM etapas x JOIN usuarios u ON u.id=x.usuario_id WHERE x.flujo_id=?", (bp['flujo_id'],))}
fila = next(f for f in datos if f[0] == 'C000101')
col = lambda nombre: exportacion.CABECERA.index(nombre)
chequear('trae los datos del protocolo', fila[col('Apellido')] == bp['apellido'] and fila[col('Médico')] == 'Dr. Gustavo Kido' and fila[col('Cobertura')] == 'MEDICUS')
chequear('trae la macro y la micro del flujo', fila[col('Descripción macro')] == macro['descripcion'] and fila[col('Descripción micro')] == micro['descripcion'] and fila[col('Conclusión')] == micro['conclusion'])
con_etapa = [k for k in exportacion.ORDEN_ETAPAS if k in etapas]
chequear('hay etapas hechas', len(con_etapa) >= 3, con_etapa)
nombres = exportacion.NOMBRE_ETAPA
chequear('cada etapa hecha trae quién y cuándo', all(fila[col(f'{nombres[k]} - usuario')] == etapas[k]['iniciales'] and fila[col(f'{nombres[k]} - fecha/hora')] for k in con_etapa))
chequear('las etapas no hechas quedan vacías', all(not fila[col(f'{nombres[k]} - usuario')] for k in exportacion.ORDEN_ETAPAS if k not in etapas))
mixto = [f for f in datos if f[0] == 'C000106']
chequear('un protocolo con varios estudios repite lo de su flujo en cada uno', len(mixto) == 3 and mixto[0][col('Descripción micro')] != '' and sum(1 for f in mixto if f[col('Descripción micro')] == mixto[0][col('Descripción micro')]) == 2)

print('2) formato')
chequear('encabezado en negrita sobre violeta', hoja['A1'].font.bold and hoja['A1'].fill.fgColor.rgb.endswith('5B2C83') and hoja['A1'].alignment.wrap_text)
chequear('primeras dos columnas y encabezado fijos', hoja.freeze_panes == 'C2', hoja.freeze_panes)
chequear('filtro automático sobre todo el rango', hoja.auto_filter.ref == f'A1:{get_column_letter(len(exportacion.CABECERA))}{len(datos) + 1}', hoja.auto_filter.ref)
chequear('columnas de 16 de ancho', hoja.column_dimensions['A'].width == 16 and hoja.column_dimensions['B'].width == 16)

print('3) por tandas da lo mismo que de una vez')
for tanda in (1, 2, 3, 1000):
    exportacion.PROTOCOLOS_POR_TANDA = tanda
    _, _, otras = descargar()
    chequear(f'con {tanda} protocolo(s) por tanda el libro es idéntico', otras == filas, (len(otras), len(filas)))
exportacion.PROTOCOLOS_POR_TANDA = 250
with A.app.app_context():
    exportacion.PROTOCOLOS_POR_TANDA = 2
    rangos = list(exportacion.tandas())
    todos = [r['id'] for r in db.q("SELECT id FROM protocolos WHERE borrador=0 ORDER BY id DESC")]
    exportacion.PROTOCOLOS_POR_TANDA = 250
chequear('las tandas cubren todos los protocolos sin pisarse', [i for i in todos if any(d <= i <= h for d, h in rangos)] == todos and all(rangos[i][0] > rangos[i + 1][1] for i in range(len(rangos) - 1)), rangos)

print('4) permiso y auditoría')
with A.app.app_context():
    n = db.uno("SELECT COUNT(*) AS n FROM auditoria WHERE accion='exportar'")['n']
chequear('cada exportación queda en el historial', n == auditorias_antes + 5, (n, auditorias_antes))
otro = A.app.test_client()
with otro.session_transaction() as s:
    s['uid'] = sin_permiso
r = otro.get('/exportar')
chequear('sin el permiso no se descarga nada', r.status_code == 302 and 'spreadsheetml' not in r.headers.get('Content-Type', ''), r.status_code)

print()
print('FALLAS:', 'ninguna' if not fallas else fallas)
sys.exit(1 if fallas else 0)
