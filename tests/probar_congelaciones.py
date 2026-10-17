import io
import os
import re
import sys
from datetime import date, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from datos_informe import preparar

A, db, ids, admin = preparar()
from vighi.servicios import congelaciones as cg

fallas = []


def chequear(nombre, ok, detalle=""):
    print(("  OK   " if ok else "  FALLA"), nombre, "" if ok else detalle)
    if not ok:
        fallas.append(nombre)


def cliente(uid):
    c = A.app.test_client()
    with c.session_transaction() as s:
        s["uid"] = uid
    c.get("/")
    with c.session_transaction() as s:
        return c, s["csrf"]


def filas(sql, params=()):
    with A.app.app_context():
        return [dict(r) for r in db.q(sql, params)]


def una(sql, params=()):
    f = filas(sql, params)
    return f[0] if f else None


def avisos(html):
    return re.findall(r'<div class="aviso (?:ok|error)">([^<]*)</div>', html)


with A.app.app_context():
    db.ex("INSERT INTO ubicaciones (abreviatura, nombre) VALUES ('IADT', 'Instituto de Diagnóstico')")
    db.ex("INSERT INTO ubicaciones (abreviatura, nombre) VALUES ('SUIZO', 'Sanatorio Suizo Argentina')")
    perfil_ingreso = db.uno("SELECT id FROM perfiles WHERE nombre='Ingreso'")["id"]
    ingreso = db.ex("INSERT INTO usuarios (iniciales, nombre, sectores, admin, activo, debe_cambiar_clave, perfil_id) VALUES ('ZZIN','Ingreso Tester','ingreso',0,1,0,?)", (perfil_ingreso,))
    cobertura = db.q("SELECT valor FROM listas WHERE nombre='cobertura' ORDER BY orden")[0]["valor"]
    protocolo_ok = db.uno("SELECT id, numero FROM protocolos WHERE numero='C000101'")
    protocolo_otro_dni = db.ex("INSERT INTO protocolos (numero, fecha_recoleccion, apellido, nombre, dni, creado_por, creado_en) VALUES ('C000900','2026-10-01','Ajeno','Pac','99887766',?,?)", (admin, db.ahora()))
    diana, andrea, otto, cito = ids["med"], ids["sin"], ids["otro"], ids["cito"]
with A.app.app_context():
    for uid_ in (diana, andrea, otto):
        db.ex("UPDATE usuarios SET permisos_mas=CASE WHEN permisos_mas='' THEN ? ELSE permisos_mas || ',' || ? END WHERE id=?", ("congelaciones_realiza", "congelaciones_realiza", uid_))
    sin_permiso = db.uno("SELECT id FROM usuarios WHERE iniciales='FF'")["id"]

ing, csrf_ing = cliente(ingreso)
dia, csrf_dia = cliente(diana)
otto_c, csrf_otto = cliente(otto)
cito_c, csrf_cito = cliente(cito)


def datos(**cambios):
    base = dict(fecha="2026-03-10", hora="09:00", tipo="GI", ubicacion="IADT", patologo_id=str(diana), medico="Dr. Pérez", material="Mama",
                cobertura=cobertura, dni="30111222", apellido="Paciente", nombre="Uno", observaciones="", estado="agendada")
    base.update(cambios)
    return {k: v for k, v in base.items() if v is not None}


def alta(cliente_, csrf, **cambios):
    r = cliente_.post("/congelaciones/nueva", data={**datos(**cambios), "_csrf": csrf}, follow_redirects=True)
    return r, avisos(r.get_data(as_text=True))


def por_id(cid):
    return una("SELECT * FROM congelaciones WHERE id=?", (cid,))


def ultima():
    return una("SELECT * FROM congelaciones ORDER BY id DESC LIMIT 1")


print("1) reglas puras")
chequear("turnos: 12:00 es mañana, 12:30 mediodía, 16:00 mediodía, 16:30 tarde, 20:00 tarde, 20:30 noche, 06:00 mañana",
         [cg.turno(h) for h in ("12:00", "12:30", "16:00", "16:30", "20:00", "20:30", "06:00")] == ["Mañana", "Mediodía", "Mediodía", "Tarde", "Tarde", "Noche", "Mañana"])
chequear("las horas van de 06:00 a 23:30 cada media hora", cg.HORAS[0] == "06:00" and cg.HORAS[-1] == "23:30" and len(cg.HORAS) == 36)
with A.app.app_context():
    feriados = db.feriados()
laboral = date(2026, 10, 14)
chequear("el 14/10/2026 es día hábil", laboral.weekday() < 5 and laboral not in feriados)
chequear("en horario: miércoles hábil 08:00 sí, 07:30 no, 19:30 sí, 20:00 no",
         [cg.en_horario("2026-10-14", h, feriados) for h in ("08:00", "07:30", "19:30", "20:00")] == [True, False, True, False])
chequear("sábado y domingo son fuera de horario", not cg.en_horario("2026-10-17", "10:00", feriados) and not cg.en_horario("2026-10-18", "10:00", feriados))
feriado = sorted(feriados)[0]
chequear("un feriado es fuera de horario", not cg.en_horario(feriado.isoformat(), "10:00", feriados))
chequear("lunes_de devuelve el lunes de la semana", cg.lunes_de(date(2026, 10, 18)) == date(2026, 10, 12) and cg.lunes_de(date(2026, 10, 12)) == date(2026, 10, 12))
octubre = cg.semanas_del_mes(2026, 10)
chequear("las semanas de octubre 2026 empiezan el lunes 28/9 y terminan el 26/10", octubre[0] == date(2026, 9, 28) and octubre[-1] == date(2026, 10, 26) and len(octubre) == 5, str(octubre))
chequear("diciembre cierra bien el año", cg.semanas_del_mes(2026, 12)[-1] == date(2026, 12, 28))
chequear("cuando escribe el día de la semana", cg.cuando("2026-10-14") == "Miércoles 14 de octubre")

print("2) permisos")
r = cito_c.get("/congelaciones")
chequear("cualquiera con sesión ve la agenda", r.status_code == 200)
r = cito_c.get("/congelaciones/nueva", follow_redirects=True)
chequear("sin permiso no se puede agendar", "No tenés permiso" in r.get_data(as_text=True))
r, a = alta(cito_c, csrf_cito)
chequear("y tampoco crear por POST", not filas("SELECT id FROM congelaciones") and any("No tenés permiso" in x for x in a), str(a))
chequear("el perfil Ingreso puede agendar y Médico firmante confirmar", "congelaciones" in una("SELECT permisos FROM perfiles WHERE nombre='Ingreso'")["permisos"].split(",")
         and "congelaciones_confirmar" in una("SELECT permisos FROM perfiles WHERE nombre='Médico firmante'")["permisos"].split(","))

print("2b) quién aparece como patólogo")
with A.app.app_context():
    habilitados = {u["iniciales"] for u in cg.patologos()}
chequear("solo los médicos firmantes con el permiso «hace congelaciones»", {"DM", "AFH", "ZZOT"} <= habilitados and "FF" not in habilitados and "ZZCT" not in habilitados and "ZZIN" not in habilitados, str(habilitados))
r, a = alta(ing, csrf_ing, patologo_id=str(sin_permiso), hora="06:00")
chequear("un firmante sin el permiso no se puede elegir", any("Elegí el patólogo" in x for x in a), str(a))
chequear("ni aparece en el formulario", f'<option value="{sin_permiso}"' not in ing.get("/congelaciones/nueva").get_data(as_text=True))

print("3) alta y validaciones")
r, a = alta(ing, csrf_ing, estado="confirmada")
primera = ultima()
chequear("Ingreso agenda una congelación", primera and primera["estado"] == "confirmada" and primera["apellido"] == "Paciente" and primera["patologo_id"] == diana and any("agendada para el 2026-03-10" in x for x in a), str(a))
chequear("la lleva a la semana de la fecha", "Marzo" in r.get_data(as_text=True) and 'value="2026-03-09"' in r.get_data(as_text=True))
chequear("queda en la auditoría", any("alta 2026-03-10 09:00 GI en IADT" in x["detalle"] for x in filas("SELECT detalle FROM auditoria WHERE accion='congelacion'")))
errores = {
    "fecha": (dict(fecha="2026-02-30"), "Elegí la fecha"),
    "hora": (dict(hora="05:00"), "Elegí la hora"),
    "tipo": (dict(tipo="ZZ"), "Elegí el tipo"),
    "ubicación inexistente": (dict(ubicacion="NO EXISTE"), "Elegí la ubicación"),
    "patólogo que no es firmante": (dict(patologo_id=str(cito)), "Elegí el patólogo de la lista"),
    "cobertura inventada": (dict(cobertura="NO EXISTE"), "Elegí la cobertura"),
    "dni con letras": (dict(dni="12a45"), "solo números"),
    "dni todo ceros": (dict(dni="00000000"), "solo números"),
    "sin apellido": (dict(apellido=""), "Falta el apellido"),
    "texto largo": (dict(observaciones="x" * 1001), "demasiado largo"),
}
antes = len(filas("SELECT id FROM congelaciones"))
for nombre, (cambio, texto) in errores.items():
    r, a = alta(ing, csrf_ing, **cambio)
    chequear(f"rechaza {nombre}", any(texto in x for x in a), str(a))
chequear("ninguna de las inválidas se guardó", len(filas("SELECT id FROM congelaciones")) == antes)
r, a = alta(ing, csrf_ing, hora="10:00", estado="agendada", patologo_id="", dni="", nombre="", cobertura="", medico="", material="", apellido="SinPatologo")
sin_datos = ultima()
chequear("alcanzan fecha, hora, tipo, ubicación y apellido", sin_datos and sin_datos["patologo_id"] is None and sin_datos["dni"] is None, str(a))
chequear("el estado inicial inválido cae en «agendada»", (alta(ing, csrf_ing, hora="11:00", estado="eliminada", patologo_id="", apellido="EstadoRaro")[0] and ultima()["estado"] == "agendada"))
r, a = alta(ing, csrf_ing, dni="30.111.222", hora="12:30", apellido="ConPuntos", nombre="Dni")
chequear("el DNI se guarda sin puntos", (una("SELECT dni FROM congelaciones WHERE apellido='ConPuntos'") or {}).get("dni") == "30111222", str(a))
chequear("el DNI sin puntos se pasa a numérico", cg.limpiar_dni("30.111.222") == "30111222" and cg.limpiar_dni(" 30-111 222") == "30111222")

print("4) choques de agenda")
r, a = alta(ing, csrf_ing, hora="09:00", apellido="Otra")
chequear("el patólogo no puede tener dos congelaciones a la misma hora", any("ya tiene la congelación" in x for x in a) and len(filas("SELECT id FROM congelaciones WHERE hora='09:00' AND fecha='2026-03-10'")) == 1, str(a))
r, a = alta(ing, csrf_ing, hora="09:00", patologo_id=str(andrea), apellido="Otra")
chequear("otro patólogo a la misma hora sí", any("agendada para" in x for x in a))
segunda = ultima()
r, a = alta(ing, csrf_ing, hora="09:30", apellido="Cancelable")
cancelable = ultima()
r = ing.post(f"/congelaciones/{cancelable['id']}/estado", data={"accion": "eliminar", "_csrf": csrf_ing}, follow_redirects=True)
chequear("eliminar la deja eliminada", por_id(cancelable["id"])["estado"] == "eliminada")
r, a = alta(ing, csrf_ing, hora="09:30", apellido="Nueva en el mismo horario")
chequear("una eliminada libera el horario", any("agendada para" in x for x in a), str(a))
r = ing.post(f"/congelaciones/{cancelable['id']}/estado", data={"accion": "restaurar", "_csrf": csrf_ing}, follow_redirects=True)
chequear("no se puede restaurar si el horario ya lo ocupó otra", por_id(cancelable["id"])["estado"] == "eliminada" and any("ya tiene la congelación" in x for x in avisos(r.get_data(as_text=True))))
db_nueva = ultima()
r = ing.post(f"/congelaciones/{db_nueva['id']}/estado", data={"accion": "eliminar", "_csrf": csrf_ing}, follow_redirects=True)
chequear("se elimina la que se cargó por error (queda en el historial)", por_id(db_nueva["id"])["estado"] == "eliminada")
r = ing.post(f"/congelaciones/{cancelable['id']}/estado", data={"accion": "restaurar", "_csrf": csrf_ing}, follow_redirects=True)
chequear("y entonces sí se puede restaurar (vuelve a Agendada)", por_id(cancelable["id"])["estado"] == "agendada")

print("5) edición y reprogramación")
cid = primera["id"]
r = ing.post(f"/congelaciones/{cid}/editar", data={**datos(material="Mama derecha", estado=None), "_csrf": csrf_ing}, follow_redirects=True)
chequear("editar el material no cambia el estado", por_id(cid)["material"] == "Mama derecha" and por_id(cid)["estado"] == "confirmada")
r = ing.post(f"/congelaciones/{cid}/editar", data={**datos(hora="14:00", estado=None), "_csrf": csrf_ing}, follow_redirects=True)
chequear("cambiar la hora de una confirmada la devuelve a «A confirmar»", por_id(cid)["estado"] == "agendada" and por_id(cid)["hora"] == "14:00" and "vuelve a A confirmar" in " ".join(avisos(r.get_data(as_text=True))))
r = ing.post(f"/congelaciones/{cid}/editar", data={**datos(hora="14:00", ubicacion="NO EXISTE", estado=None), "_csrf": csrf_ing}, follow_redirects=True)
chequear("una edición inválida muestra el error y no guarda", "Elegí la ubicación" in r.get_data(as_text=True) and por_id(cid)["ubicacion"] == "IADT")
r = ing.post(f"/congelaciones/{cid}/editar", data={**datos(hora="09:00", patologo_id=str(andrea), estado=None), "_csrf": csrf_ing}, follow_redirects=True)
chequear("tampoco se puede pisar el horario de otra congelación del patólogo", por_id(cid)["hora"] == "14:00" and any("ya tiene la congelación" in x for x in avisos(r.get_data(as_text=True))))
r = ing.post(f"/congelaciones/{cid}/estado", data={"accion": "confirmar", "_csrf": csrf_ing}, follow_redirects=True)
chequear("confirmar la deja confirmada", por_id(cid)["estado"] == "confirmada")
r = ing.post(f"/congelaciones/{cid}/estado", data={"accion": "confirmar", "_csrf": csrf_ing}, follow_redirects=True)
chequear("confirmar de nuevo avisa y no hace nada", any("no se puede pasar" in x for x in avisos(r.get_data(as_text=True))))
r = ing.post(f"/congelaciones/{cid}/estado", data={"accion": "inventada", "_csrf": csrf_ing})
chequear("una acción desconocida es un error 400", r.status_code == 400)
r = cito_c.post(f"/congelaciones/{cid}/editar", data={**datos(material="Hackeado", estado=None), "_csrf": csrf_cito}, follow_redirects=True)
chequear("sin permiso no se edita", por_id(cid)["material"] != "Hackeado")

print("6) confirmar y eliminar: el patólogo y el permiso especial")
r, a = alta(ing, csrf_ing, fecha="2026-03-25", hora="07:30", tipo="PUN", ubicacion="SUIZO", patologo_id=str(otto), apellido="De Otto")
de_otto = ultima()
r = otto_c.post(f"/congelaciones/{de_otto['id']}/estado", data={"accion": "confirmar", "_csrf": csrf_otto}, follow_redirects=True)
chequear("un patólogo sin permisos confirma la suya", por_id(de_otto["id"])["estado"] == "confirmada")
r, a = alta(ing, csrf_ing, fecha="2026-03-25", hora="08:00", tipo="QR", ubicacion="SUIZO", patologo_id=str(diana), apellido="De Diana")
de_diana = ultima()
r = otto_c.post(f"/congelaciones/{de_diana['id']}/estado", data={"accion": "confirmar", "_csrf": csrf_otto}, follow_redirects=True)
chequear("pero no la de otro", por_id(de_diana["id"])["estado"] == "agendada" and any("No tenés permiso" in x for x in avisos(r.get_data(as_text=True))))
r = dia.post(f"/congelaciones/{de_diana['id']}/estado", data={"accion": "confirmar", "_csrf": csrf_dia}, follow_redirects=True)
chequear("la suya la confirma Diana", por_id(de_diana["id"])["estado"] == "confirmada")
r = dia.post(f"/congelaciones/{de_otto['id']}/estado", data={"accion": "eliminar", "_csrf": csrf_dia}, follow_redirects=True)
chequear("con el permiso de confirmar, un firmante elimina la de cualquiera", por_id(de_otto["id"])["estado"] == "eliminada")
r = dia.post(f"/congelaciones/{de_otto['id']}/estado", data={"accion": "restaurar", "_csrf": csrf_dia}, follow_redirects=True)
chequear("pero restaurar requiere permiso de agenda", por_id(de_otto["id"])["estado"] == "eliminada")
r = ing.post(f"/congelaciones/{de_otto['id']}/estado", data={"accion": "restaurar", "_csrf": csrf_ing}, follow_redirects=True)
r = ing.post(f"/congelaciones/{de_otto['id']}/estado", data={"accion": "confirmar", "_csrf": csrf_ing}, follow_redirects=True)
chequear("restaurar y confirmar de nuevo", por_id(de_otto["id"])["estado"] == "confirmada")

print("7) vincular con un protocolo")
r = ing.post(f"/congelaciones/{cid}/protocolo", data={"numero": "C999999", "_csrf": csrf_ing}, follow_redirects=True)
chequear("un protocolo que no existe", por_id(cid)["protocolo_id"] is None and any("No encontré el protocolo" in x for x in avisos(r.get_data(as_text=True))))
r = ing.post(f"/congelaciones/{cid}/protocolo", data={"numero": "c000900", "_csrf": csrf_ing}, follow_redirects=True)
chequear("un protocolo con otro DNI se rechaza", por_id(cid)["protocolo_id"] is None and any("no coincide" in x for x in avisos(r.get_data(as_text=True))))
r = ing.post(f"/congelaciones/{cid}/protocolo", data={"numero": "c000101", "_csrf": csrf_ing}, follow_redirects=True)
chequear("el número se acepta en minúscula y vincula", por_id(cid)["protocolo_id"] == protocolo_ok["id"])
ing.post(f"/congelaciones/{cancelable['id']}/estado", data={"accion": "eliminar", "_csrf": csrf_ing})
r = ing.post(f"/congelaciones/{cancelable['id']}/protocolo", data={"numero": "C000101", "_csrf": csrf_ing}, follow_redirects=True)
chequear("una eliminada no se vincula a un protocolo", por_id(cancelable["id"])["protocolo_id"] is None and any("restaurala" in x for x in avisos(r.get_data(as_text=True))))
ing.post(f"/congelaciones/{cancelable['id']}/estado", data={"accion": "restaurar", "_csrf": csrf_ing})
chequear("la vinculación queda en la auditoría del protocolo", any("se vinculó al protocolo C000101" in x["detalle"] for x in filas("SELECT detalle FROM auditoria WHERE protocolo_id=?", (protocolo_ok["id"],))))
r = ing.get(f"/congelaciones/{cid}")
html = r.get_data(as_text=True)
chequear("la ficha muestra el protocolo, la edición y el historial", r.status_code == 200 and "C000101" in html and "Guardar cambios" in html and "Historial" in html and "vinculó al protocolo" in html)
r = cito_c.get(f"/congelaciones/{cid}")
chequear("sin permiso la ficha es de solo lectura", "Guardar cambios" not in r.get_data(as_text=True) and "<fieldset class=\"tarjeta\" disabled>" in r.get_data(as_text=True))
r = ing.post(f"/congelaciones/{cid}/protocolo", data={"numero": "", "_csrf": csrf_ing}, follow_redirects=True)
chequear("se desvincula con el número vacío", por_id(cid)["protocolo_id"] is None)
r = ing.post(f"/congelaciones/{cid}/protocolo", data={"numero": "C000101", "_csrf": csrf_ing}, follow_redirects=True)

print("8) agenda semanal")
ing.post(f"/congelaciones/{cancelable['id']}/estado", data={"accion": "eliminar", "_csrf": csrf_ing})
r = ing.get("/congelaciones?semana=2026-03-09&anio=2026&mes=3")
html = r.get_data(as_text=True)
chequear("la semana del 9 al 15 de marzo muestra las congelaciones", r.status_code == 200 and "Paciente, Uno" in html and "SinPatologo" in html and "Martes 10" in html and "9 - 15" in html, html[:200])
chequear("las eliminadas no aparecen en la agenda", "Cancelable" not in html)
chequear("una confirmada vinculada muestra el protocolo", "Confirmada · C000101" in html)
chequear("la semana siguiente muestra la de Otto y la de Diana (25/3)", "De Otto" in ing.get("/congelaciones?semana=2026-03-23&anio=2026&mes=3").get_data(as_text=True))
chequear("fuera de horario se marca", "Fuera de horario" in ing.get("/congelaciones?semana=2026-03-23&anio=2026&mes=3").get_data(as_text=True))
chequear("hay un botón para agendar en cada día con permiso", "congelaciones/nueva?fecha=2026-03-10" in html)
chequear("sin permiso no hay botón de agendar", "congelaciones/nueva?fecha" not in cito_c.get("/congelaciones?semana=2026-03-09&anio=2026&mes=3").get_data(as_text=True))
chequear("el contador de la semana cuenta las no canceladas", re.search(r"(\d+) congelaciones? en la semana \((\d+) confirmadas?\)", html) is not None)
h = ing.get("/congelaciones?semana=2026-03-23&anio=2026&mes=3&ubicacion=SUIZO").get_data(as_text=True)
chequear("filtra por procedencia", "De Diana" in h and "De Otto" in h and "Paciente, Uno" not in h)
chequear("y otra procedencia no trae las de SUIZO", "De Diana" not in ing.get("/congelaciones?semana=2026-03-23&anio=2026&mes=3&ubicacion=IADT").get_data(as_text=True))
h = ing.get(f"/congelaciones?semana=2026-03-23&anio=2026&mes=3&patologo={otto}").get_data(as_text=True)
chequear("filtra por patólogo", "De Otto" in h and "De Diana" not in h)
h = ing.get("/congelaciones?semana=2026-03-09&anio=2026&mes=3&patologo=sin").get_data(as_text=True)
chequear("filtra los sin patólogo", "SinPatologo" in h and "Paciente, Uno" not in h)
h = ing.get("/congelaciones?semana=2026-03-23&anio=2026&mes=3&tipo=PUN").get_data(as_text=True)
chequear("filtra por tipo", "De Otto" in h and "De Diana" not in h)
h = ing.get("/congelaciones?semana=2026-03-23&anio=2026&mes=3&tipo=PUN&tipo=QR").get_data(as_text=True)
chequear("varios tipos a la vez", "De Otto" in h and "De Diana" in h)
h = ing.get("/congelaciones?semana=2026-03-23&anio=2026&mes=3&horario=fuera").get_data(as_text=True)
chequear("filtra fuera de horario", "De Otto" in h and "De Diana" not in h)
h = ing.get("/congelaciones?semana=2026-03-09&anio=2026&mes=3&vinculadas=si").get_data(as_text=True)
chequear("filtra las vinculadas", "Paciente, Uno" in h and "SinPatologo" not in h)
h = ing.get("/congelaciones?semana=2026-03-09&anio=2026&mes=3&vinculadas=no").get_data(as_text=True)
chequear("y las que no", "Paciente, Uno" not in h and "SinPatologo" in h)
h = ing.get("/congelaciones?semana=2026-03-09&anio=2026&mes=3&estado=eliminada").get_data(as_text=True)
chequear("el estado «eliminada» las muestra", "Cancelable" in h and "Paciente, Uno" not in h)
h = ing.get("/congelaciones?semana=2026-03-09&anio=2026&mes=3&estado=confirmada").get_data(as_text=True)
chequear("filtra por estado", "Paciente, Uno" in h and "SinPatologo" not in h)
for malo in ("semana=basura", "anio=9999999&mes=40", "tipo=XX&horario=raro&vinculadas=tal&estado=tal"):
    chequear(f"parámetros inválidos no rompen ({malo})", ing.get("/congelaciones?" + malo).status_code == 200)
chequear("sin parámetros muestra la semana de hoy", ing.get("/congelaciones").status_code == 200)

print("9) tablero de congelaciones")
hoy = date.today()
r, a = alta(ing, csrf_ing, fecha=(hoy - timedelta(days=1)).isoformat(), hora="10:00", estado="confirmada", patologo_id=str(andrea), apellido="Ayer", dni="30111222")
ayer = ultima()
r, a = alta(ing, csrf_ing, fecha=(hoy + timedelta(days=1)).isoformat(), hora="10:00", estado="confirmada", patologo_id=str(andrea), apellido="Mañana")
manana = ultima()
r, a = alta(ing, csrf_ing, fecha=(hoy + timedelta(days=3)).isoformat(), hora="11:00", estado="agendada", patologo_id=str(andrea), apellido="PorConfirmar")
por_confirmar = ultima()
r, a = alta(ing, csrf_ing, fecha=(hoy - timedelta(days=1)).isoformat(), hora="12:00", estado="confirmada", patologo_id=str(andrea), apellido="YaVinculada")
vinculada = ultima()
ing.post(f"/congelaciones/{vinculada['id']}/protocolo", data={"numero": "C000102", "_csrf": csrf_ing})
with A.app.app_context():
    actividades = {c["apellido"]: c["actividad"] for c in cg.listar(cg.filtros_vacios(), con_eliminadas=True)}
chequear("actividades derivadas", actividades.get("Ayer") == "vincular" and actividades.get("Mañana") == "realizar" and actividades.get("PorConfirmar") == "confirmar"
         and actividades.get("YaVinculada") == "vinculada" and actividades.get("Cancelable") == "eliminada" and actividades.get("SinPatologo") == "alerta", str(actividades))
h = ing.get("/congelaciones?vista=tablero").get_data(as_text=True)
chequear("el tablero muestra por defecto las pendientes", "Ayer" in h and "Mañana" in h and "PorConfirmar" in h and "YaVinculada" not in h and "Cancelable" not in h)
chequear("la que falta vincular tiene el formulario rápido", "vincular-rapido" in h and "Falta vincular el protocolo" in h)
h_pendientes = ing.get("/congelaciones?vista=tablero").get_data(as_text=True)
orden = [list(cg.ACTIVIDADES).index(x) for x in re.findall(r"chip act-(\w+)", h_pendientes)]
chequear("primero lo que pide acción: alerta, confirmar, vincular y realizar", orden == sorted(orden) and 0 in orden and len(set(orden)) >= 3, str(orden))
h = ing.get("/congelaciones?vista=tablero&actividad=todas").get_data(as_text=True)
chequear("todas incluye las vinculadas y eliminadas", "YaVinculada" in h and "Cancelable" in h)
h = ing.get("/congelaciones?vista=tablero&actividad=vinculada").get_data(as_text=True)
chequear("el firmante del protocolo vinculado aparece en MF", "YaVinculada" in h and "C000102" in h)
with A.app.app_context():
    firmantes = cg.firmantes_de([una("SELECT id FROM protocolos WHERE numero='C000102'")["id"], None])
chequear("MF es el firmante responsable del protocolo", list(firmantes.values()) == ["AFH"], str(firmantes))
h = ing.get(f"/congelaciones?vista=tablero&actividad=todas&patologo={andrea}").get_data(as_text=True)
chequear("el tablero filtra por patólogo", "Ayer" in h and "De Otto" not in h and "Paciente, Uno" not in h)
h = ing.get("/congelaciones?vista=tablero&actividad=inventada").get_data(as_text=True)
chequear("una actividad inválida vuelve a las pendientes", "Ayer" in h and "YaVinculada" not in h)
chequear("vincular desde el tablero", ing.post(f"/congelaciones/{ayer['id']}/protocolo", data={"numero": "C000101", "_csrf": csrf_ing}, headers={"Referer": "http://localhost/congelaciones?vista=tablero"}).status_code == 302
         and por_id(ayer["id"])["protocolo_id"] == protocolo_ok["id"])
chequear("y vuelve al tablero", ing.post(f"/congelaciones/{ayer['id']}/protocolo", data={"numero": "", "_csrf": csrf_ing}, headers={"Referer": "http://localhost/congelaciones?vista=tablero"}).headers["Location"].endswith("vista=tablero"))

print("10) dashboard")
h = ing.get("/congelaciones?vista=dashboard&anio=2026").get_data(as_text=True)
with A.app.app_context():
    tabla, meses, total = cg.dashboard(cg.filtros_vacios(), 2026, "confirmadas")
por = {t["patologo"]: t for t in tabla}
nombre_diana = una("SELECT COALESCE(nombre, iniciales) AS n FROM usuarios WHERE id=?", (diana,))["n"]
chequear("cuenta las confirmadas por patólogo y mes", por[nombre_diana]["meses"][2] == 2 and sum(por["Otto Tester"]["meses"]) == 1 and por["Otto Tester"]["meses"][2] == 1, str(por))
chequear("los totales por mes y el general suman", meses[2] == sum(t["meses"][2] for t in tabla) and total == sum(t["total"] for t in tabla))
with A.app.app_context():
    tabla_todas, _, total_todas = cg.dashboard(cg.filtros_vacios(), 2026, "todas")
chequear("«Confirmadas y a confirmar» suma además las agendadas y nunca las eliminadas", total_todas > total and
         sum(t["total"] for t in tabla_todas) == len([c for c in filas("SELECT estado FROM congelaciones WHERE fecha LIKE '2026-%'") if c["estado"] != "eliminada"]))
with A.app.app_context():
    tabla_fuera, _, total_fuera = cg.dashboard({**cg.filtros_vacios(), "patologo": str(otto), "horario": "fuera"}, 2026, "confirmadas")
    _, _, fuera_diana = cg.dashboard({**cg.filtros_vacios(), "patologo": str(diana), "horario": "fuera"}, 2026, "confirmadas")
    _, _, en_diana = cg.dashboard({**cg.filtros_vacios(), "patologo": str(diana), "horario": "en"}, 2026, "confirmadas")
chequear("filtra fuera de horario", [t["patologo"] for t in tabla_fuera] == ["Otto Tester"] and total_fuera == 1 and fuera_diana == 0 and en_diana == 2, str((tabla_fuera, fuera_diana, en_diana)))
with A.app.app_context():
    tabla_tipo, _, _ = cg.dashboard({**cg.filtros_vacios(), "tipos": ["PUN"]}, 2026, "todas")
chequear("filtra por tipo", [t["patologo"] for t in tabla_tipo] == ["Otto Tester"])
chequear("la pantalla del dashboard responde y muestra los patólogos", "Otto Tester" in h and "Total" in h and ">Mar<" in h)
chequear("un año sin datos lo dice", "No hay congelaciones en 2030" in ing.get("/congelaciones?vista=dashboard&anio=2030").get_data(as_text=True))
chequear("los parámetros inválidos del dashboard no rompen", ing.get("/congelaciones?vista=dashboard&anio=x&contar=y").status_code == 200)

print("11) agenda del patólogo")
h = dia.get("/congelaciones?vista=patologo").get_data(as_text=True)
chequear("un patólogo entra directo a su agenda (sin las pasadas)", "Elegí un patólogo para ver su agenda" not in h and "De Diana" not in h and "Paciente, Uno" not in h)
h = dia.get("/congelaciones?vista=patologo&pasadas=1").get_data(as_text=True)
chequear("con «incluir anteriores» aparecen las suyas, no las de otros", "De Diana" in h and "De Otto" not in h)
chequear("tiene botones de confirmar y rechazar", "Rechazar" in h and "Confirmar" in h)
h = ing.get("/congelaciones?vista=patologo").get_data(as_text=True)
chequear("quien no es patólogo tiene que elegir uno", "Elegí un patólogo para ver su agenda" in h)
h = ing.get(f"/congelaciones?vista=patologo&patologo={andrea}").get_data(as_text=True)
chequear("se ve la agenda de otro patólogo", "Mañana" in h and "PorConfirmar" in h and "Ayer" not in h)
chequear("con conteos por estado y procedencia", "Agendada (2)" in h and "Confirmada (1)" in h and "Todas (3)" in h and "IADT (3)" in h, h[h.find("agenda-chips"):h.find("agenda-chips") + 600])
h = ing.get(f"/congelaciones?vista=patologo&patologo={andrea}&estado=agendada").get_data(as_text=True)
chequear("filtra por estado", "PorConfirmar" in h and "Mañana" not in h)
r = ing.post(f"/congelaciones/{por_confirmar['id']}/estado", data={"accion": "eliminar", "_csrf": csrf_ing}, headers={"Referer": f"http://localhost/congelaciones?vista=patologo&patologo={andrea}"}, follow_redirects=True)
chequear("rechazar desde la agenda la elimina y vuelve a la agenda", por_id(por_confirmar["id"])["estado"] == "eliminada" and "vista=patologo" in r.request.url)
h = ing.get(f"/congelaciones?vista=patologo&patologo={andrea}").get_data(as_text=True)
chequear("las eliminadas ya no están en su agenda", "PorConfirmar" not in h)

print("12) descarga a Excel y consulta de paciente")
import openpyxl

r = ing.get("/congelaciones/descargar?desde=2026-03-09&hasta=2026-03-15")
chequear("descarga un .xlsx", r.status_code == 200 and "spreadsheetml" in r.mimetype and "congelaciones_" in r.headers["Content-Disposition"])
hoja = openpyxl.load_workbook(io.BytesIO(r.get_data())).active
cabecera = [c.value for c in hoja[1]]
registros = [[c.value for c in fila] for fila in hoja.iter_rows(min_row=2)]
chequear("trae la cabecera y solo las de esa semana", cabecera[:4] == ["Fecha", "Hora", "Turno", "Horario"] and registros and all("2026-03-09" <= x[0] <= "2026-03-15" for x in registros), str(registros[:2]))
chequear("con el estado y el protocolo", any(x[4] == "Confirmada" and x[14] == "C000101" for x in registros), str(registros))
r = ing.get(f"/congelaciones/descargar?patologo={otto}&tipo=PUN")
hoja = openpyxl.load_workbook(io.BytesIO(r.get_data())).active
chequear("respeta los filtros", hoja.max_row == 2 and hoja.cell(2, 13).value == "De Otto")
chequear("queda en la auditoría de exportaciones", any("Congelaciones" in x["detalle"] for x in filas("SELECT detalle FROM auditoria WHERE accion='exportar'")))
r = ing.get("/api/congelaciones/paciente?dni=30.111.222")
chequear("el DNI trae el paciente de un protocolo", r.get_json().get("apellido") and r.get_json().get("cobertura") is not None, str(r.get_json()))
chequear("un DNI desconocido devuelve vacío", ing.get("/api/congelaciones/paciente?dni=11111111").get_json() == {})
chequear("un DNI vacío o de ceros también", ing.get("/api/congelaciones/paciente?dni=0000").get_json() == {} and ing.get("/api/congelaciones/paciente").get_json() == {})
chequear("sin sesión la API no responde", A.app.test_client().get("/api/congelaciones/paciente?dni=30111222").status_code == 302)

print("13) médicos según la ubicación")
with A.app.app_context():
    uno = db.ex("INSERT INTO medicos (nombre) VALUES ('Dr. Filtro Uno')")
    db.ex("INSERT INTO medicos (nombre) VALUES ('Dr. Filtro Dos')")
    iadt = db.uno("SELECT id FROM ubicaciones WHERE abreviatura='IADT'")["id"]
    db.ex("INSERT INTO ubicacion_medicos (ubicacion_id, medico_id) VALUES (?,?)", (iadt, uno))
de_iadt = ing.get("/api/medicos?ubicacion=IADT").get_json()
chequear("con una ubicación se ofrecen solo sus médicos", de_iadt == ["Dr. Filtro Uno"], str(de_iadt))
todos = ing.get("/api/medicos").get_json()
chequear("sin ubicación, todos", "Dr. Filtro Uno" in todos and "Dr. Filtro Dos" in todos and len(todos) > 2)
chequear("una ubicación sin médicos asociados ofrece todos", set(ing.get("/api/medicos?ubicacion=SUIZO").get_json()) == set(todos))
chequear("y una que no existe también", set(ing.get("/api/medicos?ubicacion=NOEXISTE").get_json()) == set(todos))
chequear("el formulario toma los médicos de la ubicación que se elige", 'data-medicos-de="ubicacion"' in ing.get("/congelaciones/nueva").get_data(as_text=True))
alta(ing, csrf_ing, fecha="2026-03-26", hora="13:00", ubicacion="SUIZO", medico="Dra. Suiza", patologo_id=str(andrea), apellido="MedicoSuizo")
with A.app.app_context():
    de_suizo, de_iadt_usados = cg.medicos_usados("SUIZO"), cg.medicos_usados("IADT")
chequear("el filtro de médicos de la agenda sigue la procedencia", "Dra. Suiza" in de_suizo and "Dra. Suiza" not in de_iadt_usados and "Dr. Pérez" in de_iadt_usados, str((de_suizo, de_iadt_usados)))
h = ing.get("/congelaciones?semana=2026-03-23&anio=2026&mes=3&ubicacion=IADT").get_data(as_text=True)
chequear("y la pantalla lo respeta", 'value="Dra. Suiza"' not in h and 'value="Dr. Pérez"' in h)

print("14) alerta de 24 horas")
from datetime import datetime

ahora_fija = datetime(2026, 3, 9, 10, 0)
limite = cg.limite_alerta(ahora_fija)
chequear("el límite es 24 horas después", limite == "2026-03-10 10:00", limite)


def fila_de(estado="agendada", fecha="2026-03-10", hora="09:30", protocolo=None):
    return {"estado": estado, "fecha": fecha, "hora": hora, "protocolo_id": protocolo}


chequear("agendada con menos de 24 hs: alerta", cg.en_alerta(fila_de(), limite))
chequear("justo a las 24 hs todavía no", not cg.en_alerta(fila_de(hora="10:00"), limite))
chequear("con más de 24 hs, no", not cg.en_alerta(fila_de(hora="10:30"), limite) and not cg.en_alerta(fila_de(fecha="2026-03-20"), limite))
chequear("una agendada que ya pasó sigue en alerta", cg.en_alerta(fila_de(fecha="2026-03-01"), limite))
chequear("confirmada o eliminada nunca", not cg.en_alerta(fila_de("confirmada"), limite) and not cg.en_alerta(fila_de("eliminada"), limite))
chequear("la actividad pasa a «alerta»", cg.actividad_de(fila_de(), "2026-03-09 10:00", limite) == "alerta" and cg.actividad_de(fila_de(hora="12:00"), "2026-03-09 10:00", limite) == "confirmar")
chequear("los textos de la alerta dicen cuánto falta", cg.texto_de_alerta(fila_de(hora="09:30"), ahora_fija) == "Sin confirmar: faltan 23 h"
         and cg.texto_de_alerta(fila_de(fecha="2026-03-09", hora="10:30"), ahora_fija) == "Sin confirmar: falta menos de 1 h"
         and cg.texto_de_alerta(fila_de(fecha="2026-03-08"), ahora_fija) == "Ya pasó la hora y no se confirmó")
ayer_iso, lejos_iso = (hoy - timedelta(days=1)).isoformat(), (hoy + timedelta(days=10)).isoformat()
with A.app.app_context():
    for apellido_, estado_, fecha_ in (("AlertaAyer", "agendada", ayer_iso), ("ConfirmadaAyer", "confirmada", ayer_iso), ("FuturoLejano", "agendada", lejos_iso)):
        db.ex("INSERT INTO congelaciones (fecha, hora, estado, tipo, ubicacion, patologo_id, apellido, creado_en) VALUES (?,?,?,?,?,?,?,?)", (fecha_, "10:00", estado_, "GI", "IADT", otto, apellido_, db.ahora()))
    alerta_id = db.uno("SELECT id FROM congelaciones WHERE apellido='AlertaAyer'")["id"]
h = otto_c.get("/").get_data(as_text=True)
m = re.search(r'title="Congelaciones tuyas sin confirmar[^"]*">(\d+)</span>', h)
chequear("el patólogo ve en el menú cuántas tiene sin confirmar", m and m.group(1) == "1", str(m))
chequear("quien no tiene ninguna no ve el número", "Congelaciones tuyas sin confirmar" not in cito_c.get("/").get_data(as_text=True))
h = ing.get(f"/congelaciones?semana={ayer_iso}&anio={hoy.year}&mes={hoy.month}&estado=alerta").get_data(as_text=True)
chequear("el filtro «En alerta» muestra solo las agendadas por vencer", "AlertaAyer" in h and "ConfirmadaAyer" not in h and "¡Sin confirmar!" in h)
chequear("el filtro está en la lista de estados", 'value="alerta"' in h)
h = ing.get("/congelaciones?vista=tablero").get_data(as_text=True)
chequear("el tablero la muestra primero, como alerta", "Alerta: sin confirmar" in h and h.index("AlertaAyer") < h.index("FuturoLejano"))
h = ing.get(f"/congelaciones/{alerta_id}").get_data(as_text=True)
chequear("la ficha avisa que no se confirmó", "Ya pasó la hora y no se confirmó" in h and "Le aparece al patólogo" in h)
h = otto_c.get("/congelaciones?vista=patologo").get_data(as_text=True)
chequear("la agenda del patólogo trae las vencidas aunque sean de ayer, con el aviso", "AlertaAyer" in h and "ConfirmadaAyer" not in h and "sin confirmar con menos de 24 hs" in h)
otto_c.post(f"/congelaciones/{alerta_id}/estado", data={"accion": "confirmar", "_csrf": csrf_otto})
chequear("al confirmarla desaparece el número del menú", "Congelaciones tuyas sin confirmar" not in otto_c.get("/").get_data(as_text=True))
h = otto_c.get("/congelaciones?vista=patologo").get_data(as_text=True)
chequear("y el aviso de la agenda", "sin confirmar con menos de 24 hs" not in h)

print("15) un patólogo al que se le quita el permiso")
with A.app.app_context():
    huerfana = db.ex("INSERT INTO congelaciones (fecha, hora, estado, tipo, ubicacion, patologo_id, apellido, creado_en) VALUES ('2026-05-05','10:00','confirmada','GI','IADT',?,'DeOtroFirmante',?)", (sin_permiso, db.ahora()))
h = ing.get(f"/congelaciones/{huerfana}").get_data(as_text=True)
chequear("la ficha de una congelación con un patólogo sin permiso lo sigue mostrando", f'<option value="{sin_permiso}" selected>' in h)
r = ing.post(f"/congelaciones/{huerfana}/editar", data={**datos(fecha="2026-05-05", hora="10:00", patologo_id=str(sin_permiso), material="Cambiado", estado=None), "_csrf": csrf_ing}, follow_redirects=True)
chequear("y se puede editar sin cambiarle el patólogo", por_id(huerfana)["material"] == "Cambiado" and por_id(huerfana)["patologo_id"] == sin_permiso)
h = ing.get("/congelaciones?semana=2026-05-04&anio=2026&mes=5").get_data(as_text=True)
chequear("los filtros lo ofrecen mientras tenga congelaciones", f'<option value="{sin_permiso}"' in h)

print("16) menú y migración de permisos")
h = ing.get("/").get_data(as_text=True)
chequear("el menú tiene Congelaciones", "Congelaciones" in h and "#i-copo" in h)
chequear("la ficha de una inexistente es 404", ing.get("/congelaciones/99999").status_code == 404)
with A.app.app_context():
    db.ex("INSERT INTO congelaciones (fecha, hora, estado, tipo, ubicacion, apellido, creado_en) VALUES ('2026-01-05','10:00','cancelada','GI','IADT','Vieja',?)", (db.ahora(),))
    db.inicializar()
    db.olvidar_cache()
    vieja = db.uno("SELECT estado FROM congelaciones WHERE apellido='Vieja'")["estado"]
chequear("una congelación «cancelada» de antes pasa a «eliminada»", vieja == "eliminada", vieja)
with A.app.app_context():
    db.ex("UPDATE perfiles SET permisos=REPLACE(REPLACE(permisos, ',congelaciones_confirmar', ''), ',congelaciones', '')")
    db.ex("DELETE FROM congelaciones")
    db.ex("DROP TABLE congelaciones")
    quedaron = [r["permisos"] for r in db.q("SELECT permisos FROM perfiles WHERE permisos LIKE '%congelaciones%'")]
    db.inicializar()
    db.olvidar_cache()
    ingreso_p = db.uno("SELECT permisos FROM perfiles WHERE nombre='Ingreso'")["permisos"].split(",")
    firmante_p = db.uno("SELECT permisos FROM perfiles WHERE nombre='Médico firmante'")["permisos"].split(",")
    existe = db.uno("SELECT COUNT(*) AS n FROM congelaciones")["n"]
chequear("sin la tabla, un arranque la crea y da los permisos a los perfiles que corresponden", not quedaron and "congelaciones" in ingreso_p and "congelaciones_confirmar" in firmante_p and existe == 0,
         f"{quedaron} {ingreso_p} {firmante_p}")
with A.app.app_context():
    db.inicializar()
    db.olvidar_cache()
    repetidos = [p for p in db.uno("SELECT permisos FROM perfiles WHERE nombre='Ingreso'")["permisos"].split(",") if p == "congelaciones"]
chequear("un segundo arranque no repite el permiso", len(repetidos) == 1)

print("FALLAS:", ", ".join(fallas) if fallas else "ninguna")
sys.exit(1 if fallas else 0)
